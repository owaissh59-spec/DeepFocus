"""
docx_style.py
=============
Styling / layout engine for the JKSSB Junior Pharmacist study-notes book.

Design goals
------------
* A4 page, 0.5 inch margins on every side.
* "Topper's handwritten book" feel: colour-coded heading hierarchy, ruled
  dividers, shaded definition boxes, mnemonic boxes, high-yield boxes,
  professional tables with shaded headers and alternating row bands.
* Tables use PROPORTIONAL column widths driven by the length of the text that
  each column has to carry, so text always fits and nothing is squeezed.
* Page breaks occur ONLY at the end of a chapter.
* A real Word TOC field is inserted so page numbers can be baked in.
"""

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Inches, Pt, RGBColor

# ---------------------------------------------------------------------------
# COLOUR PALETTE
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x0B, 0x31, 0x5E)        # part titles
BLUE = RGBColor(0x0F, 0x4C, 0x81)        # H1 / chapter titles
TEAL = RGBColor(0x0E, 0x6B, 0x6B)        # H2
GREEN = RGBColor(0x1B, 0x6E, 0x3C)       # H3
PURPLE = RGBColor(0x5B, 0x2C, 0x83)      # H4 / run-in heads
MAROON = RGBColor(0x8B, 0x1A, 0x1A)      # warnings / must-remember
ORANGE = RGBColor(0xA8, 0x4B, 0x00)      # exam-tip text
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Hex shading strings (no leading #)
SH_HEADER_BLUE = "0F4C81"
SH_HEADER_TEAL = "0E6B6B"
SH_HEADER_GREEN = "1B6E3C"
SH_HEADER_PURPLE = "5B2C83"
SH_BAND = "EDF3F9"        # alternating row band
SH_BAND_GREEN = "EDF6F0"
SH_DEF = "E8F1FA"         # definition box
SH_MNEMONIC = "FFF6DA"    # mnemonic box
SH_HIGHYIELD = "FDECEC"   # high-yield box
SH_NOTE = "F0EAF8"        # note / remember box
SH_COMPARE = "EAF6F6"     # comparison box

BODY_FONT = "Cambria"
HEAD_FONT = "Calibri"
MONO_FONT = "Courier New"

BODY_SIZE = Pt(10.5)
TABLE_SIZE = Pt(9)
TABLE_HEAD_SIZE = Pt(9)


# ---------------------------------------------------------------------------
# LOW-LEVEL XML HELPERS
# ---------------------------------------------------------------------------
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def shade(element, fill):
    """Apply solid shading to a paragraph's pPr, a cell's tcPr, or a row."""
    pr = element
    sd = _el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill})
    pr.append(sd)


def cell_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    tcPr.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))


def para_shade(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    for old in pPr.findall(qn("w:shd")):
        pPr.remove(old)
    pPr.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))


def para_borders(paragraph, edges=("top", "bottom", "left", "right"),
                 sz=6, color="0F4C81", val="single", space=4):
    pPr = paragraph._p.get_or_add_pPr()
    for old in pPr.findall(qn("w:pBdr")):
        pPr.remove(old)
    bdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        if edge in edges:
            bdr.append(_el(f"w:{edge}", **{"w:val": val, "w:sz": sz,
                                           "w:space": space, "w:color": color}))
    pPr.append(bdr)


def keep_with_next(paragraph, on=True):
    pPr = paragraph._p.get_or_add_pPr()
    for old in pPr.findall(qn("w:keepNext")):
        pPr.remove(old)
    k = OxmlElement("w:keepNext")
    if not on:
        k.set(qn("w:val"), "0")
    pPr.append(k)


def no_widow_control(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:widowControl"))


def set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    """Margins in twentieths of a point (dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:tcMar")):
        tcPr.remove(old)
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("start", left), ("bottom", bottom), ("end", right)):
        mar.append(_el(f"w:{tag}", **{"w:w": val, "w:type": "dxa"}))
    # legacy names for wider compatibility
    for tag, val in (("left", left), ("right", right)):
        mar.append(_el(f"w:{tag}", **{"w:w": val, "w:type": "dxa"}))
    tcPr.append(mar)


def cell_valign(cell, align="center"):
    cell.vertical_alignment = {
        "top": WD_ALIGN_VERTICAL.TOP,
        "center": WD_ALIGN_VERTICAL.CENTER,
        "bottom": WD_ALIGN_VERTICAL.BOTTOM,
    }[align]


def repeat_header_row(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(_el("w:tblHeader", **{"w:val": "true"}))


def row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


# ---------------------------------------------------------------------------
# DOCUMENT SET-UP
# ---------------------------------------------------------------------------
def new_document():
    doc = Document()
    _setup_page(doc)
    _setup_base_styles(doc)
    _enable_field_update(doc)
    return doc


def _setup_page(doc):
    for section in doc.sections:
        section.page_width = Cm(21.0)      # A4
        section.page_height = Cm(29.7)
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.header_distance = Inches(0.25)
        section.footer_distance = Inches(0.25)


def usable_width(doc):
    s = doc.sections[0]
    return Emu(int(s.page_width - s.left_margin - s.right_margin))


def _setup_base_styles(doc):
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = BODY_SIZE
    st.font.color.rgb = BLACK
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), BODY_FONT)
    rfonts.set(qn("w:hAnsi"), BODY_FONT)
    rfonts.set(qn("w:cs"), BODY_FONT)
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.06
    pf.widow_control = True

    # List Paragraph tightened
    try:
        lp = doc.styles["List Paragraph"]
        lp.font.name = BODY_FONT
        lp.font.size = BODY_SIZE
        lp.paragraph_format.space_after = Pt(2)
        lp.paragraph_format.line_spacing = 1.05
    except KeyError:
        pass


def _enable_field_update(doc):
    """Tell Word to refresh fields (TOC page numbers) when the file opens."""
    settings = doc.settings.element
    for old in settings.findall(qn("w:updateFields")):
        settings.remove(old)
    settings.append(_el("w:updateFields", **{"w:val": "true"}))


# ---------------------------------------------------------------------------
# RUN / TEXT HELPERS  (supports lightweight inline markup)
# ---------------------------------------------------------------------------
def _style_run(run, bold=False, italic=False, size=None, color=None,
               font=None, underline=False, small_caps=False, highlight=None):
    run.bold = bold
    run.italic = italic
    if underline:
        run.underline = True
    if size is not None:
        run.font.size = size
    if color is not None:
        run.font.color.rgb = color
    f = font or BODY_FONT
    run.font.name = f
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), f)
    rfonts.set(qn("w:hAnsi"), f)
    rfonts.set(qn("w:cs"), f)
    if small_caps:
        rpr.append(OxmlElement("w:smallCaps"))
    if highlight:
        rpr.append(_el("w:highlight", **{"w:val": highlight}))
    return run


import re as _re

#: Inline markup tokens. The parser is recursive, so markers may be nested -
#: e.g. ^^a key term with **an emphasised part** inside^^.
_MARKUP = _re.compile(r"(\*\*.+?\*\*|__.+?__|``.+?``|\^\^.+?\^\^|~~.+?~~)", _re.S)

_MARKER_STYLES = {
    "**": {"bold": True},
    "__": {"italic": True},
    "``": {"mono": True},
    "^^": {"bold": True, "color": PURPLE},
    "~~": {"bold": True, "color": MAROON},
}


def _emit(paragraph, text, style, size):
    """Walk `text`, applying nested inline markup on top of `style`."""
    for part in _MARKUP.split(text):
        if not part:
            continue
        marker = part[:2]
        if marker in _MARKER_STYLES and part.endswith(marker) and len(part) > 4:
            inner = dict(style)
            inner.update(_MARKER_STYLES[marker])
            _emit(paragraph, part[2:-2], inner, size)
        else:
            run_size = Pt(size.pt - 0.5) if style.get("mono") else size
            _style_run(
                paragraph.add_run(part),
                bold=bool(style.get("bold")),
                italic=bool(style.get("italic")),
                size=run_size,
                color=style.get("color"),
                font=MONO_FONT if style.get("mono") else style.get("font"),
            )


def rich(paragraph, text, size=None, color=None, font=None, base_bold=False):
    """
    Add text to a paragraph honouring tiny inline markup:
        **bold**      -> bold
        __italic__    -> italic
        ``code``      -> monospace
        ^^term^^      -> bold + purple (key term)
        ~~hot~~       -> bold + maroon (must-remember value)
    Markers may be nested; inner markers override outer ones.
    """
    size = size or BODY_SIZE
    _emit(paragraph, text,
          {"bold": base_bold, "italic": False, "color": color, "font": font,
           "mono": False},
          size)
    return paragraph


def plain_len(text):
    """Length of text once inline markup is stripped - used for column sizing."""
    import re
    return len(re.sub(r"\*\*|__|``|\^\^|~~", "", text or ""))



# ---------------------------------------------------------------------------
# HEADINGS
# ---------------------------------------------------------------------------
def _bookmark(paragraph, name, bid):
    start = _el("w:bookmarkStart", **{"w:id": bid, "w:name": name})
    end = _el("w:bookmarkEnd", **{"w:id": bid})
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


_OUTLINE_COUNTER = {"n": 1000}


def _set_outline_level(paragraph, level):
    """Mark paragraph with an outline level so the TOC field can pick it up."""
    pPr = paragraph._p.get_or_add_pPr()
    for old in pPr.findall(qn("w:outlineLvl")):
        pPr.remove(old)
    pPr.append(_el("w:outlineLvl", **{"w:val": level}))


def part_title(doc, kicker, title, subtitle=None):
    """Full-width banner announcing a PART of the book."""
    band = doc.add_paragraph()
    band.alignment = WD_ALIGN_PARAGRAPH.CENTER
    band.paragraph_format.space_before = Pt(10)
    band.paragraph_format.space_after = Pt(0)
    para_shade(band, "0B315E")
    para_borders(band, sz=0, color="0B315E", val="none")
    _style_run(band.add_run(kicker.upper()), bold=True, size=Pt(11),
               color=WHITE, font=HEAD_FONT, small_caps=False)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(0)
    t.paragraph_format.space_after = Pt(0)
    para_shade(t, "0B315E")
    _style_run(t.add_run(title.upper()), bold=True, size=Pt(21), color=WHITE,
               font=HEAD_FONT)

    if subtitle:
        s = doc.add_paragraph()
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        s.paragraph_format.space_before = Pt(0)
        s.paragraph_format.space_after = Pt(4)
        para_shade(s, "0B315E")
        _style_run(s.add_run(subtitle), italic=True, size=Pt(10), color=WHITE,
                   font=HEAD_FONT)
    rule(doc, color="C9A227", sz=14, space_after=8)


def chapter_title(doc, number, title):
    """
    Chapter heading (TOC level 1). Uses a shaded bar with the chapter number
    followed by the chapter name.
    """
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    para_shade(p, "0F4C81")
    para_borders(p, edges=("left",), sz=24, color="C9A227")
    _style_run(p.add_run(f"  CHAPTER {number}   "), bold=True, size=Pt(11),
               color=RGBColor(0xC9, 0xA2, 0x27), font=HEAD_FONT)
    _style_run(p.add_run(f"{title.upper()}  "), bold=True, size=Pt(14.5),
               color=WHITE, font=HEAD_FONT)
    _OUTLINE_COUNTER["n"] += 1
    _bookmark(p, f"_Toc_ch{number}", _OUTLINE_COUNTER["n"])
    _set_outline_level(p, 0)
    keep_with_next(p)
    return p


def h2(doc, text, color=TEAL):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    para_borders(p, edges=("bottom",), sz=8, color="0E6B6B", space=2)
    _style_run(p.add_run(text), bold=True, size=Pt(12.5), color=color,
               font=HEAD_FONT)
    _OUTLINE_COUNTER["n"] += 1
    _set_outline_level(p, 1)
    keep_with_next(p)
    return p


def h3(doc, text, color=GREEN):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    _style_run(p.add_run("\u25B6  "), bold=True, size=Pt(8), color=color,
               font=HEAD_FONT)
    _style_run(p.add_run(text), bold=True, size=Pt(11.2), color=color,
               font=HEAD_FONT)
    _OUTLINE_COUNTER["n"] += 1
    _set_outline_level(p, 2)
    keep_with_next(p)
    return p


def h4(doc, text, color=PURPLE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    _style_run(p.add_run(text), bold=True, italic=True, size=Pt(10.6),
               color=color, font=HEAD_FONT)
    keep_with_next(p)
    return p



# ---------------------------------------------------------------------------
# BODY TEXT, RULES, LISTS
# ---------------------------------------------------------------------------
def para(doc, text, size=None, color=None, indent=0.0, space_after=3,
         align="justify", italic_all=False, first_line=0.0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.left_indent = Inches(indent)
    if first_line:
        pf.first_line_indent = Inches(first_line)
    pf.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
    }[align]
    rich(p, text, size=size, color=color, base_bold=False)
    if italic_all:
        for r in p.runs:
            r.italic = True
    return p


def rule(doc, color="0F4C81", sz=8, space_before=2, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.0
    _style_run(p.add_run(""), size=Pt(1))
    para_borders(p, edges=("bottom",), sz=sz, color=color, space=0)
    return p


BULLETS = ["\u25CF", "\u25CB", "\u25AA", "\u2013"]


def bullet(doc, text, level=0, size=None, color=None, space_after=2,
           marker=None, align="justify"):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.left_indent = Inches(0.22 + 0.20 * level)
    pf.first_line_indent = Inches(-0.16)
    pf.line_spacing = 1.05
    pf.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
                    "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    mk = marker if marker is not None else BULLETS[min(level, len(BULLETS) - 1)]
    bullet_colors = [BLUE, TEAL, GREEN, PURPLE]
    _style_run(p.add_run(f"{mk}  "), bold=True, size=Pt(7.5),
               color=bullet_colors[min(level, 3)], font=HEAD_FONT)
    rich(p, text, size=size, color=color)
    return p


def bullets(doc, items, level=0, size=None, space_after=2):
    for it in items:
        if isinstance(it, (list, tuple)):
            bullets(doc, it, level=level + 1, size=size, space_after=space_after)
        else:
            bullet(doc, it, level=level, size=size, space_after=space_after)


def numbered(doc, items, level=0, size=None, start=1, space_after=2):
    for i, it in enumerate(items, start=start):
        bullet(doc, it, level=level, size=size, marker=f"{i}.",
               space_after=space_after)


def defn(doc, term, text, size=None):
    """Run-in definition: bold purple term, then the explanation."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    pf.left_indent = Inches(0.22)
    pf.first_line_indent = Inches(-0.22)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _style_run(p.add_run(f"{term} \u2014 "), bold=True, size=size or BODY_SIZE,
               color=PURPLE, font=HEAD_FONT)
    rich(p, text, size=size)
    return p


# ---------------------------------------------------------------------------
# CALLOUT BOXES  (single-cell tables => reliable full-width shaded boxes)
# ---------------------------------------------------------------------------
_BOX_PRESETS = {
    "definition": (SH_DEF, "0F4C81", "DEFINITION", BLUE),
    "mnemonic": (SH_MNEMONIC, "C9A227", "MNEMONIC", RGBColor(0x8A, 0x6D, 0x00)),
    "highyield": (SH_HIGHYIELD, "C0392B", "HIGH-YIELD FOR EXAM", MAROON),
    "note": (SH_NOTE, "5B2C83", "REMEMBER", PURPLE),
    "compare": (SH_COMPARE, "0E6B6B", "AT A GLANCE", TEAL),
    "exam": (SH_MNEMONIC, "A84B00", "EXAM POINT", ORANGE),
}


def box(doc, kind, lines, title=None, size=Pt(9.8)):
    """
    Shaded, bordered callout box spanning the full text width.
    `lines` may be a string or a list of strings (rendered as bullets).
    """
    fill, border, default_title, title_color = _BOX_PRESETS[kind]
    label = default_title if title is None else title

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    _table_borders(tbl, sz=8, color=border, insideH=0, insideV=0)
    _table_width(tbl, [1.0], usable_width(doc))
    _table_spacing_before(tbl, Pt(5))

    cell = tbl.cell(0, 0)
    cell_shade(cell, fill)
    set_cell_margins(cell, top=70, bottom=70, left=110, right=110)

    first = cell.paragraphs[0]
    first.paragraph_format.space_before = Pt(0)
    first.paragraph_format.space_after = Pt(2)
    first.paragraph_format.line_spacing = 1.0
    if label:
        _style_run(first.add_run(label), bold=True, size=Pt(8.6),
                   color=title_color, font=HEAD_FONT, small_caps=True)
        holder = None
    else:
        holder = first

    if isinstance(lines, str):
        lines = [lines]

    for i, ln in enumerate(lines):
        if holder is not None and i == 0:
            p = holder
        else:
            p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if len(lines) > 1:
            p.paragraph_format.left_indent = Inches(0.16)
            p.paragraph_format.first_line_indent = Inches(-0.16)
            _style_run(p.add_run("\u25AA  "), bold=True, size=Pt(7.5),
                       color=title_color, font=HEAD_FONT)
        rich(p, ln, size=size)
    _after_table_gap(doc, Pt(4))
    return tbl


# ---------------------------------------------------------------------------
# TABLES WITH PROPORTIONAL COLUMN WIDTHS
# ---------------------------------------------------------------------------
def _table_borders(tbl, sz=6, color="9DB8D2", insideH=4, insideV=4,
                   inside_color=None):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    ic = inside_color or color
    for edge in ("top", "left", "bottom", "right"):
        b.append(_el(f"w:{edge}", **{"w:val": "single", "w:sz": sz,
                                     "w:space": 0, "w:color": color}))
    b.append(_el("w:insideH", **{"w:val": "single" if insideH else "none",
                                 "w:sz": insideH or 0, "w:space": 0,
                                 "w:color": ic}))
    b.append(_el("w:insideV", **{"w:val": "single" if insideV else "none",
                                 "w:sz": insideV or 0, "w:space": 0,
                                 "w:color": ic}))
    tblPr.append(b)


def _table_layout_fixed(tbl):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(old)
    tblPr.append(_el("w:tblLayout", **{"w:type": "fixed"}))


def _table_width(tbl, fractions, total):
    """Assign explicit widths; `fractions` must sum to 1.0."""
    _table_layout_fixed(tbl)
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn("w:tblW")):
        tblPr.remove(old)
    tblPr.append(_el("w:tblW", **{"w:w": int(total.twips), "w:type": "dxa"}))

    widths = [Emu(int(total * f)) for f in fractions]
    # rebuild tblGrid
    for old in tbl._tbl.findall(qn("w:tblGrid")):
        tbl._tbl.remove(old)
    grid = OxmlElement("w:tblGrid")
    for w in widths:
        grid.append(_el("w:gridCol", **{"w:w": int(w.twips)}))
    tbl._tbl.insert(1, grid)

    for row in tbl.rows:
        for idx, cell in enumerate(row.cells):
            if idx < len(widths):
                cell.width = widths[idx]
    return widths


def _table_spacing_before(tbl, pts):
    """python-docx has no table spacing; emulate via preceding empty paragraph."""
    prev = tbl._tbl.getprevious()
    if prev is not None and prev.tag == qn("w:p"):
        pPr = prev.get_or_add_pPr()
        spacing = pPr.find(qn("w:spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            pPr.append(spacing)
        spacing.set(qn("w:after"), str(int(pts.pt * 20)))


def _after_table_gap(doc, pts=Pt(4)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    _style_run(p.add_run(""), size=pts)
    return p


def _longest_word(text):
    import re
    clean = re.sub(r"\*\*|__|``|\^\^|~~", "", text or "")
    words = re.split(r"[\s/\-]+", clean)
    return max((len(w) for w in words), default=4)


#: Approximate number of characters of 9 pt Cambria that fit across the full
#: usable text width of an A4 page with 0.5 inch margins, after subtracting
#: typical cell padding. Used to convert character counts into page fractions.
CHARS_PER_FULL_WIDTH = 116.0


def compute_fractions(headers, rows, damp=0.72, min_chars=4, max_share=0.62,
                      chars_per_width=CHARS_PER_FULL_WIDTH):
    """
    Proportional column widths driven by the amount of text each column carries.

    Method
    ------
    1. A raw appetite per column from the average and maximum cell length
       (max dominates, because that is what has to wrap).
    2. `damp` (< 1) compresses extremes so one verbose column cannot starve
       its neighbours, while still receiving the largest share.
    3. Every column is then clamped between
         floor   = width of its longest single word (+ padding)  -> text never
                   breaks mid-word or overflows, and
         ceiling = width of its longest cell (+ padding)          -> a column of
                   short values such as "No." or "121 C" can never be inflated
                   by redistribution.
       Residual width is handed back to the columns that are still hungry.
    """
    ncol = len(headers)
    if ncol == 0:
        return []
    cols = [[headers[c]] + [(r[c] if c < len(r) else "") for r in rows]
            for c in range(ncol)]

    weights, floors, ceils = [], [], []
    for c in range(ncol):
        lens = [plain_len(t) for t in cols[c]]
        maxl = max(lens) if lens else 4
        avgl = (sum(lens) / len(lens)) if lens else 4
        lw = max(_longest_word(t) for t in cols[c])
        appetite = max(0.42 * avgl + 0.58 * maxl, min_chars, lw)
        weights.append(appetite ** damp)
        floors.append(min((lw + 2.5) / chars_per_width, 0.90))
        ceils.append(min((maxl + 2.5) / chars_per_width, max_share))

    # make the clamp range coherent
    for c in range(ncol):
        ceils[c] = max(ceils[c], floors[c])

    total_w = sum(weights) or 1.0
    fr = [w / total_w for w in weights]

    fixed = [False] * ncol
    for _ in range(24):
        residual = 0.0
        changed = False
        for i in range(ncol):
            if fixed[i]:
                continue
            if fr[i] > ceils[i]:
                residual += fr[i] - ceils[i]
                fr[i] = ceils[i]
                fixed[i] = True
                changed = True
            elif fr[i] < floors[i]:
                residual -= floors[i] - fr[i]
                fr[i] = floors[i]
                fixed[i] = True
                changed = True
        free = [i for i in range(ncol) if not fixed[i]]
        if abs(residual) > 1e-9 and free:
            pool = sum(fr[i] for i in free) or 1.0
            for i in free:
                fr[i] += residual * (fr[i] / pool)
            changed = True
        elif abs(residual) > 1e-9 and not free:
            # every column pinned: distribute proportionally to appetite
            pool = sum(weights)
            for i in range(ncol):
                fr[i] += residual * (weights[i] / pool)
            break
        if not changed:
            break

    fr = [max(f, 0.03) for f in fr]
    s = sum(fr) or 1.0
    return [f / s for f in fr]


def table(doc, headers, rows, header_fill=SH_HEADER_BLUE, band=SH_BAND,
          fractions=None, size=TABLE_SIZE, head_size=TABLE_HEAD_SIZE,
          align_center_cols=None, caption=None, first_col_bold=True,
          damp=0.72, max_share=0.46, min_chars=7):
    """
    Build a professional table.

    headers : list[str]
    rows    : list[list[str]]  (inline markup supported in every cell)
    """
    if caption:
        cap = doc.add_paragraph()
        cap.paragraph_format.space_before = Pt(6)
        cap.paragraph_format.space_after = Pt(1)
        cap.paragraph_format.line_spacing = 1.0
        _style_run(cap.add_run(caption), bold=True, size=Pt(9.2), color=NAVY,
                   font=HEAD_FONT, small_caps=True)
        keep_with_next(cap)

    ncol = len(headers)
    tbl = doc.add_table(rows=1, cols=ncol)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    _table_borders(tbl, sz=6, color="8FAAC6", insideH=4, insideV=4,
                   inside_color="C3D3E5")

    if fractions is None:
        fractions = compute_fractions(headers, rows, damp=damp,
                                      min_chars=min_chars, max_share=max_share)
    _table_width(tbl, fractions, usable_width(doc))
    _table_spacing_before(tbl, Pt(5))

    align_center_cols = set(align_center_cols or [])

    # ---- header row
    hdr = tbl.rows[0]
    row_cant_split(hdr)
    repeat_header_row(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        cell_shade(cell, header_fill)
        set_cell_margins(cell, top=50, bottom=50, left=75, right=75)
        cell_valign(cell, "center")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _style_run(p.add_run(text), bold=True, size=head_size, color=WHITE,
                   font=HEAD_FONT)

    # ---- body rows
    for r_i, row_data in enumerate(rows):
        row = tbl.add_row()
        row_cant_split(row)
        fill = band if (r_i % 2 == 0) else None
        for c_i in range(ncol):
            cell = row.cells[c_i]
            if fill:
                cell_shade(cell, fill)
            set_cell_margins(cell, top=42, bottom=42, left=75, right=75)
            cell_valign(cell, "top")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.03
            if c_i in align_center_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            txt = row_data[c_i] if c_i < len(row_data) else ""
            bold = first_col_bold and c_i == 0
            rich(p, txt, size=size, base_bold=bold,
                 color=NAVY if bold else None)
    _after_table_gap(doc, Pt(4))
    return tbl


# ---------------------------------------------------------------------------
# PAGE BREAK  (only used at the end of a chapter)
# ---------------------------------------------------------------------------
def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


def chapter_end(doc, last=False):
    """Decorative end-of-chapter rule followed by the ONLY page break used."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    _style_run(p.add_run("\u2756  \u2756  \u2756"), bold=True, size=Pt(8),
               color=RGBColor(0xC9, 0xA2, 0x27), font=HEAD_FONT)
    if not last:
        page_break(doc)



# ---------------------------------------------------------------------------
# FOOTER / HEADER WITH PAGE NUMBERS
# ---------------------------------------------------------------------------
def _add_field(paragraph, instr, initial="1"):
    r1 = paragraph.add_run()
    fld = _el("w:fldChar", **{"w:fldCharType": "begin"})
    r1._r.append(fld)

    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)

    r3 = paragraph.add_run()
    r3._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))

    r4 = paragraph.add_run(initial)

    r5 = paragraph.add_run()
    r5._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))
    return [r1, r2, r3, r4, r5]


def build_footer(doc, book_title):
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    para_borders(p, edges=("top",), sz=6, color="0F4C81", space=2)

    # left: book title | right: Page X of Y
    pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(usable_width(doc), WD_TAB_ALIGNMENT.RIGHT)
    _style_run(p.add_run(book_title), size=Pt(7.8), color=GREY, font=HEAD_FONT,
               italic=True)
    _style_run(p.add_run("\t"), size=Pt(7.8))
    _style_run(p.add_run("Page "), size=Pt(8.2), color=NAVY, font=HEAD_FONT,
               bold=True)
    for r in _add_field(p, " PAGE "):
        _style_run(r, bold=True, size=Pt(8.2), color=NAVY, font=HEAD_FONT)
    _style_run(p.add_run(" of "), size=Pt(8.2), color=GREY, font=HEAD_FONT)
    for r in _add_field(p, " NUMPAGES "):
        _style_run(r, bold=False, size=Pt(8.2), color=GREY, font=HEAD_FONT)


def suppress_first_page_footer(doc):
    """Keep the cover page clean."""
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    fp = section.first_page_footer
    fp.is_linked_to_previous = False
    if fp.paragraphs:
        fp.paragraphs[0].text = ""


# ---------------------------------------------------------------------------
# TABLE OF CONTENTS (real Word field -> real page numbers)
# ---------------------------------------------------------------------------
def toc_field(doc, levels="1-3"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)

    r1 = p.add_run()
    r1._r.append(_el("w:fldChar", **{"w:fldCharType": "begin",
                                     "w:dirty": "true"}))
    r2 = p.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f' TOC \\o "{levels}" \\h \\z \\u '
    r2._r.append(it)

    r3 = p.add_run()
    r3._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    r4 = p.add_run("Right-click here and choose \u201cUpdate Field\u201d "
                   "to refresh the contents.")
    _style_run(r4, italic=True, size=Pt(9), color=GREY)
    r5 = p.add_run()
    r5._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))
    return p


def cover_page(doc, title, subtitle, exam, topics, meta_lines):
    """Clean, exam-guide style cover."""
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(0)
    _style_run(sp.add_run(""), size=Pt(14))

    top = doc.add_paragraph()
    top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    top.paragraph_format.space_after = Pt(2)
    _style_run(top.add_run(exam.upper()), bold=True, size=Pt(12), color=MAROON,
               font=HEAD_FONT)

    rule(doc, color="C9A227", sz=18, space_after=10)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_after = Pt(2)
    _style_run(t.add_run(title), bold=True, size=Pt(34), color=NAVY,
               font=HEAD_FONT)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.space_after = Pt(10)
    _style_run(s.add_run(subtitle), italic=True, size=Pt(13), color=TEAL,
               font=HEAD_FONT)

    rule(doc, color="0F4C81", sz=10, space_after=14)

    for i, (num, name, desc) in enumerate(topics):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.autofit = False
        _table_borders(tbl, sz=6, color="0F4C81", insideH=0, insideV=0)
        _table_width(tbl, [1.0], usable_width(doc))
        cell = tbl.cell(0, 0)
        cell_shade(cell, SH_DEF if i % 2 == 0 else SH_BAND_GREEN)
        set_cell_margins(cell, top=90, bottom=90, left=130, right=130)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.0
        _style_run(p.add_run(f"UNIT {num}   "), bold=True, size=Pt(9.5),
                   color=MAROON, font=HEAD_FONT)
        _style_run(p.add_run(name), bold=True, size=Pt(15),
                   color=NAVY if i % 2 == 0 else GREEN, font=HEAD_FONT)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.05
        rich(p2, desc, size=Pt(9.5), color=GREY)
        _after_table_gap(doc, Pt(7))

    rule(doc, color="C9A227", sz=10, space_before=8, space_after=8)

    for line in meta_lines:
        m = doc.add_paragraph()
        m.alignment = WD_ALIGN_PARAGRAPH.CENTER
        m.paragraph_format.space_after = Pt(2)
        rich(m, line, size=Pt(9.8), color=GREY)
