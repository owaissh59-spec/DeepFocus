#!/usr/bin/env python3
"""
Render the Perioperative Care supplementary notes .docx into reading-order
7 x 8.5 in booklet pages, preserving the document's colour-coded structure.

The source .docx carries no Heading styles -- every paragraph is "Normal" and
structure is encoded in direct formatting. Roles are recovered from the
(bold, size, colour) signature:

    bold 17.0  navy  #1F3B73  -> document title
    bold 10.5  green #1B6B3A  -> subtitle
         9.0   grey  #555555  -> scope note
    bold 14.5  navy  #1F3B73  -> section heading   (18x)
    bold 11.5  green #1B6B3A  -> subsection        (60x)
         9.5   navy  #1F3B73  -> MCQ question stem (77x)
         9.0   plain          -> MCQ options       (77x)
         9.0   red   #9C2710  -> MCQ answer        (77x)
         9.5   plain          -> body text
    '\u25aa' prefix           -> bullet            (93x)

Single-cell tables (12 of them) are callout boxes; the rest are real tables.
Each MCQ is wrapped in KeepTogether so a question never separates from its
answer across the fold.

Usage:
    python render_docx.py periop.docx pages.pdf [--scale 1.0]
"""

import argparse
import html
import os
import re

import docx
from docx.oxml.ns import qn
from docx.table import Table as DTable
from docx.text.paragraph import Paragraph as DPara
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    HRFlowable,
    KeepTogether,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

IN = 72.0
PAGE_W, PAGE_H = 7 * IN, 8.5 * IN
OUTER, GUTTER = 0.34 * IN, 0.52 * IN
TOP, BOTTOM = 0.34 * IN, 0.40 * IN

NAVY = colors.HexColor("#1F3B73")
GREEN = colors.HexColor("#1B6B3A")
RED = colors.HexColor("#9C2710")
GREY = colors.HexColor("#555555")
RULE = colors.HexColor("#B8C4D9")
BOXBG = colors.HexColor("#F0F3F9")
THBG = colors.HexColor("#E3E9F4")

FONT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fonts")
BASE, BOLD, ITAL, BITAL = "DJV", "DJV-Bold", "DJV-Obl", "DJV-BoldObl"


def register_fonts():
    """DejaVu Sans Condensed -- covers all 38 non-ASCII glyphs the notes use."""
    for name, fn in (
        (BASE, "DejaVuSansCondensed.ttf"),
        (BOLD, "DejaVuSansCondensed-Bold.ttf"),
        (ITAL, "DejaVuSansCondensed-Oblique.ttf"),
        (BITAL, "DejaVuSansCondensed-BoldOblique.ttf"),
    ):
        pdfmetrics.registerFont(TTFont(name, os.path.join(FONT_DIR, fn)))
    pdfmetrics.registerFontFamily(
        BASE, normal=BASE, bold=BOLD, italic=ITAL, boldItalic=BITAL
    )


# ------------------------------------------------------------ classification

def _runs(p):
    return [r for r in p.runs if r.text.strip()]


def signature(p):
    rs = _runs(p)
    if not rs:
        return None
    bold = sum(1 for r in rs if r.bold) / len(rs) >= 0.5
    sizes = [r.font.size.pt for r in rs if r.font.size is not None]
    size = max(sizes) if sizes else 9.5
    col = None
    for r in rs:
        c = r.font.color
        if c is not None and c.type is not None and c.rgb is not None:
            col = str(c.rgb).upper()
            break
    return bold, size, col


def classify(p):
    sig = signature(p)
    if sig is None:
        return None
    bold, size, col = sig
    text = p.text.strip()

    if size >= 16:
        return "title"
    if col == "9C2710":
        return "answer"
    if col == "555555":
        return "scope"
    if col == "1B6B3A":
        return "h2" if size >= 11 else "subtitle"
    if col == "1F3B73":
        if size >= 13:
            return "h1"
        return "closing" if size < 9.4 else "question"
    if text.startswith("\u25aa"):
        return "bullet"
    if re.match(r"^\d+\.\s", text):
        return "numbered"
    if size <= 9.2:
        return "option"
    return "body_bold" if bold else "body"


# ------------------------------------------------------------ inline markup

def inline(p, force=None):
    """Rebuild a paragraph's runs as reportlab markup, keeping bold/italic."""
    out = []
    for r in p.runs:
        t = html.escape(r.text, quote=False)
        if not t:
            continue
        if force == "plain":
            out.append(t)
            continue
        if r.bold and r.italic:
            t = f"<b><i>{t}</i></b>"
        elif r.bold:
            t = f"<b>{t}</b>"
        elif r.italic:
            t = f"<i>{t}</i>"
        if r.underline:
            t = f"<u>{t}</u>"
        out.append(t)
    return "".join(out).strip()


def plain(text):
    return html.escape(text, quote=False)


# ------------------------------------------------------------ styles

def build_styles(s):
    body = ParagraphStyle(
        "body", fontName=BASE, fontSize=9.2 * s, leading=11.6 * s,
        alignment=TA_LEFT, spaceAfter=3.4 * s, allowWidows=0, allowOrphans=0,
    )
    S = {"body": body}
    S["body_bold"] = ParagraphStyle("bb", parent=body, fontName=BOLD)
    S["title"] = ParagraphStyle(
        "title", parent=body, fontName=BOLD, fontSize=15 * s,
        leading=17.5 * s, textColor=NAVY, spaceAfter=1.5 * s,
    )
    S["subtitle"] = ParagraphStyle(
        "sub", parent=body, fontName=BOLD, fontSize=9.8 * s,
        leading=12 * s, textColor=GREEN, spaceAfter=2 * s,
    )
    S["scope"] = ParagraphStyle(
        "scope", parent=body, fontSize=8.2 * s, leading=10.2 * s,
        textColor=GREY, spaceAfter=4 * s,
    )
    S["h1"] = ParagraphStyle(
        "h1", parent=body, fontName=BOLD, fontSize=11.8 * s,
        leading=13.6 * s, textColor=NAVY,
        spaceBefore=9 * s, spaceAfter=3.2 * s, keepWithNext=1,
    )
    S["h2"] = ParagraphStyle(
        "h2", parent=body, fontName=BOLD, fontSize=10 * s,
        leading=12 * s, textColor=GREEN,
        spaceBefore=6 * s, spaceAfter=2.4 * s, keepWithNext=1,
    )
    S["bullet"] = ParagraphStyle(
        "bul", parent=body, leftIndent=9.5 * s, bulletIndent=1.5 * s,
        spaceAfter=2.2 * s,
    )
    S["numbered"] = ParagraphStyle(
        "num", parent=body, leftIndent=13 * s, firstLineIndent=-13 * s,
        spaceAfter=2.2 * s,
    )
    S["question"] = ParagraphStyle(
        "q", parent=body, fontName=BOLD, fontSize=9 * s, leading=11.2 * s,
        textColor=NAVY, spaceBefore=4.5 * s, spaceAfter=1.4 * s, keepWithNext=1,
    )
    S["option"] = ParagraphStyle(
        "opt", parent=body, fontSize=8.6 * s, leading=10.6 * s,
        leftIndent=7 * s, spaceAfter=1.4 * s, keepWithNext=1,
    )
    S["answer"] = ParagraphStyle(
        "ans", parent=body, fontSize=8.6 * s, leading=10.6 * s,
        leftIndent=7 * s, textColor=RED, spaceAfter=4.5 * s,
    )
    S["closing"] = ParagraphStyle(
        "close", parent=body, fontSize=8.6 * s, leading=10.6 * s,
        textColor=NAVY, spaceBefore=6 * s,
    )
    S["cell"] = ParagraphStyle(
        "cell", parent=body, fontSize=8.1 * s, leading=10 * s, spaceAfter=0
    )
    S["cellh"] = ParagraphStyle("ch", parent=S["cell"], fontName=BOLD)
    S["callout"] = ParagraphStyle(
        "co", parent=body, fontSize=8.4 * s, leading=10.5 * s, spaceAfter=0
    )
    S["callout_h"] = ParagraphStyle("coh", parent=S["callout"], fontName=BOLD)
    return S


# ------------------------------------------------------------ tables

def cell_markup(cell, style_key, S):
    parts = []
    for p in cell.paragraphs:
        t = inline(p)
        if t:
            if p.text.strip().startswith("\u25aa"):
                t = "\u2022 " + t.lstrip("\u25aa ").strip()
            parts.append(t)
    return Paragraph("<br/>".join(parts) or "&nbsp;", S[style_key])


def col_weights(rows, ncol):
    """Width proportional to the longest cell in each column, clamped."""
    w = []
    for c in range(ncol):
        longest = max((len(r[c]) for r in rows), default=1)
        w.append(max(longest, 6) ** 0.72)
    total = sum(w)
    frac = [x / total for x in w]
    lo = 0.13
    frac = [max(f, lo) for f in frac]
    t = sum(frac)
    return [f / t for f in frac]


def make_table(t, S, avail_w):
    raw = [[c.text.strip() for c in r.cells] for r in t.rows]
    if not raw:
        return None
    ncol = max(len(r) for r in raw)

    # single cell -> callout box
    if ncol == 1 and len(raw) == 1:
        lines = [p for p in t.rows[0].cells[0].paragraphs]
        flow = []
        for i, p in enumerate(lines):
            txt = inline(p)
            if not txt:
                continue
            if p.text.strip().startswith("\u25aa"):
                txt = "\u2022 " + txt.lstrip("\u25aa ").strip()
            flow.append(Paragraph(txt, S["callout_h"] if i == 0 else S["callout"]))
        if not flow:
            return None
        box = Table([[flow]], colWidths=[avail_w], hAlign="LEFT")
        box.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), BOXBG),
                ("BOX", (0, 0), (-1, -1), 0.5, RULE),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ])
        )
        return box

    raw = [r + [""] * (ncol - len(r)) for r in raw]
    fracs = col_weights(raw, ncol)
    widths = [f * avail_w for f in fracs]

    data = []
    for i, row in enumerate(t.rows):
        cells = list(row.cells) + [None] * (ncol - len(row.cells))
        data.append([
            cell_markup(c, "cellh" if i == 0 else "cell", S)
            if c is not None else Paragraph("", S["cell"])
            for c in cells
        ])

    tbl = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    tbl.setStyle(
        TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.35, RULE),
            ("BACKGROUND", (0, 0), (-1, 0), THBG),
            ("TEXTCOLOR", (0, 0), (-1, 0), NAVY),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3.5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3.5),
            ("TOPPADDING", (0, 0), (-1, -1), 2.4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.4),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#F7F9FC")]),
        ])
    )
    return tbl


# ------------------------------------------------------------ document

class BookDoc(BaseDocTemplate):
    def handle_pageBegin(self):
        self._handle_pageBegin()
        self._handle_nextPageTemplate("verso" if self.page % 2 else "recto")


def make_footer(scale):
    def footer(canv, doc):
        canv.saveState()
        canv.setFont(BASE, 7.2 * scale)
        canv.setFillColor(GREY)
        recto = doc.page % 2 == 1
        y = BOTTOM * 0.42
        if recto:
            canv.drawRightString(PAGE_W - OUTER, y, str(doc.page))
        else:
            canv.drawString(OUTER, y, str(doc.page))
        canv.restoreState()
    return footer


def iter_blocks(doc):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield DPara(child, doc)
        elif child.tag == qn("w:tbl"):
            yield DTable(child, doc)


def build_flow(path, S, avail_w):
    d = docx.Document(path)
    flow = []
    counts = {}

    for block in iter_blocks(d):
        if isinstance(block, DTable):
            t = make_table(block, S, avail_w)
            if t is not None:
                flow += [Spacer(1, 2.5), t, Spacer(1, 4.5)]
                counts["table"] = counts.get("table", 0) + 1
            continue

        role = classify(block)
        if role is None:
            continue
        counts[role] = counts.get(role, 0) + 1
        text = block.text.strip()

        if role == "bullet":
            body = inline(block).lstrip("\u25aa").strip()
            flow.append(Paragraph(body, S["bullet"], bulletText="\u2022"))
        elif role == "numbered":
            flow.append(Paragraph(inline(block), S["numbered"]))
        elif role == "title":
            flow.append(Paragraph(inline(block), S["title"]))
        elif role == "h1":
            flow.append(Paragraph(inline(block), S["h1"]))
        else:
            flow.append(Paragraph(inline(block), S[role]))

        if role == "subtitle":
            flow.append(HRFlowable(width="100%", thickness=0.7,
                                   color=RULE, spaceAfter=5))

    return group_mcqs(flow), counts


def group_mcqs(flow):
    """Keep each question with its options and answer on one page."""
    out = []
    i = 0
    n = len(flow)
    while i < n:
        f = flow[i]
        if isinstance(f, Paragraph) and f.style.name == "q":
            grp = [f]
            j = i + 1
            while j < n and isinstance(flow[j], Paragraph) and \
                    flow[j].style.name in ("opt", "ans"):
                grp.append(flow[j])
                if flow[j].style.name == "ans":
                    j += 1
                    break
                j += 1
            out.append(KeepTogether(grp))
            i = j
        else:
            out.append(f)
            i += 1
    return out


def render(in_path, out_path, scale=1.0, pad_to_4=True):
    register_fonts()
    S = build_styles(scale)
    avail_w = PAGE_W - OUTER - GUTTER
    avail_h = PAGE_H - TOP - BOTTOM

    doc = BookDoc(out_path, pagesize=(PAGE_W, PAGE_H),
                  title="Perioperative Care - Supplementary Notes")
    foot = make_footer(scale)
    doc.addPageTemplates([
        PageTemplate(id="recto", onPage=foot, frames=[
            Frame(GUTTER, BOTTOM, avail_w, avail_h, id="r",
                  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)]),
        PageTemplate(id="verso", onPage=foot, frames=[
            Frame(OUTER, BOTTOM, avail_w, avail_h, id="v",
                  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)]),
    ])

    flow, counts = build_flow(in_path, S, avail_w)
    doc.build(flow)

    from pypdf import PdfReader, PdfWriter
    r = PdfReader(out_path)
    real = len(r.pages)
    total = real
    if pad_to_4 and real % 4:
        w = PdfWriter()
        for p in r.pages:
            w.add_page(p)
        while len(w.pages) % 4:
            w.add_blank_page(width=PAGE_W, height=PAGE_H)
        with open(out_path, "wb") as fh:
            w.write(fh)
        total = len(w.pages)
    return real, total, counts


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("output")
    ap.add_argument("--scale", type=float, default=1.0)
    ap.add_argument("--no-pad", action="store_true")
    a = ap.parse_args()

    real, total, counts = render(a.input, a.output, a.scale, not a.no_pad)
    print("recovered blocks:")
    for k in sorted(counts, key=lambda k: -counts[k]):
        print(f"  {k:<12} {counts[k]}")
    pad = f"  (+{total - real} blank)" if total != real else ""
    print(f"\nwrote {a.output}: {total} pages at 7 x 8.5 in{pad}")


if __name__ == "__main__":
    main()
