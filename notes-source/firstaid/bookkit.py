"""
bookkit.py  --  A small typesetting engine on top of python-docx used to build
the "First Aid - Complete Study Notes" book (A4, topper-style visual design).

Design goals
------------
* A4 page, tight but readable margins, page numbers in the footer.
* Real Word Heading 1/2/3/4 styles so that a field based Table of Contents
  picks them up with correct page numbers.
* Colour coded headings, shaded/bordered callout boxes, professional tables
  with proportional column widths and repeating header rows.
* Page breaks ONLY before a new chapter (i.e. at the end of a chapter).
"""

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

# --------------------------------------------------------------------------
# Palette
# --------------------------------------------------------------------------
NAVY = RGBColor(0x0B, 0x31, 0x61)      # chapter titles
BLUE = RGBColor(0x11, 0x55, 0x9E)      # H2
GREEN = RGBColor(0x0E, 0x6B, 0x48)     # H3
PURPLE = RGBColor(0x6A, 0x1B, 0x9A)    # H4
MAROON = RGBColor(0x9E, 0x1B, 0x32)    # warnings
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x5A, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HX_NAVY = "0B3161"
HX_BLUE = "11559E"
HX_GREEN = "0E6B48"
HX_PURPLE = "6A1B9A"
HX_MAROON = "9E1B32"
HX_AMBER = "B26A00"
HX_TEAL = "0F6E6E"

BODY_FONT = "Cambria"
SANS_FONT = "Calibri"

USABLE_CM = 21.0 - 1.6 - 1.6   # page width minus L/R margins  -> 17.8 cm

BOX_KINDS = {
    # kind        border/accent  fill      title colour   default label
    "key":      (HX_BLUE,   "E8F0FB", HX_NAVY,   "KEY POINTS"),
    "exam":     (HX_MAROON, "FDECEF", HX_MAROON, "EXAM FOCUS"),
    "mnemonic": (HX_PURPLE, "F5ECFA", HX_PURPLE, "MNEMONIC"),
    "def":      (HX_GREEN,  "E9F5EF", HX_GREEN,  "DEFINITION"),
    "warn":     (HX_AMBER,  "FFF6E5", HX_AMBER,  "CAUTION / DO NOT"),
    "note":     (HX_TEAL,   "E9F4F4", HX_TEAL,   "NOTE"),
    "steps":    (HX_NAVY,   "EEF1F6", HX_NAVY,   "PROCEDURE"),
}


# --------------------------------------------------------------------------
# low level XML helpers
# --------------------------------------------------------------------------
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


def shade(element, fill):
    """Apply solid shading to a paragraph-properties or cell-properties host."""
    pr = element
    shd = _el("w:shd", val="clear", color="auto", fill=fill)
    pr.append(shd)


def para_shading(paragraph, fill):
    shade(paragraph._p.get_or_add_pPr(), fill)


def para_borders(paragraph, **kw):
    """kw: top/bottom/left/right = (size_eighth_pt, colour_hex) or None."""
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = _el("w:pBdr")
        pPr.append(pbdr)
    for side in ("top", "left", "bottom", "right"):
        if side in kw and kw[side]:
            size, colour = kw[side]
            pbdr.append(_el("w:" + side, val="single", sz=size, space="4", color=colour))


def cell_shading(cell, fill):
    shade(cell._tc.get_or_add_tcPr(), fill)


def cell_margins(cell, top=40, start=80, bottom=40, end=80):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = _el("w:tcMar")
    for tag, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        mar.append(_el("w:" + tag, w=val, type="dxa"))
    tcPr.append(mar)


def cell_valign(cell, val="center"):
    cell._tc.get_or_add_tcPr().append(_el("w:vAlign", val=val))


def keep_with_next(paragraph, on=True):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(_el("w:keepNext", val="1" if on else "0"))


def no_widow_control(paragraph):
    paragraph._p.get_or_add_pPr().append(_el("w:widowControl", val="0"))


def add_field(paragraph, instr, placeholder=""):
    """Insert a Word field code (used for PAGE / TOC)."""
    r1 = paragraph.add_run()
    r1._r.append(_el("w:fldChar", fldCharType="begin", dirty="true"))
    r2 = paragraph.add_run()
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = instr
    r2._r.append(t)
    r3 = paragraph.add_run()
    r3._r.append(_el("w:fldChar", fldCharType="separate"))
    r4 = paragraph.add_run(placeholder)
    r5 = paragraph.add_run()
    r5._r.append(_el("w:fldChar", fldCharType="end"))
    return [r1, r2, r3, r4, r5]


# --------------------------------------------------------------------------
# inline markup:  **bold**   *italic*   ==highlight==   ~term~ (colour accent)
# --------------------------------------------------------------------------
import re

_TOKEN = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|==.+?==|~.+?~|\*.+?\*)", re.S)


def write_runs(paragraph, text, size=None, font=None, colour=None, bold=False,
               italic=False, highlight=False):
    """Add runs to *paragraph* honouring the light-weight inline markup.

    Markup may be nested, e.g. ``**bold with *italic* inside**``.
    """
    if text is None:
        return
    text = str(text)
    parts = _TOKEN.split(text)
    if len(parts) == 1:                      # plain text -- emit a single run
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        if font:
            run.font.name = font
        if colour:
            run.font.color.rgb = colour
        if highlight:
            run._r.get_or_add_rPr().append(
                _el("w:shd", val="clear", color="auto", fill="FFF2A8"))
        return
    for part in parts:
        if not part:
            continue
        if part.startswith("***") and part.endswith("***") and len(part) > 6:
            write_runs(paragraph, part[3:-3], size, font, colour, True, True,
                       highlight)
        elif part.startswith("**") and part.endswith("**") and len(part) > 4:
            write_runs(paragraph, part[2:-2], size, font, colour, True, italic,
                       highlight)
        elif part.startswith("==") and part.endswith("==") and len(part) > 4:
            write_runs(paragraph, part[2:-2], size, font, colour, True, italic,
                       True)
        elif part.startswith("~") and part.endswith("~") and len(part) > 2:
            write_runs(paragraph, part[1:-1], size, font, MAROON, True, italic,
                       highlight)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            write_runs(paragraph, part[1:-1], size, font, colour, bold, True,
                       highlight)
        else:
            write_runs(paragraph, part, size, font, colour, bold, italic,
                       highlight)


# --------------------------------------------------------------------------
# The Book
# --------------------------------------------------------------------------
class Book:
    def __init__(self, base_size=10.0):
        self.doc = Document()
        self.base = base_size
        self.chapter_no = 0
        self.section_no = 0
        self.toc_entries = []          # [{kind,label,title,key}]
        self._first_chapter = True
        self._setup_page()
        self._setup_styles()
        self._setup_document_settings()

    # ---------------- page & style set-up ---------------------------------
    def _setup_page(self):
        s = self.doc.sections[0]
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)
        s.left_margin = Cm(1.6)
        s.right_margin = Cm(1.6)
        s.top_margin = Cm(1.5)
        s.bottom_margin = Cm(1.4)
        s.header_distance = Cm(0.8)
        s.footer_distance = Cm(0.7)
        s.different_first_page_header_footer = True

    def _setup_document_settings(self):
        settings = self.doc.settings.element
        settings.append(_el("w:updateFields", val="false"))
        settings.append(_el("w:autoHyphenation", val="false"))
        settings.append(_el("w:evenAndOddHeaders", val="false"))

    def _style(self, name, size, colour, bold=True, font=None, space_before=0,
               space_after=0, italic=False, caps=False):
        st = self.doc.styles[name]
        st.font.name = font or SANS_FONT
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.italic = italic
        st.font.color.rgb = colour
        pf = st.paragraph_format
        pf.space_before = Pt(space_before)
        pf.space_after = Pt(space_after)
        pf.line_spacing = 1.0
        pf.keep_with_next = True
        rPr = st.element.get_or_add_rPr()
        rPr.append(_el("w:rFonts", ascii=font or SANS_FONT, hAnsi=font or SANS_FONT,
                       eastAsia=font or SANS_FONT, cs=font or SANS_FONT))
        if caps:
            rPr.append(_el("w:caps", val="1"))
        return st

    def _setup_styles(self):
        n = self.doc.styles["Normal"]
        n.font.name = BODY_FONT
        n.font.size = Pt(self.base)
        n.font.color.rgb = BLACK
        pf = n.paragraph_format
        pf.line_spacing = 1.04
        pf.space_before = Pt(0)
        pf.space_after = Pt(2.5)
        pf.widow_control = True
        rPr = n.element.get_or_add_rPr()
        rPr.append(_el("w:rFonts", ascii=BODY_FONT, hAnsi=BODY_FONT,
                       eastAsia=BODY_FONT, cs=BODY_FONT))

        self._style("Heading 1", 21, NAVY, space_before=2, space_after=4)
        self._style("Heading 2", 13.5, BLUE, space_before=9, space_after=3)
        self._style("Heading 3", 11.5, GREEN, space_before=7, space_after=2)
        self._style("Heading 4", 10.5, PURPLE, space_before=5, space_after=1,
                    italic=False)

    # ---------------- header / footer -------------------------------------
    def build_running_heads(self, title="FIRST AID  |  Complete Study Notes"):
        sec = self.doc.sections[0]

        hdr = sec.header
        p = hdr.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(title)
        r.font.size = Pt(7.5)
        r.font.name = SANS_FONT
        r.font.color.rgb = GREY
        r.font.all_caps = True
        para_borders(p, bottom=(6, "B8C4D9"))
        p.paragraph_format.space_after = Pt(1)

        ftr = sec.footer
        p = ftr.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para_borders(p, top=(6, "B8C4D9"))
        p.paragraph_format.space_before = Pt(1)
        r = p.add_run("Page  ")
        r.font.size = Pt(8)
        r.font.name = SANS_FONT
        r.font.color.rgb = GREY
        runs = add_field(p, " PAGE ", "1")
        for rr in runs:
            rr.font.size = Pt(9)
            rr.font.name = SANS_FONT
            rr.font.bold = True
            rr.font.color.rgb = NAVY

    # ---------------- generic paragraph helpers ---------------------------
    def _new(self, style=None):
        return self.doc.add_paragraph(style=style)

    def p(self, text="", size=None, align="just", space_after=2.5,
          space_before=0, indent=0, italic=False, colour=None, bold=False,
          keep=False):
        par = self._new()
        amap = {"just": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT,
                "center": WD_ALIGN_PARAGRAPH.CENTER,
                "right": WD_ALIGN_PARAGRAPH.RIGHT}
        par.alignment = amap[align]
        pf = par.paragraph_format
        pf.space_after = Pt(space_after)
        pf.space_before = Pt(space_before)
        if indent:
            pf.left_indent = Cm(indent)
        write_runs(par, text, size=size or self.base, italic=italic,
                   colour=colour, bold=bold)
        if keep:
            keep_with_next(par)
        return par

    # bullets -------------------------------------------------------------
    _MARKS = ["\u25aa", "\u2013", "\u25e6", "\u00b7"]

    def bullets(self, items, level=0, size=None, space_after=1.2, tight=True):
        """items: list of strings, or (string, [sub-items]) tuples."""
        for it in items:
            if isinstance(it, (tuple, list)):
                head, subs = it[0], it[1]
                self._bullet(head, level, size, space_after)
                self.bullets(subs, level + 1, size, space_after)
            else:
                self._bullet(it, level, size, space_after)

    def _bullet(self, text, level, size, space_after):
        par = self._new()
        pf = par.paragraph_format
        base_ind = 0.45 + 0.42 * level
        pf.left_indent = Cm(base_ind)
        pf.first_line_indent = Cm(-0.36)
        pf.space_after = Pt(space_after)
        pf.space_before = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.tab_stops.add_tab_stop(Cm(base_ind), WD_TAB_ALIGNMENT.LEFT)
        mark = self._MARKS[min(level, 3)]
        r = par.add_run(mark + "\t")
        r.font.size = Pt((size or self.base) - 0.5)
        r.font.name = SANS_FONT
        r.font.bold = True
        r.font.color.rgb = [NAVY, BLUE, GREEN, PURPLE][min(level, 3)]
        write_runs(par, text, size=size or self.base)
        return par

    def numbered(self, items, size=None, space_after=1.4, start=1, indent=0.0,
                 colour=None):
        for i, it in enumerate(items, start):
            par = self._new()
            pf = par.paragraph_format
            pf.left_indent = Cm(0.62 + indent)
            pf.first_line_indent = Cm(-0.62)
            pf.space_after = Pt(space_after)
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.tab_stops.add_tab_stop(Cm(0.62 + indent), WD_TAB_ALIGNMENT.LEFT)
            r = par.add_run("%d." % i + "\t")
            r.font.bold = True
            r.font.size = Pt(size or self.base)
            r.font.name = SANS_FONT
            r.font.color.rgb = colour or NAVY
            write_runs(par, it, size=size or self.base)

    def dl(self, pairs, size=None, space_after=1.6, sep="  \u2014  "):
        """Definition list: bold term, then description on the same line."""
        for term, desc in pairs:
            par = self._new()
            pf = par.paragraph_format
            pf.left_indent = Cm(0.45)
            pf.first_line_indent = Cm(-0.45)
            pf.space_after = Pt(space_after)
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = par.add_run(term)
            r.font.bold = True
            r.font.size = Pt(size or self.base)
            r.font.color.rgb = NAVY
            r2 = par.add_run(sep)
            r2.font.size = Pt(size or self.base)
            r2.font.color.rgb = GREY
            write_runs(par, desc, size=size or self.base)

    # ---------------- structural elements --------------------------------
    def part(self, label, title):
        """A part divider strip (no page break -- flows inline)."""
        self.toc_entries.append({"kind": "part", "label": label,
                                 "title": title, "key": None})
        par = self._new()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(6)
        par.paragraph_format.space_after = Pt(8)
        para_shading(par, HX_NAVY)
        r = par.add_run(label + "   \u2022   " + title.upper())
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.name = SANS_FONT
        r.font.color.rgb = WHITE
        keep_with_next(par)

    def chapter(self, title, subtitle=None, syllabus=None, tag=None):
        """Start a new chapter on a fresh page."""
        self.chapter_no += 1
        self.section_no = 0
        label = tag or ("CHAPTER %d" % self.chapter_no)
        self.toc_entries.append({"kind": "chapter",
                                 "label": label,
                                 "title": title,
                                 "key": "%s %s" % (label, title)})
        if not self._first_chapter:
            self.page_break()
        self._first_chapter = False

        tagp = self._new()
        tagp.paragraph_format.space_after = Pt(1)
        tagp.paragraph_format.space_before = Pt(0)
        tag = tagp
        r = tag.add_run(label)
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.name = SANS_FONT
        r.font.color.rgb = MAROON
        keep_with_next(tag)

        h = self.doc.add_paragraph(title, style="Heading 1")
        para_borders(h, bottom=(18, HX_NAVY))
        h.paragraph_format.space_after = Pt(4)
        if subtitle:
            sp = self._new()
            sp.paragraph_format.space_before = Pt(3)
            sp.paragraph_format.space_after = Pt(6)
            r = sp.add_run(subtitle)
            r.font.size = Pt(9.5)
            r.font.italic = True
            r.font.name = SANS_FONT
            r.font.color.rgb = GREY
            keep_with_next(sp)
        if syllabus:
            self.box("SYLLABUS COVERED IN THIS CHAPTER", syllabus, kind="note",
                     size=9)
        return h

    def h2(self, text, num=None):
        """Auto-numbered section heading (e.g. '3.4  Roller Bandages')."""
        if num is None:
            self.section_no += 1
            num = "%d.%d" % (self.chapter_no, self.section_no)
        label = "%s  %s" % (num, text)
        self.toc_entries.append({"kind": "section", "label": num,
                                 "title": text, "key": label})
        h = self.doc.add_paragraph(label, style="Heading 2")
        para_borders(h, bottom=(6, "C6D3E6"))
        return h

    def h3(self, text):
        return self.doc.add_paragraph(text, style="Heading 3")

    def h4(self, text):
        return self.doc.add_paragraph(text, style="Heading 4")

    def page_break(self):
        p = self._new()
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_break(WD_BREAK.PAGE)

    def divider(self):
        p = self._new()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        para_borders(p, bottom=(6, "C6D3E6"))

    # ---------------- callout box ----------------------------------------
    def box(self, title, lines, kind="key", size=None, bullet=True):
        accent, fill, tcol, default = BOX_KINDS[kind]
        if title is None:
            title = default
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        self._table_width(tbl, [1.0])
        self._table_borders(tbl, sz=6, colour=accent, insideH=None, insideV=None)
        self._tbl_spacing_before(tbl, 4)
        cell = tbl.rows[0].cells[0]
        cell_shading(cell, fill)
        cell_margins(cell, top=70, bottom=70, start=130, end=110)
        self._cell_left_accent(cell, accent)

        cell.paragraphs[0].text = ""
        tp = cell.paragraphs[0]
        tp.paragraph_format.space_after = Pt(2.5)
        tp.paragraph_format.line_spacing = 1.0
        r = tp.add_run(title)
        r.font.size = Pt((size or self.base) - 0.5)
        r.font.bold = True
        r.font.name = SANS_FONT
        r.font.color.rgb = RGBColor.from_string(tcol)
        r.font.all_caps = True

        if isinstance(lines, str):
            lines = [lines]
        for ln in lines:
            par = cell.add_paragraph()
            pf = par.paragraph_format
            pf.space_after = Pt(1.5)
            pf.line_spacing = 1.02
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if bullet:
                pf.left_indent = Cm(0.4)
                pf.first_line_indent = Cm(-0.4)
                rr = par.add_run("\u25aa  ")
                rr.font.size = Pt((size or self.base) - 1)
                rr.font.bold = True
                rr.font.color.rgb = RGBColor.from_string(accent)
            write_runs(par, ln, size=size or self.base)
        self._after_table_gap(3)
        return tbl

    def _cell_left_accent(self, cell, colour):
        tcPr = cell._tc.get_or_add_tcPr()
        bdr = _el("w:tcBorders")
        bdr.append(_el("w:left", val="single", sz=24, space="0", color=colour))
        tcPr.append(bdr)

    # ---------------- tables ---------------------------------------------
    def _table_width(self, tbl, weights):
        """Fixed layout; column widths proportional to *weights*."""
        total = float(sum(weights))
        usable = Cm(USABLE_CM)
        tblPr = tbl._tbl.tblPr
        tblPr.append(_el("w:tblLayout", type="fixed"))
        tblPr.append(_el("w:tblW", w=str(int(usable.twips)), type="dxa"))
        grid = tbl._tbl.find(qn("w:tblGrid"))
        widths = []
        for w in weights:
            widths.append(int(round(usable.twips * (w / total))))
        # correct rounding drift on the last column
        widths[-1] += int(usable.twips) - sum(widths)
        if grid is not None:
            for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
                gc.set(qn("w:w"), str(w))
        for row in tbl.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = Emu(int(w * 635))
        return widths

    def _table_borders(self, tbl, sz=4, colour="9DB2CE", insideH="D5DEEB",
                       insideV="D5DEEB"):
        tblPr = tbl._tbl.tblPr
        bdr = _el("w:tblBorders")
        for side in ("top", "left", "bottom", "right"):
            bdr.append(_el("w:" + side, val="single", sz=sz, space="0",
                           color=colour))
        if insideH:
            bdr.append(_el("w:insideH", val="single", sz=2, space="0",
                           color=insideH))
        if insideV:
            bdr.append(_el("w:insideV", val="single", sz=2, space="0",
                           color=insideV))
        tblPr.append(bdr)

    def _tbl_spacing_before(self, tbl, pt):
        pass  # spacing handled by surrounding paragraphs

    def _after_table_gap(self, pt=3):
        p = self._new()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        for r in p.runs:
            r.font.size = Pt(pt)
        run = p.add_run("")
        run.font.size = Pt(pt)

    def table(self, headers, rows, weights=None, caption=None, size=None,
              header_fill=HX_NAVY, zebra="F1F5FA", align=None, first_bold=True,
              font_size_body=None):
        """Create a styled table.

        headers  : list of column headings (or None for a head-less table)
        rows     : list of row value lists
        weights  : relative column widths; if omitted they are derived from
                   the average rendered text length of each column so that the
                   text fits comfortably.
        align    : optional list of 'l'/'c'/'r' per column
        """
        ncols = len(headers) if headers else len(rows[0])
        size = size or (self.base - 1.2)
        body_size = font_size_body or size
        if caption:
            cp = self.p("", space_after=1.5, space_before=5)
            r = cp.add_run(caption)
            r.font.size = Pt(size)
            r.font.bold = True
            r.font.name = SANS_FONT
            r.font.color.rgb = NAVY
            keep_with_next(cp)

        if weights is None:
            weights = self._auto_weights(headers, rows, ncols)

        tbl = self.doc.add_table(rows=0, cols=ncols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        self._table_borders(tbl)

        if headers:
            hr = tbl.add_row()
            trPr = hr._tr.get_or_add_trPr()
            trPr.append(_el("w:tblHeader", val="true"))
            trPr.append(_el("w:cantSplit", val="true"))
            for i, htxt in enumerate(headers):
                c = hr.cells[i]
                cell_shading(c, header_fill)
                cell_margins(c, top=50, bottom=50, start=75, end=75)
                cell_valign(c, "center")
                par = c.paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.space_before = Pt(0)
                par.paragraph_format.line_spacing = 1.0
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = par.add_run(str(htxt))
                r.font.size = Pt(size)
                r.font.bold = True
                r.font.name = SANS_FONT
                r.font.color.rgb = WHITE

        amap = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER,
                "r": WD_ALIGN_PARAGRAPH.RIGHT, "j": WD_ALIGN_PARAGRAPH.JUSTIFY}
        for ri, row in enumerate(rows):
            tr = tbl.add_row()
            tr._tr.get_or_add_trPr().append(_el("w:cantSplit", val="true"))
            for ci in range(ncols):
                val = row[ci] if ci < len(row) else ""
                c = tr.cells[ci]
                cell_margins(c, top=42, bottom=42, start=75, end=75)
                cell_valign(c, "top")
                if zebra and ri % 2 == 1:
                    cell_shading(c, zebra)
                par = c.paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.space_before = Pt(0)
                par.paragraph_format.line_spacing = 1.0
                if align:
                    par.alignment = amap.get(align[ci], WD_ALIGN_PARAGRAPH.LEFT)
                else:
                    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
                bold = first_bold and ci == 0
                colour = NAVY if bold else None
                # allow multi-line cells with "\n"
                chunks = str(val).split("\n")
                write_runs(par, chunks[0], size=body_size, bold=bold,
                           colour=colour)
                for extra in chunks[1:]:
                    np = c.add_paragraph()
                    np.paragraph_format.space_after = Pt(0)
                    np.paragraph_format.line_spacing = 1.0
                    if align:
                        np.alignment = amap.get(align[ci],
                                                WD_ALIGN_PARAGRAPH.LEFT)
                    write_runs(np, extra, size=body_size, bold=bold,
                               colour=colour)
        self._table_width(tbl, weights)
        self._after_table_gap(4)
        return tbl

    @staticmethod
    def _plain(s):
        return re.sub(r"\*\*|==|~|\*", "", str(s))

    def _auto_weights(self, headers, rows, ncols):
        """Derive proportional column widths from content length."""
        maxw = [0.0] * ncols
        avgw = [0.0] * ncols
        for ci in range(ncols):
            lens = []
            for row in rows:
                v = self._plain(row[ci]) if ci < len(row) else ""
                longest_word = max([len(w) for w in v.replace("\n", " ").split()] or [1])
                lens.append(len(v))
                maxw[ci] = max(maxw[ci], longest_word)
            avgw[ci] = (sum(lens) / float(len(lens))) if lens else 1.0
            if headers:
                h = self._plain(headers[ci])
                maxw[ci] = max(maxw[ci], max([len(w) for w in h.split()] or [1]))
                avgw[ci] = max(avgw[ci], len(h) * 0.85)
        # weight = sqrt-damped average length, but never narrower than the
        # longest single word in the column
        weights = []
        for ci in range(ncols):
            w = max(avgw[ci] ** 0.62, maxw[ci] * 0.62, 3.0)
            weights.append(w)
        return weights

    # ---------------- front matter ---------------------------------------
    def cover(self, title_lines, subtitle, exam, tagline, bullets_):
        band = self._new()
        band.paragraph_format.space_before = Pt(26)
        band.paragraph_format.space_after = Pt(0)
        band.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para_shading(band, HX_MAROON)
        r = band.add_run(exam)
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.name = SANS_FONT
        r.font.color.rgb = WHITE

        for i, ln in enumerate(title_lines):
            p = self._new()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(24 if i == 0 else 2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(ln)
            r.font.size = Pt(46 if i == 0 else 26)
            r.font.bold = True
            r.font.name = SANS_FONT
            r.font.color.rgb = NAVY if i == 0 else BLUE

        rule = self._new()
        rule.paragraph_format.space_before = Pt(6)
        rule.paragraph_format.space_after = Pt(10)
        para_borders(rule, bottom=(24, HX_MAROON))

        p = self._new()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(subtitle)
        r.font.size = Pt(14)
        r.font.name = SANS_FONT
        r.font.color.rgb = GREEN
        r.font.bold = True

        p = self._new()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(16)
        r = p.add_run(tagline)
        r.font.size = Pt(11)
        r.font.italic = True
        r.font.color.rgb = GREY

        self.box("WHAT THIS BOOK GIVES YOU", bullets_, kind="key", size=10)
        self.page_break()

    # ---------------- static (pre-paginated) table of contents ------------
    def static_toc(self, entries, page_map=None, heading="TABLE OF CONTENTS",
                   note=None):
        """Render a hand-built contents page with dot leaders and real page
        numbers.  *page_map* maps entry['key'] -> printed page number."""
        from docx.enum.text import WD_TAB_LEADER
        page_map = page_map or {}
        h = self._new()
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.paragraph_format.space_after = Pt(3)
        h.paragraph_format.space_before = Pt(2)
        para_shading(h, HX_NAVY)
        r = h.add_run(heading)
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.name = SANS_FONT
        r.font.color.rgb = WHITE
        if note:
            self.p(note, size=8.5, align="center", italic=True, colour=GREY,
                   space_after=5)

        right = Cm(USABLE_CM)
        for e in entries:
            if e["kind"] == "part":
                par = self._new()
                par.paragraph_format.space_before = Pt(7)
                par.paragraph_format.space_after = Pt(2.5)
                para_shading(par, "E3E9F2")
                para_borders(par, left=(18, HX_MAROON))
                par.paragraph_format.left_indent = Cm(0.15)
                rr = par.add_run("  %s  \u2022  %s" % (e["label"],
                                                       e["title"].upper()))
                rr.font.size = Pt(9.5)
                rr.font.bold = True
                rr.font.name = SANS_FONT
                rr.font.color.rgb = NAVY
                keep_with_next(par)
                continue

            par = self._new()
            pf = par.paragraph_format
            pf.space_after = Pt(1.2)
            pf.space_before = Pt(2.5 if e["kind"] == "chapter" else 0)
            pf.tab_stops.add_tab_stop(right, WD_TAB_ALIGNMENT.RIGHT,
                                      WD_TAB_LEADER.DOTS)
            if e["kind"] == "chapter":
                pf.left_indent = Cm(0.2)
                r1 = par.add_run("%s.  " % e["label"].replace("CHAPTER ", "Ch. ")
                                 .replace("APPENDIX", "App."))
                r1.font.size = Pt(9.5)
                r1.font.bold = True
                r1.font.name = SANS_FONT
                r1.font.color.rgb = MAROON
                r2 = par.add_run(e["title"])
                r2.font.size = Pt(10.5)
                r2.font.bold = True
                r2.font.name = SANS_FONT
                r2.font.color.rgb = NAVY
                num_size, num_bold, num_col = 10.5, True, NAVY
            else:
                pf.left_indent = Cm(1.1)
                r1 = par.add_run("%s  " % e["label"])
                r1.font.size = Pt(9)
                r1.font.bold = True
                r1.font.name = SANS_FONT
                r1.font.color.rgb = BLUE
                r2 = par.add_run(e["title"])
                r2.font.size = Pt(9)
                r2.font.name = BODY_FONT
                r2.font.color.rgb = BLACK
                num_size, num_bold, num_col = 9, False, GREY
            pg = page_map.get(e["key"], "")
            r3 = par.add_run("\t%s" % (pg if pg else "\u2013"))
            r3.font.size = Pt(num_size)
            r3.font.bold = num_bold
            r3.font.name = SANS_FONT
            r3.font.color.rgb = num_col
        self.page_break()

    def toc(self, levels="1-2", heading="TABLE OF CONTENTS",
            note=None):
        h = self._new()
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.paragraph_format.space_after = Pt(2)
        h.paragraph_format.space_before = Pt(2)
        para_shading(h, HX_NAVY)
        r = h.add_run(heading)
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.name = SANS_FONT
        r.font.color.rgb = WHITE
        if note:
            self.p(note, size=8.5, align="center", italic=True, colour=GREY,
                   space_after=6)
        p = self._new()
        p.paragraph_format.space_before = Pt(4)
        add_field(p, ' TOC \\o "%s" \\h \\z \\u ' % levels,
                  "Right-click and choose 'Update Field' to build the contents.")
        self.page_break()

    def save(self, path):
        self.doc.save(path)
