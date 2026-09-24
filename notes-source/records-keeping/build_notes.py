# -*- coding: utf-8 -*-
"""
Builds the DOCX study notes.

Two-pass strategy so that the Table of Contents carries REAL page numbers:
  pass 1 : build with blank page numbers  ->  convert to PDF  ->  read which
           page each chapter/section heading landed on  ->  save pagemap.json
  pass 2 : build again, this time printing the measured page numbers.
Because the TOC lines are identical in both passes (only the right-aligned
number changes), pagination does not shift between the passes.
"""
import json
import os
import re
import subprocess
import sys

from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER

import docx_kit as K
from docx_kit import Book, add_runs, para_shading, para_borders, _norm

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'JKSSB_Junior_Pharmacist_Stores_Records_and_Procedures_Notes.docx')
PAGEMAP = os.path.join(HERE, 'pagemap.json')
SOFFICE = '/opt/libreoffice25.8/program/soffice'

TITLE = 'RECORDS KEEPING'
SUBTITLE = 'Stores Records & Procedures'
RUNNING = 'Records Keeping \u2014 Stores Records & Procedures  |  JKSSB Junior Pharmacist'

SYLLABUS = [
    '**Stores Records & Procedures** \u2014 complete coverage of the prescribed topic.',
    '**Clerical procedure in the goods inward section.**',
    '**Records and procedures in main stores.**',
    '**Classification and codification** of stores/drug items.',
    '**Keeping of stock books.**',
    '**Preparation of indents.**',
    '**Methods of storing drugs.**',
]

META = [
    '**Subject:** Hospital & Clinical Pharmacy / Drug Store & Business Management',
    '**Exam:** JKSSB Junior Pharmacist',
    '**Paper coverage:** Records Keeping \u2014 Stores Records & Procedures',
    '14 chapters \u2022 fully tabulated \u2022 high-yield boxes \u2022 quick-revision capsule',
]


# ------------------------------------------------------------------ TOC
def write_toc(b, marker):
    """Write the Table of Contents and move it in front of `marker`."""
    body = b.doc.element.body
    before = list(body)          # keep the lxml proxies alive for identity tests

    p = b._p(align=WD_ALIGN_PARAGRAPH.CENTER, after=1, before=0)
    para_shading(p, K.NAVY)
    para_borders(p, left=(24, K.MAROON, 4))
    add_runs(p, 'TABLE OF CONTENTS', size=17, colour='FFFFFF', bold=True)
    sp = b._p('Chapter-wise contents with page numbers', size=9, colour=K.GREY,
              italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    para_borders(sp, bottom=(6, 'F2C14E', 3))

    hdr = b._p(after=3)
    hdr.paragraph_format.tab_stops.add_tab_stop(
        Cm(K.USABLE_CM), WD_TAB_ALIGNMENT.RIGHT)
    add_runs(hdr, 'TOPIC', size=8.5, colour=K.MAROON, bold=True)
    add_runs(hdr, '\tPAGE', size=8.5, colour=K.MAROON, bold=True)
    para_borders(hdr, bottom=(6, K.GRID, 2))

    for level, number, title, key in b.toc_entries:
        page = b.page_map.get(key, '')
        if level == 'part':
            q = b._p(before=7, after=3)
            para_shading(q, K.LTGREY)
            add_runs(q, '   %s  \u2022  %s' % (number, title.upper()), size=9.5,
                     colour=K.MAROON, bold=True)
            continue
        q = b._p(after=1)
        pf = q.paragraph_format
        pf.line_spacing = 1.04
        if level == 'ch':
            pf.left_indent = Cm(0.25)
            pf.space_before = Pt(4)
            pf.tab_stops.add_tab_stop(Cm(K.USABLE_CM), WD_TAB_ALIGNMENT.RIGHT,
                                      WD_TAB_LEADER.DOTS)
            add_runs(q, 'Chapter %s.  ' % number, size=10.5, colour=K.MAROON, bold=True)
            add_runs(q, title, size=10.5, colour=K.NAVY, bold=True)
            add_runs(q, '\t%s' % page, size=10.5, colour=K.NAVY, bold=True)
        else:
            pf.left_indent = Cm(1.15)
            pf.tab_stops.add_tab_stop(Cm(K.USABLE_CM), WD_TAB_ALIGNMENT.RIGHT,
                                      WD_TAB_LEADER.DOTS)
            add_runs(q, '%s  ' % number, size=9.5, colour=K.BLUE, bold=True)
            add_runs(q, title, size=9.5)
            add_runs(q, '\t%s' % page, size=9.5, colour=K.BLUE)

    b.spacer(6)
    note = b._p(after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_runs(note, 'How to use this book: read the chapter, then re-read only the '
                   'coloured boxes and tables the night before the examination. '
                   'Chapter 13 is a complete revision capsule.',
             size=8.5, colour=K.GREY, italic=True)
    b.page_break()

    moved = [e for e in list(body) if not any(e is x for x in before)]
    for el in moved:
        marker._p.addprevious(el)


# ------------------------------------------------------------------ build
def build(page_map):
    b = Book(RUNNING, page_map=page_map)
    b.cover(TITLE, SUBTITLE, SYLLABUS, META)
    marker = b._p()               # TOC will be injected here
    marker.paragraph_format.space_after = Pt(0)

    import content_a, content_b, content_c, content_d, content_e, content_f, \
        content_g, content_h
    content_a.chapter1(b)
    content_a.chapter2(b)
    content_b.chapter3(b)
    content_b.chapter4(b)
    content_c.chapter5(b)
    content_c.chapter6(b)
    content_d.chapter7(b)
    content_d.chapter8(b)
    content_e.chapter9(b)
    content_f.chapter10(b)
    content_g.chapter11(b)
    content_g.chapter12(b)
    content_h.chapter13(b)
    content_h.chapter14(b)

    write_toc(b, marker)
    # remove the marker paragraph
    marker._p.getparent().remove(marker._p)
    b.doc.save(OUT)
    return b


# ------------------------------------------------------------------ measure
def to_pdf():
    outdir = os.path.join(HERE, 'pdf')
    os.makedirs(outdir, exist_ok=True)
    subprocess.run([SOFFICE, '--headless', '--norestore', '--convert-to', 'pdf',
                    '--outdir', outdir, OUT],
                   check=True, capture_output=True, timeout=900)
    return os.path.join(outdir, os.path.splitext(os.path.basename(OUT))[0] + '.pdf')


def page_texts(pdf, lower=True):
    txt = subprocess.run(['pdftotext', '-layout', pdf, '-'],
                         check=True, capture_output=True).stdout.decode('utf-8', 'ignore')
    pages = [re.sub(r'\s+', ' ', pg).strip() for pg in txt.split('\f')]
    return [p.lower() for p in pages] if lower else pages


def measure(b, pdf):
    lo = page_texts(pdf)
    raw = page_texts(pdf, lower=False)
    # the TOC is the last of the first few pages that still shows dot leaders
    toc_end = 1
    for i, t in enumerate(lo[:10], 1):
        if re.search(r'\.{6,}', t):
            toc_end = i
    pm, missing, order = {}, [], []
    for level, number, title, key in b.toc_entries:
        if not key:
            continue
        cased = key.startswith('CH::')
        parts = (key[4:] if cased else key).split('||')
        hay = raw if cased else lo
        found = None
        for i in range(toc_end, len(hay)):
            if all(p in hay[i] for p in parts):
                found = i + 1
                break
        if found:
            pm[key] = found
            order.append((found, key))
        else:
            missing.append(key)
    # sanity: page numbers must never go backwards in contents order
    for (p1, k1), (p2, k2) in zip(order, order[1:]):
        if p2 < p1:
            print('  !! out of order: %s (p%d) then %s (p%d)' % (k1[:50], p1, k2[:50], p2))
    return pm, missing, len(lo), toc_end


if __name__ == '__main__':
    sys.path.insert(0, HERE)
    pm = {}
    if '--use-cache' in sys.argv and os.path.exists(PAGEMAP):
        pm = json.load(open(PAGEMAP))
    b = build(pm)
    print('pass 1 written')
    pdf = to_pdf()
    pm2, missing, npages, toc_end = measure(b, pdf)
    print('pdf pages: %d   toc ends on page %d   resolved %d/%d headings'
          % (npages, toc_end, len(pm2), len([e for e in b.toc_entries if e[3]])))
    if missing:
        print('UNRESOLVED:')
        for m in missing:
            print('   ', m[:90])
    json.dump(pm2, open(PAGEMAP, 'w'), indent=1)
    b2 = build(pm2)
    print('pass 2 written ->', OUT)
    pdf = to_pdf()
    pages = page_texts(pdf)
    print('final pdf pages:', len(pages))
