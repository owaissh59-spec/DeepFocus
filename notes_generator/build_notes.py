# -*- coding: utf-8 -*-
"""
Generator for M-Pharmacy notes:
"Formulation Development of Pharmaceutical and Cosmetic Products (MPH204T)"

Replicates the visual design of the existing repo notes
(Cosmetics and Cosmeceuticals.pdf) WITHOUT the diagonal watermark.

Design system
-------------
- Page   : US Letter (8.5 x 11 in), ~1 in margins.
- Header : "CAREWELL PHARMACY" centered bold, thick brown rule below.
- Footer : brown rule, "For More Notes ..." left + "Page X" right,
           red bold Telegram line below.
- Body   : Comic Sans MS 14 pt, 1.5 line spacing, justified, black.
- H1 Unit: centered, red (#FF0000), bold, 20 pt.
- H2     : left, red (#FF0000), bold, 16 pt.
- H3     : left, black, bold, 14 pt.
- Bullets: arrow (Wingdings F0D8); sub-bullets small dots.
- Tables : maroon (#5B1A38) header w/ white bold text, alternating gray rows.
- Boxes  : shaded callout boxes for High-Yield / Mnemonic / Important notes.
- Figures: embedded PNG diagrams with numbered captions.
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------- palette
RED       = RGBColor(0xFF, 0x00, 0x00)
TG_RED    = RGBColor(0xCC, 0x00, 0x00)
BLACK     = RGBColor(0x00, 0x00, 0x00)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
BLUE_LINK = RGBColor(0x05, 0x63, 0xC1)
GREEN     = RGBColor(0x1F, 0x7A, 0x33)
PURPLE    = RGBColor(0x6A, 0x1B, 0x9A)
BROWN     = "823B0A"
MAROON    = "5B1A38"
ROW_A     = "EDE7EA"
ROW_B     = "D9D9D9"
BOX_FILL  = "FFF3D6"   # light amber for highlight boxes
BOX_EDGE  = "E0A800"
GREY_TXT  = RGBColor(0x59, 0x59, 0x59)

BODY_FONT = "Comic Sans MS"
BODY_SIZE = 14
FIG_DIR = None  # set at runtime


# ---------------------------------------------------------------- low-level helpers
def _set_font(run, name=BODY_FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rfonts.set(qn(a), name)


def _shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_fill)
    tcpr.append(shd)


def _cell_margins(cell, top=50, bottom=50, left=100, right=100):
    tcpr = cell._tc.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for tag, val in (('top', top), ('bottom', bottom), ('start', left), ('end', right)):
        e = OxmlElement('w:' + tag); e.set(qn('w:w'), str(val)); e.set(qn('w:type'), 'dxa')
        m.append(e)
    tcpr.append(m)


def _para_border(p, edge='bottom', color=BROWN, sz=18, space=1):
    ppr = p._p.get_or_add_pPr()
    pbdr = ppr.find(qn('w:pBdr'))
    if pbdr is None:
        pbdr = OxmlElement('w:pBdr'); ppr.append(pbdr)
    e = OxmlElement('w:' + edge)
    e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
    e.set(qn('w:space'), str(space)); e.set(qn('w:color'), color)
    pbdr.append(e)


def _add_field(paragraph, field):
    run = paragraph.add_run()
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = field
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end')
    run._r.append(b); run._r.append(instr); run._r.append(e)
    _set_font(run); run.font.size = Pt(9)
    return run


def _table_borders(table, color="BFBFBF", sz=4):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        borders.append(e)
    tblPr.append(borders)


def _cell_box(cell, fill, edge_color):
    """Shade a cell and give it a coloured border (for callout boxes)."""
    _shade(cell, fill)
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge, sz in (('top', 6), ('bottom', 6), ('start', 24), ('end', 6)):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), edge_color)
        borders.append(e)
    tcpr.append(borders)


def _keep_with_next(p):
    ppr = p._p.get_or_add_pPr()
    ppr.append(OxmlElement('w:keepNext'))


def _set_font_style(style, name):
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.append(rfonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rfonts.set(qn(a), name)


# ---------------------------------------------------------------- doc builder
class Notes:
    def __init__(self):
        self.doc = Document()
        self._base_style()
        self._page_setup()
        self._header_footer()
        self.fig_no = 0
        self.mcq_no = 0

    def _base_style(self):
        st = self.doc.styles['Normal']
        st.font.name = BODY_FONT
        st.font.size = Pt(BODY_SIZE)
        st.font.color.rgb = BLACK
        _set_font_style(st, BODY_FONT)
        pf = st.paragraph_format
        pf.line_spacing = 1.5
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.space_before = Pt(0)
        pf.space_after = Pt(8)

    def _page_setup(self):
        sec = self.doc.sections[0]
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11)
        sec.top_margin = Inches(0.9)
        sec.bottom_margin = Inches(0.9)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)
        sec.header_distance = Inches(0.5)
        sec.footer_distance = Inches(0.4)

    def _header_footer(self):
        sec = self.doc.sections[0]
        hdr = sec.header
        hdr.is_linked_to_previous = False
        hp = hdr.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = hp.add_run("CAREWELL PHARMACY")
        r.bold = True; r.font.size = Pt(14); r.font.color.rgb = BLACK; _set_font(r)
        _para_border(hp, 'bottom', BROWN, sz=20, space=1)
        hp.paragraph_format.space_after = Pt(2)
        hp.paragraph_format.line_spacing = 1.0

        ftr = sec.footer
        ftr.is_linked_to_previous = False
        fp = ftr.paragraphs[0]
        _para_border(fp, 'top', BROWN, sz=20, space=1)
        fp.paragraph_format.space_before = Pt(2)
        fp.paragraph_format.space_after = Pt(0)
        fp.paragraph_format.line_spacing = 1.0
        fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_ALIGN_PARAGRAPH.RIGHT)
        r1 = fp.add_run("For More Notes - Visit "); r1.font.size = Pt(9); _set_font(r1)
        r2 = fp.add_run("www.carewellpharmacy.in")
        r2.font.size = Pt(9); r2.font.color.rgb = BLUE_LINK; r2.underline = True; _set_font(r2)
        r3 = fp.add_run("\tPage "); r3.font.size = Pt(9); _set_font(r3)
        _add_field(fp, "PAGE")
        tp = ftr.add_paragraph()
        tp.paragraph_format.space_before = Pt(2); tp.paragraph_format.space_after = Pt(0)
        tp.paragraph_format.line_spacing = 1.0
        tr = tp.add_run("Buy Courses From Telegram Id = @M_Pharmacy_OwnerBot")
        tr.bold = True; tr.font.size = Pt(13); tr.font.color.rgb = TG_RED; _set_font(tr, "Arial")

    # ------------------------------------------------------------ building blocks
    def title_page(self):
        d = self.doc
        for _ in range(5):
            d.add_paragraph()
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("M Pharmacy Notes")
        r.bold = True; r.font.size = Pt(32); r.font.color.rgb = BLACK; _set_font(r)
        p.paragraph_format.space_after = Pt(12)

        for line, col in [("Formulation Development of", RED),
                          ("Pharmaceutical and", RED),
                          ("Cosmetic Products", RED)]:
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line); r.bold = True; r.font.size = Pt(26)
            r.font.color.rgb = col; _set_font(r)
            p.paragraph_format.space_after = Pt(2)
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("(MPH204T)")
        r.bold = True; r.font.size = Pt(18); r.font.color.rgb = BLACK; _set_font(r)
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("M. Pharmacy \u2013 Pharmaceutics \u2013 Semester II")
        r.italic = True; r.font.size = Pt(14); r.font.color.rgb = GREY_TXT; _set_font(r)

        for _ in range(6):
            d.add_paragraph()
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Connect With Us")
        r.bold = True; r.underline = True; r.font.size = Pt(16); _set_font(r)
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Website - "); r.bold = True; r.font.size = Pt(14); _set_font(r)
        r = p.add_run("www.carewellpharmacy.in")
        r.font.size = Pt(14); r.font.color.rgb = BLUE_LINK; _set_font(r)
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Telegram - "); r.bold = True; r.font.size = Pt(14); _set_font(r)
        r = p.add_run("Carewell Pharmacy"); r.font.size = Pt(14); _set_font(r)
        self.page_break()

    def page_break(self):
        self.doc.add_page_break()

    def unit_title(self, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.line_spacing = 1.15
        _para_border(p, 'bottom', MAROON, sz=12, space=4)
        _keep_with_next(p)
        r = p.add_run(text)
        r.bold = True; r.font.size = Pt(20); r.font.color.rgb = RED; _set_font(r)
        return p

    def h2(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        _keep_with_next(p)
        r = p.add_run(text)
        r.bold = True; r.font.size = Pt(16); r.font.color.rgb = RED; _set_font(r)
        return p

    def h3(self, text, color=BLACK):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        _keep_with_next(p)
        r = p.add_run(text)
        r.bold = True; r.font.size = Pt(14); r.font.color.rgb = color; _set_font(r)
        return p

    def para(self, text=None, segments=None, justify=True):
        p = self.doc.add_paragraph()
        if justify:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(8)
        if segments is None:
            segments = [(text, False, False)]
        for seg in segments:
            t, b, i = (seg + (False, False))[:3] if isinstance(seg, tuple) else (seg, False, False)
            r = p.add_run(t); r.bold = b; r.italic = i
            r.font.size = Pt(BODY_SIZE); r.font.color.rgb = BLACK; _set_font(r)
        return p

    def formula(self, text):
        """Centered emphasised formula line."""
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        r.bold = True; r.font.size = Pt(14); r.font.color.rgb = PURPLE; _set_font(r)
        return p

    def term(self, term, definition):
        return self.para(segments=[(term + " ", True, False), (definition, False, False)])

    def bullets(self, items, level=0):
        for it in items:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(4)
            if level == 0:
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.first_line_indent = Inches(-0.28)
                mark = p.add_run("\uf0d8")
                mark.font.size = Pt(12); mark.font.color.rgb = BLACK; _set_font(mark, "Wingdings")
                tab = p.add_run("  "); _set_font(tab); tab.font.size = Pt(BODY_SIZE)
            else:
                p.paragraph_format.left_indent = Inches(0.8)
                p.paragraph_format.first_line_indent = Inches(-0.22)
                mark = p.add_run("\u2022  ")
                mark.font.size = Pt(12); mark.font.color.rgb = TG_RED; _set_font(mark, "Arial")
            self._fill_runs(p, it)

    def numbered(self, items):
        for idx, it in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.first_line_indent = Inches(-0.28)
            num = p.add_run("%d. " % idx)
            num.bold = True; num.font.size = Pt(BODY_SIZE); _set_font(num)
            self._fill_runs(p, it)

    def _fill_runs(self, p, it):
        if isinstance(it, tuple) and len(it) == 2 and isinstance(it[1], str):
            r = p.add_run(it[0] + " "); r.bold = True
            r.font.size = Pt(BODY_SIZE); r.font.color.rgb = BLACK; _set_font(r)
            r2 = p.add_run(it[1]); r2.font.size = Pt(BODY_SIZE)
            r2.font.color.rgb = BLACK; _set_font(r2)
        elif isinstance(it, list):
            for seg in it:
                t, b, i = (seg + (False, False))[:3]
                r = p.add_run(t); r.bold = b; r.italic = i
                r.font.size = Pt(BODY_SIZE); r.font.color.rgb = BLACK; _set_font(r)
        else:
            r = p.add_run(it); r.font.size = Pt(BODY_SIZE)
            r.font.color.rgb = BLACK; _set_font(r)

    def table(self, headers, rows, widths=None, fontsize=11):
        n = len(headers)
        t = self.doc.add_table(rows=1, cols=n)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.allow_autofit = False
        _table_borders(t)
        hcells = t.rows[0].cells
        for j, htext in enumerate(headers):
            _shade(hcells[j], MAROON); _cell_margins(hcells[j])
            hcells[j].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cp = hcells[j].paragraphs[0]
            cp.paragraph_format.space_after = Pt(0); cp.paragraph_format.line_spacing = 1.0
            r = cp.add_run(htext)
            r.bold = True; r.font.size = Pt(fontsize + 0.5); r.font.color.rgb = WHITE; _set_font(r)
        for ridx, row in enumerate(rows):
            cells = t.add_row().cells
            fill = ROW_A if ridx % 2 == 0 else ROW_B
            for j, val in enumerate(row):
                _shade(cells[j], fill); _cell_margins(cells[j])
                cells[j].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                cp = cells[j].paragraphs[0]
                cp.paragraph_format.space_after = Pt(0); cp.paragraph_format.line_spacing = 1.05
                r = cp.add_run(str(val))
                r.font.size = Pt(fontsize); r.font.color.rgb = BLACK; _set_font(r)
        if widths:
            for j, w in enumerate(widths):
                for row in t.rows:
                    row.cells[j].width = Inches(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def box(self, title, lines, kind="HIGH-YIELD"):
        """A shaded callout box with a coloured left border."""
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.rows[0].cells[0]
        cell.width = Inches(6.4)
        _cell_box(cell, BOX_FILL, BOX_EDGE)
        _cell_margins(cell, top=90, bottom=90, left=160, right=140)
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_after = Pt(3); cp.paragraph_format.line_spacing = 1.2
        r = cp.add_run(("\u2605 " + title) if title else "\u2605")
        r.bold = True; r.font.size = Pt(13); r.font.color.rgb = TG_RED; _set_font(r)
        for ln in lines:
            bp = cell.add_paragraph()
            bp.paragraph_format.space_after = Pt(2); bp.paragraph_format.line_spacing = 1.25
            bp.paragraph_format.left_indent = Inches(0.2)
            bp.paragraph_format.first_line_indent = Inches(-0.18)
            m = bp.add_run("\u2022  "); m.font.size = Pt(11); _set_font(m, "Arial")
            self._fill_runs(bp, ln)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def figure(self, filename, caption, width=5.2):
        import os
        path = os.path.join(FIG_DIR, filename)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        if os.path.exists(path):
            run.add_picture(path, width=Inches(width))
        self.fig_no += 1
        cp = self.doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(8); cp.paragraph_format.line_spacing = 1.1
        r = cp.add_run("Fig. %d : %s" % (self.fig_no, caption))
        r.italic = True; r.bold = True; r.font.size = Pt(11); r.font.color.rgb = GREY_TXT; _set_font(r)

    def mcq_header(self, text):
        self.h2(text)

    def mcq(self, question, options, answer_idx, explanation=None):
        """question: str; options: list[str]; answer_idx: 0-based; explanation: str."""
        self.mcq_no += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run("Q%d. " % self.mcq_no); r.bold = True
        r.font.size = Pt(BODY_SIZE); r.font.color.rgb = PURPLE; _set_font(r)
        r2 = p.add_run(question); r2.font.size = Pt(BODY_SIZE); _set_font(r2)
        letters = "abcdefgh"
        for i, opt in enumerate(options):
            op = self.doc.add_paragraph()
            op.paragraph_format.left_indent = Inches(0.5)
            op.paragraph_format.space_after = Pt(1); op.paragraph_format.line_spacing = 1.2
            rr = op.add_run("(%s) %s" % (letters[i], opt))
            rr.font.size = Pt(13); _set_font(rr)
        ap = self.doc.add_paragraph()
        ap.paragraph_format.left_indent = Inches(0.5)
        ap.paragraph_format.space_after = Pt(6); ap.paragraph_format.line_spacing = 1.2
        ans = ap.add_run("Ans : (%s) %s" % (letters[answer_idx], options[answer_idx]))
        ans.bold = True; ans.font.size = Pt(13); ans.font.color.rgb = GREEN; _set_font(ans)
        if explanation:
            er = ap.add_run("   \u2013 " + explanation)
            er.italic = True; er.font.size = Pt(12); er.font.color.rgb = GREY_TXT; _set_font(er)

    def spacer(self, pts=4):
        self.doc.add_paragraph().paragraph_format.space_after = Pt(pts)

    def save(self, path):
        self.doc.save(path)


if __name__ == "__main__":
    import os
    import content
    here = os.path.dirname(os.path.abspath(__file__))
    FIG_DIR = os.path.join(here, "figures")   # set global on the running module
    n = Notes()
    n.title_page()
    content.build(n)
    root = os.path.dirname(here)
    out = os.path.join(root, "Formulation Development of Pharmaceutical and Cosmetic Products.docx")
    n.save(out)
    print("Saved", out, "| figures:", n.fig_no, "| MCQs:", n.mcq_no)
