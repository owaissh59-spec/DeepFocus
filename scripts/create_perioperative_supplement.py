#!/usr/bin/env python3
"""
Generates supplementary study notes (DOCX) for the syllabus unit:
  "Preparation of Patient for Operation, Pre & Post Operative Patient Care"
covering ONLY what is missing / thin in the already-studied PDF
(Chapters 1-3: Perioperative Care of the Patient / Protection of the Patient in
Surgery / Safety Measures for Operating Room Personnel).

Layout spec: LEGAL page (8.5 x 14 in), all margins 1.27 cm, compact spacing,
coloured headings, shaded tables, highlight boxes.
"""

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------- palette
C_H1 = RGBColor(0x1F, 0x3B, 0x73)      # deep blue
C_H2 = RGBColor(0x1B, 0x6B, 0x3A)      # green
C_H3 = RGBColor(0x6A, 0x28, 0x8F)      # purple
C_ANS = RGBColor(0x9C, 0x27, 0x10)     # rust (answers)
C_TXT = RGBColor(0x20, 0x20, 0x20)

SH_HDR = "1F3B73"      # table header fill
SH_ALT = "EEF3FA"      # alternating row fill
SH_CLIN = "FFF4E5"     # clinical box
SH_MNE = "E8F5E9"      # mnemonic box
SH_HY = "FDECEF"       # high-yield box
SH_GAP = "F3E9FB"      # gap-alert box

BODY = 9.5
TBL = 8.0

doc = Document()

# ---------------------------------------------------------------- page setup
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(14)
for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(sec, m, Cm(1.27))

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(BODY)
st.font.color.rgb = C_TXT
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.space_after = Pt(1.5)
st.paragraph_format.line_spacing = 1.0
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


# ---------------------------------------------------------------- helpers
def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _cell_margins(table, top=12, bottom=12, left=60, right=60):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement("w:" + tag)
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)


def _borders(table, color="9DB3D0", sz=4):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for tag in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + tag)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def rich(p, text, size=BODY, color=None):
    """Inline markup: **bold**, __italic__, ~~bold-italic~~"""
    import re
    tokens = re.split(r"(\*\*.+?\*\*|__.+?__|~~.+?~~)", text)
    for t in tokens:
        if not t:
            continue
        if t.startswith("**") and t.endswith("**"):
            r = p.add_run(t[2:-2]); r.bold = True
        elif t.startswith("__") and t.endswith("__"):
            r = p.add_run(t[2:-2]); r.italic = True
        elif t.startswith("~~") and t.endswith("~~"):
            r = p.add_run(t[2:-2]); r.bold = True; r.italic = True
        else:
            r = p.add_run(t)
        r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
    return p


def h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    r.font.size = Pt(15)
    r.bold = True
    r.font.color.rgb = C_H1
    # bottom rule
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "12")
    bt.set(qn("w:space"), "2"); bt.set(qn("w:color"), "1F3B73")
    pbdr.append(bt); pPr.append(pbdr)
    return p


def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = C_H2
    return p


def h3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.bold = True
    r.font.color.rgb = C_H3
    return p


def para(text, size=BODY, before=0, after=1.5, justify=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, text, size=size)
    return p


def bullets(items, size=BODY, indent=0.42, marker="\u25aa"):
    for it in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(indent)
        p.paragraph_format.first_line_indent = Cm(-0.28)
        p.paragraph_format.space_after = Pt(0.8)
        rich(p, marker + "  " + it, size=size)


def numbered(items, size=BODY, indent=0.5):
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(indent)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.space_after = Pt(0.8)
        rich(p, "**%d.**  %s" % (i, it), size=size)


def table(headers, rows, widths=None, size=TBL, header_fill=SH_HDR):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _borders(t)
    _cell_margins(t)
    hdr = t.rows[0]
    for i, htxt in enumerate(headers):
        c = hdr.cells[i]
        c.text = ""
        _shade(c, header_fill)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(htxt)
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            c = cells[ci]
            c.text = ""
            if ri % 2 == 1:
                _shade(c, SH_ALT)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            rich(p, str(val), size=size)
    if widths:
        for ci, w in enumerate(widths):
            for row in t.rows:
                row.cells[ci].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def box(title, lines, fill=SH_CLIN, border="D9A15B", title_color=None):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    _borders(t, color=border, sz=8)
    _cell_margins(t, top=40, bottom=40, left=100, right=100)
    c = t.rows[0].cells[0]
    c.text = ""
    _shade(c, fill)
    c.width = Cm(19.0)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = title_color or C_H1
    for ln in lines:
        q = c.add_paragraph()
        q.paragraph_format.space_after = Pt(0.5)
        q.paragraph_format.left_indent = Cm(0.15)
        rich(q, ln, size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def divider():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "6")
    bt.set(qn("w:space"), "1"); bt.set(qn("w:color"), "B9A0D6")
    pbdr.append(bt); pPr.append(pbdr)


MCQ_N = [0]


def mcq(q, opts, ans, expl):
    """opts: list of 4 strings; ans: index 0-3"""
    MCQ_N[0] += 1
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0.5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("Q%d. " % MCQ_N[0]); r.bold = True; r.font.size = Pt(9.5)
    r.font.color.rgb = C_H1
    rich(p, q, size=9.5)
    letters = "abcd"
    o = doc.add_paragraph()
    o.paragraph_format.left_indent = Cm(0.45)
    o.paragraph_format.space_after = Pt(0.5)
    o.paragraph_format.keep_with_next = True
    rich(o, "   ".join("(%s) %s" % (letters[i], opts[i]) for i in range(len(opts))), size=9)
    a = doc.add_paragraph()
    a.paragraph_format.left_indent = Cm(0.45)
    a.paragraph_format.space_after = Pt(2)
    ra = a.add_run("Ans: (%s) %s \u2014 " % (letters[ans], opts[ans]))
    ra.bold = True; ra.font.size = Pt(9); ra.font.color.rgb = C_ANS
    rich(a, expl, size=9, color=C_ANS)


# ================================================================ TITLE
tp = doc.add_paragraph()
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
tp.paragraph_format.space_after = Pt(1)
r = tp.add_run("SUPPLEMENTARY NOTES \u2014 WHAT YOUR PDF DID NOT COVER")
r.bold = True; r.font.size = Pt(17); r.font.color.rgb = C_H1

sp = doc.add_paragraph()
sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
sp.paragraph_format.space_after = Pt(1)
r = sp.add_run("Preparation of Patient for Operation \u2022 Pre- & Post-Operative Patient Care \u2022 "
               "Protection of the Patient in Surgery \u2022 OT Safety for the Medical Assistant / Pharmacist")
r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = C_H2

sp2 = doc.add_paragraph()
sp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
sp2.paragraph_format.space_after = Pt(4)
r = sp2.add_run("Gap-analysis of the studied PDF (Ch. 1 Perioperative Care of the Patient \u2022 "
                "Ch. 2 Protection of the Patient in Surgery / Patient Safety \u2022 "
                "Ch. 3 Safety Measures for Operating Room Personnel) + complete MCQ-proofing")
r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

box("HOW TO USE THIS DOCUMENT",
    ["Your PDF is **strong** on: psychological support, patient identification/Universal Protocol, the perioperative "
     "record forms, surgical positions, counting procedure, the four sterilization modalities, ESU precautions, "
     "orientation / in-service education, body mechanics, fatigue and waste anaesthetic gases. **Do not re-read those "
     "in depth \u2014 revise them from the PDF.**",
     "This document contains **only the deficits**: every syllabus sub-point where the PDF is silent or gives "
     "qualitative text where examiners ask for **numbers, classifications, doses and mechanisms**.",
     "\u25b6 **Biggest single gap: POST-OPERATIVE PATIENT CARE.** The syllabus heading says \u201cPre & Post Operative "
     "Patient Care\u201d, but the PDF stops at handing the patient over to PACU. Part C is therefore essential.",
     "Sources followed: AORN Guidelines for Perioperative Practice; Berry & Kohn's Operating Room Technique; "
     "CDC/HICPAC Guidelines (SSI, disinfection & sterilization, isolation precautions); WHO Surgical Safety Checklist "
     "& Hand Hygiene guidelines; AAGBI/ASA fasting & MH guidance; ICRP/AERB radiation limits; OSHA/NIOSH exposure "
     "limits; Biomedical Waste Management Rules 2016 (India); Bailey & Love; Lawrence's Essential Surgical Specialties."],
    fill=SH_GAP, border="9B6FC4", title_color=C_H3)


# ================================================================ COVERAGE AUDIT
h1("0.  Coverage Audit \u2014 Syllabus vs. Your PDF")

para("Read this table first. \u2705 = safe to revise from the PDF alone. \u26a0 = PDF is partial, the exam-relevant "
     "detail is in this document. \u274c = essentially absent from the PDF.")

table(
    ["Syllabus sub-topic", "Status", "Where the PDF covers it", "What is still missing \u2192 read here"],
    [
        ["Psychological support of the surgical patient", "\u2705 + \u26a0",
         "Ch.1, pp. 1\u201311 (very detailed)",
         "Anxiety assessment scales, coping mechanisms, pre-op teaching models, special groups \u2192 **\u00a7A1**"],
        ["Preparation of patient for operation (physical prep)", "\u274c",
         "Only the checklist form (Form 1-1)",
         "Fasting rules, hair removal, bathing, bowel/bladder prep, premedication, ASA grading, investigations, "
         "VTE & antibiotic prophylaxis, ERAS \u2192 **\u00a7A2\u2013A9**"],
        ["Informed consent", "\u26a0",
         "Forms 1-2, 1-3 reproduced",
         "Legal elements, types, who may consent, age, IPC sections, exceptions, medico-legal doctrines \u2192 **\u00a7A10**"],
        ["Admission procedure / patient identification", "\u2705",
         "Ch.2, pp. 18\u201319",
         "Only add the **WHO Surgical Safety Checklist** (3 phases) which the PDF omits \u2192 **\u00a7B1**"],
        ["Transfer procedure & position", "\u2705 + \u26a0",
         "Ch.2, pp. 19\u201330, Fig. 2-1 A\u2013J",
         "Position-specific **complications**, nerve-injury map, physiological effects, angles/degrees, "
         "transfer-device load limits \u2192 **\u00a7B6\u2013B7**"],
        ["Environmental controls", "\u26a0",
         "Ch.2, pp. 33\u201335 (qualitative)",
         "OR temperature/humidity/air-change/pressure numbers, HEPA & laminar flow, zoning, cleaning schedule, "
         "disinfectant table, SSI classification, hand-hygiene timings, sterile-field rules \u2192 **\u00a7B2\u2013B5**"],
        ["Electro surgery", "\u26a0",
         "Ch.2, pp. 38\u201339 (5 precautions)",
         "Frequency/waveforms, monopolar vs bipolar circuits, alternate-site burns, capacitive & direct coupling, "
         "insulation failure, REM/AEM, pacemaker-ICD protocol, advanced energy devices, surgical smoke \u2192 **\u00a7B8**"],
        ["Operative records", "\u2705 + \u26a0",
         "Forms on pp. 3\u20134, 12\u201317",
         "Specimen handling & labelling, implant log, incident report rules, retention, what makes a record "
         "legally defensible \u2192 **\u00a7B11**"],
        ["Counting procedure", "\u2705",
         "Ch.2, pp. 39\u201341 (9 rules)",
         "Retained surgical items epidemiology, gossypiboma, RF-detection technology \u2192 **\u00a7B10**"],
        ["Sterilization", "\u26a0",
         "Ch.2, pp. 41\u201345 (4 modalities)",
         "Definitions & Spaulding classification, **exact time\u2013temperature\u2013pressure tables**, dry heat, "
         "radiation, filtration, physical/chemical/biological indicators, SAL, packaging & shelf-life, "
         "levels of disinfection \u2192 **\u00a7B12**"],
        ["Emergencies and disasters", "\u26a0",
         "Ch.2, pp. 47\u201349 (CPR cart, power failure, MH, death in OR)",
         "BLS/ACLS numbers, anaphylaxis doses, full MH protocol with dantrolene dosing, LAST, airway fire drill, "
         "extinguisher classes, RACE/PASS, mass-casualty triage, code colours \u2192 **\u00a7B13**"],
        ["**POST-OPERATIVE PATIENT CARE**", "\u274c",
         "PACU hand-over paragraph only (p. 8) + PACU form",
         "PACU monitoring, **Aldrete & PADSS scores**, immediate complications, PONV, pain, post-op fever timeline, "
         "wound healing & dehiscence, drains, fluids, DVT, ileus, discharge counselling \u2192 **entire Part C**"],
        ["In-service education", "\u2705", "Ch.3, pp. 50\u201351", "Nothing significant \u2014 revise from PDF"],
        ["Body mechanics", "\u2705 + \u26a0", "Ch.3, p. 51",
         "NIOSH lifting limits, safe-patient-handling numbers \u2192 **\u00a7D1**"],
        ["Fatigue factors", "\u2705", "Ch.3, pp. 51\u201352", "Nothing significant \u2014 revise from PDF"],
        ["Radiation safety", "\u26a0", "Ch.3, p. 52 (4 measures)",
         "Units (Gy/Sv/rem), **dose limits**, ALARA, inverse-square law, lead-equivalence, dosimeter types, "
         "stochastic vs deterministic effects, pregnancy limits \u2192 **\u00a7D2**"],
        ["Infection control", "\u26a0", "Ch.3, pp. 53\u201354",
         "Standard vs transmission-based precautions, PPE donning/doffing order, needle-stick transmission risks & "
         "PEP, immunisation, **biomedical waste colour coding** \u2192 **\u00a7D3**"],
        ["Chemical hazards", "\u26a0", "Ch.3, pp. 54\u201356",
         "OSHA/NIOSH **exposure limit numbers**, glutaraldehyde/formaldehyde/EO, methyl methacrylate, cytotoxic drug "
         "handling, SDS & GHS, spill management \u2192 **\u00a7D4**"],
    ],
    widths=[4.6, 1.5, 4.4, 8.5])

box("EXAM-STRATEGY NOTE",
    ["Roughly **60\u201365 %** of MCQs from this unit in Indian paper patterns come from just five clusters: "
     "**(1)** sterilization time\u2013temperature\u2013indicator numbers, **(2)** post-operative complications & scoring "
     "systems, **(3)** electrosurgery hazards, **(4)** patient positioning nerve injuries, "
     "**(5)** radiation / chemical / biomedical-waste numbers.",
     "Four of those five are the PDF's weak areas. Prioritise **\u00a7B12, Part C, \u00a7B8, \u00a7B6, \u00a7D2\u2013D4**."],
    fill=SH_HY, border="D98BA0", title_color=C_ANS)


# ================================================================ PART A
h1("Part A  \u2014  Preparation of the Patient for Operation (Pre-Operative)")

h2("A1.  Psychological Support \u2014 the frameworks your PDF omits")
para("The PDF describes __what__ staff do (telephone call, warm blanket, holding the hand, acting as advocate). "
     "Examiners also ask __why__ and by __what instrument__.")
table(["Concept", "Key content"],
      [["Sources of pre-operative anxiety (classic 7)",
        "Fear of the unknown \u2022 fear of death \u2022 fear of anaesthesia / not waking up \u2022 fear of pain \u2022 "
        "fear of disfigurement & altered body image \u2022 fear of loss of control / dependence \u2022 "
        "separation from family + financial concern"],
       ["Assessment tools",
        "**STAI** (State\u2013Trait Anxiety Inventory) \u2022 **APAIS** (Amsterdam Preoperative Anxiety & Information "
        "Scale) \u2022 **HADS** \u2022 Visual Analogue Anxiety Scale \u2022 for children the **m-YPAS** "
        "(modified Yale Preoperative Anxiety Scale)"],
       ["Peak anxiety times",
        "On being told surgery is needed \u2022 night before \u2022 **immediately before transfer to OR (highest)** "
        "\u2022 in the holding area"],
       ["Effects of unrelieved anxiety",
        "\u2191 catecholamines \u2192 tachycardia, hypertension, arrhythmia \u2022 \u2191 anaesthetic requirement "
        "\u2022 \u2191 post-op pain scores & analgesic need \u2022 \u2191 PONV \u2022 delayed wound healing "
        "(cortisol) \u2022 longer hospital stay"],
       ["Evidence-based interventions",
        "Structured **pre-operative teaching/counselling** (single most effective) \u2022 information giving matched "
        "to coping style (\u201cmonitors\u201d want detail, \u201cblunters\u201d want less) \u2022 relaxation & guided "
        "imagery \u2022 music therapy \u2022 parental presence at induction / midazolam premedication in children "
        "\u2022 pre-op visit to the OR for children \u2022 allowing spectacles, hearing aids and dentures to remain "
        "until induction"],
       ["Special groups",
        "**Paediatric** \u2014 fear of separation (1\u20133 y), fear of mutilation (3\u20136 y), fear of loss of "
        "control (school age); **Geriatric** \u2014 fear of dependency & death, sensory deficits, needs slower, "
        "repeated explanation; **Cancer/amputation** \u2014 anticipatory grief, body-image work; "
        "**Jehovah's Witness** \u2014 document blood refusal (PDF Form 1-1 \u201crefuses transfusions\u201d box)"],
       ["Legal-ethical anchor",
        "The patient's **autonomy** is protected by informed consent; the perioperative nurse/practitioner is the "
        "**patient's advocate** \u2014 the only team member whose sole duty is the patient's interest while the "
        "patient is anaesthetised"]],
      widths=[4.4, 14.6])

h2("A2.  Pre-Operative Fasting (\u201cNPO\u201d) \u2014 the PDF only has a blank \u201cNPO since ___\u201d line")
para("**ASA / ESA / AAGBI minimum fasting intervals before elective anaesthesia** (healthy patients, no risk of "
     "delayed gastric emptying). Remember as **2\u20134\u20136\u20138**.")
table(["Ingested material", "Minimum fast", "Notes"],
      [["Clear liquids (water, pulp-free juice, black tea/coffee, clear carbohydrate drink)", "**2 hours**",
        "Volume is less important than type; ERAS actively encourages a CHO drink 2 h pre-op"],
       ["Breast milk", "**4 hours**", "Neonates & infants"],
       ["Infant formula, non-human milk, light meal (toast + clear fluid)", "**6 hours**",
        "Milk behaves like a solid (curdles)"],
       ["Fried / fatty food, meat, full meal", "**8 hours**", "May need longer"],
       ["Chewing gum / boiled sweet", "Stop before OR (do not cancel)", "Not an indication to postpone in adults"]],
      widths=[7.6, 3.2, 8.2])
box("WHY NPO? \u2014 CLINICAL CORRELATION",
    ["Purpose = prevent **pulmonary aspiration of gastric contents (Mendelson's syndrome)** during induction, when "
     "protective airway reflexes are abolished.",
     "Risk becomes significant when gastric volume is **> 0.4 mL/kg (\u2248 25 mL)** and **pH < 2.5**.",
     "**Full-stomach / high-risk patients** (emergency, trauma, pregnancy \u2265 2nd trimester, obesity, diabetic "
     "gastroparesis, bowel obstruction, GERD, opioids) \u2192 treat as **not fasted** \u2192 rapid sequence induction "
     "with cricoid pressure (**Sellick's manoeuvre**), aspiration prophylaxis (H\u2082 blocker / PPI + sodium citrate "
     "\u00b1 metoclopramide).",
     "Prolonged (\u201cmidnight\u201d) fasting is now considered **harmful** \u2014 dehydration, hypoglycaemia, "
     "insulin resistance, irritability, and it does **not** reduce gastric volume."],
    fill=SH_CLIN)

h2("A3.  Routine Pre-Operative Investigations")
para("The PDF lists \u201chistory & physical, haematology, blood chemistry, urinalysis, chest X-ray, EKG\u201d. "
     "Modern practice is **selective, not routine** \u2014 a classic MCQ discriminator.")
table(["Investigation", "Indication / current recommendation"],
      [["Hb / PCV", "All major surgery; women of child-bearing age; anticipated blood loss > 500 mL"],
       ["Blood grouping & cross-match / \u201cgroup & save\u201d",
        "Per the surgical blood-order schedule; availability **confirmed by the circulator before surgery starts**"],
       ["Coagulation profile (PT/INR, aPTT)",
        "Only if bleeding history, liver disease, anticoagulants, or major/cardiac surgery \u2014 **not routine**"],
       ["Blood sugar / HbA1c", "Diabetics, > 40 y, obesity; HbA1c ideally < 8 % for elective surgery"],
       ["Renal function, electrolytes", "> 60 y, hypertension, diabetes, diuretics, renal disease, major surgery"],
       ["Urinalysis", "Urological surgery, implant surgery, symptoms; not routine otherwise"],
       ["ECG", "Age > 40\u201350 y (institution-specific), cardiac disease, hypertension, diabetes, "
                "electrolyte disturbance"],
       ["Chest X-ray", "Cardio-respiratory disease, > 60 y, thoracic surgery, suspected TB/malignancy \u2014 "
                       "**not routine** in healthy young adults"],
       ["Pulmonary function tests / ABG", "Thoracic resection, severe COPD, obesity-hypoventilation"],
       ["Pregnancy test", "**All women of child-bearing potential must be asked; test if any doubt** "
                          "(also a radiation-safety requirement \u2014 see \u00a7D2)"],
       ["Viral markers (HBsAg, anti-HCV, HIV)", "As per institutional policy; result does **not** change universal "
                                                "application of Standard Precautions"]],
      widths=[5.2, 13.8])

h2("A4.  Risk Stratification \u2014 ASA grade, Mallampati, and the scores examiners love")
h3("ASA Physical Status Classification (ASA-PS)")
table(["Class", "Definition", "Example"],
      [["**ASA I**", "Normal healthy patient", "Fit non-smoker, no or minimal alcohol"],
       ["**ASA II**", "Mild systemic disease **without** functional limitation",
        "Controlled hypertension or diabetes, current smoker, pregnancy, obesity BMI 30\u201340"],
       ["**ASA III**", "Severe systemic disease with functional limitation",
        "Poorly controlled DM/HTN, COPD, BMI \u2265 40, MI/CVA > 3 months ago, dialysis"],
       ["**ASA IV**", "Severe disease that is a **constant threat to life**",
        "MI/CVA/stent **< 3 months**, ongoing cardiac ischaemia, sepsis, ejection fraction < 25 %"],
       ["**ASA V**", "Moribund, not expected to survive without the operation",
        "Ruptured aneurysm, massive trauma, intracranial bleed with mass effect"],
       ["**ASA VI**", "Declared brain-dead, organs being removed for donation", "Organ harvest"],
       ["**suffix E**", "Emergency operation", "ASA II E = otherwise healthy patient for emergency appendicectomy"]],
      widths=[1.8, 7.2, 10.0])
box("HIGH-YIELD TRAPS",
    ["ASA grade reflects **only the patient's physiological status** \u2014 not the magnitude of surgery, not age, "
     "not the anaesthetic risk per se.",
     "Pregnancy = **ASA II** (it is a physiological, not pathological, state but not \u2018normal healthy\u2019).",
     "**Mallampati** grades the __oropharyngeal view__ (airway), not overall fitness: **I** soft palate + uvula + "
     "fauces + pillars; **II** soft palate + uvula + fauces; **III** soft palate + **base of uvula only**; "
     "**IV** soft palate **not** visible. Grades III\u2013IV predict difficult intubation.",
     "Other scores worth a line each: **Goldman / Lee's Revised Cardiac Risk Index** (cardiac risk), "
     "**Caprini score** (VTE risk), **Apfel score** (PONV \u2013 see \u00a7C4), **Aldrete score** "
     "(PACU discharge \u2013 \u00a7C1), **Glasgow Coma Scale** (neuro)."],
    fill=SH_HY, border="D98BA0", title_color=C_ANS)

h2("A5.  Skin Preparation & Hair Removal \u2014 evidence the PDF predates")
table(["Question", "Current evidence-based answer"],
      [["Should hair be removed?",
        "**No \u2014 do not remove hair unless it physically interferes** with the procedure or wound closure"],
       ["If it must be removed, how?",
        "**Electric clippers with a single-use head** (or depilatory cream). **Razors are contraindicated** \u2014 "
        "shaving causes micro-abrasions that get colonised"],
       ["When?",
        "**Immediately before surgery, outside the operating room.** Risk of SSI: no removal \u2248 clipping < "
        "shaving immediately before < shaving the night before (**highest**)"],
       ["Pre-operative bath",
        "Whole-body wash with soap or **4 % chlorhexidine gluconate** the night before **and** on the day of surgery"],
       ["Best intra-operative skin antiseptic",
        "**Alcohol-based 2 % chlorhexidine gluconate** (e.g. ChloraPrep) is superior to aqueous povidone-iodine for "
        "clean-contaminated surgery. Povidone-iodine 10 % (1 % free iodine) is the alternative"],
       ["Alcohol + electrosurgery",
        "Alcoholic preps must be allowed to **dry completely (\u2265 3 min, longer in hair)** and must **never pool** "
        "\u2014 pooled alcohol is the classic **fuel** in OR fires (\u00a7B13)"],
       ["Direction of painting",
        "From the **incision line outward** in concentric circles/strokes; the sponge is **never returned** to the "
        "clean centre; the most contaminated area (stoma, sinus, axilla, perineum, umbilicus) is prepped **last**"],
       ["Contraindications",
        "Chlorhexidine: avoid in **middle ear** (ototoxic), **eye/cornea** (keratitis), meninges/neural tissue; "
        "iodophors: neonates (thyroid, burns), iodine allergy"],
       ["Extent of prep",
        "Wide enough to allow **extension of the incision, new incisions and drain sites** without re-prepping"]],
      widths=[4.6, 14.4])

h2("A6.  Surgical Antibiotic Prophylaxis")
table(["Parameter", "Standard"],
      [["Timing", "Single IV dose **within 60 minutes before skin incision** (within **120 min** for vancomycin and "
                  "fluoroquinolones, which need slow infusion). For tourniquet surgery, complete infusion **before "
                  "tourniquet inflation**"],
       ["Usual drug", "1st/2nd-generation cephalosporin \u2014 **cefazolin 2 g** (3 g if \u2265 120 kg); add "
                      "**metronidazole** for colorectal/appendicular; **vancomycin/clindamycin** if MRSA-colonised "
                      "or \u03b2-lactam allergic"],
       ["Re-dosing intra-op", "If procedure exceeds **2 half-lives** of the drug (\u2248 4 h for cefazolin) or blood "
                              "loss **> 1500 mL**"],
       ["Duration", "**Stop within 24 h** of surgery end (\u2264 48 h for cardiac surgery). Prolonging prophylaxis "
                    "does **not** reduce SSI and breeds resistance + C. difficile"],
       ["Prophylaxis vs therapy",
        "Prophylaxis = clean/clean-contaminated cases and implant surgery. **Dirty/infected wounds get therapeutic** "
        "antibiotics, not prophylaxis"],
       ["Other proven SSI bundles", "Maintain **normothermia > 36 \u00b0C**; peri-op **glucose < 180 mg/dL**; "
                                    "adequate tissue **oxygenation (FiO\u2082 0.8)**; avoid razors; CHG-alcohol prep"]],
      widths=[3.6, 15.4])

h2("A7.  Venous Thromboembolism (VTE) Prophylaxis & Peri-Operative Drug Management")
table(["Item", "Practice points"],
      [["Virchow's triad", "**Stasis + endothelial injury + hypercoagulability** \u2014 all three are produced by "
                           "surgery + immobility"],
       ["Mechanical prophylaxis", "Graduated compression (**TED**) stockings, **intermittent pneumatic compression / "
                                  "sequential compression devices**, early ambulation. (The PDF's forms mention TED & "
                                  "sequential hose but never explain why.) Contraindicated in peripheral arterial "
                                  "disease"],
       ["Pharmacological prophylaxis", "**LMWH** (enoxaparin 40 mg SC OD) or UFH 5000 U SC BD/TDS, fondaparinux, or "
                                       "DOACs after major orthopaedic surgery; start 12 h pre- or 6\u201312 h post-op"],
       ["Highest-risk surgery", "Hip & knee arthroplasty, hip fracture, major pelvic/abdominal cancer surgery, "
                                "neurosurgery, major trauma, spinal cord injury"],
       ["Aspirin", "Usually **continued** for secondary cardiovascular prevention; stop 7 days only for very "
                   "high-bleeding-risk surgery"],
       ["Clopidogrel / prasugrel / ticagrelor", "Stop **5\u20137 days** pre-op (ticagrelor 3\u20135 d) if bleeding "
                                                "risk outweighs stent thrombosis risk \u2014 discuss with cardiology"],
       ["Warfarin", "Stop **5 days**; check INR < 1.5; bridge with LMWH if high thrombotic risk (mechanical valve, "
                    "recent VTE, AF with prior stroke)"],
       ["DOACs (dabigatran, rivaroxaban, apixaban)", "Stop **24\u201348 h** (longer with renal impairment / high "
                                                     "bleed risk)"],
       ["Metformin", "Omit on the **morning of surgery** (lactic acidosis / contrast risk); restart when eating"],
       ["Insulin", "Reduce/omit short-acting; give ~half to two-thirds of usual basal; start "
                   "glucose\u2013insulin\u2013potassium infusion for major surgery; **monitor capillary glucose "
                   "hourly**"],
       ["Continue on the day of surgery", "\u03b2-blockers (abrupt stop \u2192 rebound ischaemia), statins, "
                                          "anti-epileptics, anti-parkinsonian drugs, inhalers, thyroxine, "
                                          "anti-hypertensives except ACE-I/ARB"],
       ["Withhold", "ACE inhibitors / ARBs on the morning of surgery (refractory hypotension), oral hypoglycaemics, "
                    "diuretics, herbal supplements **2 weeks** before (see box)"]],
      widths=[4.6, 14.4])
box("PHARMACIST'S HIGH-YIELD BOX \u2014 herbal products and surgery (\u201cthe 4 G's + E\u201d)",
    ["Stop **2 weeks** before elective surgery: **G**arlic, **G**inger, **G**inkgo biloba, **G**inseng "
     "(all \u2192 bleeding / platelet inhibition); **E**phedra (\u2192 hypertension, arrhythmia, interacts with "
     "anaesthetics); **St John's wort** (CYP3A4 induction \u2192 drug failure, serotonin syndrome); "
     "**Kava & Valerian** (\u2192 potentiate sedatives, withdrawal); **Vitamin E** and **fish oil** (\u2192 bleeding).",
     "This is exactly why the PDF's Form 1-1 insists that medications listed on the front of the chart "
     "**include OTC and herbal products**."],
    fill=SH_MNE, border="6DA97A", title_color=C_H2)

h2("A8.  Pre-Anaesthetic Medication (Premedication) \u2014 pharmacist-relevant")
table(["Purpose (\u201c7 A's\u201d)", "Drug class", "Examples"],
      [["**A**nxiolysis / amnesia", "Benzodiazepines", "Midazolam 0.02\u20130.05 mg/kg IV / 0.5 mg/kg oral (children), "
                                                       "diazepam, lorazepam"],
       ["**A**nalgesia", "Opioids, NSAIDs, paracetamol, gabapentinoids",
        "Fentanyl, morphine; pre-emptive paracetamol 1 g, gabapentin 300\u2013600 mg"],
       ["**A**ntisialagogue (dry secretions)", "Antimuscarinics", "Glycopyrrolate 0.2 mg (does not cross BBB \u2014 "
                                                                  "preferred), atropine, hyoscine"],
       ["**A**nti-emetic", "5-HT\u2083 antagonists, steroids, D\u2082 blockers, NK\u2081",
        "Ondansetron 4\u20138 mg, dexamethasone 4\u20138 mg, metoclopramide, aprepitant"],
       ["**A**nti-acid / aspiration prophylaxis", "H\u2082 blockers, PPIs, non-particulate antacid, prokinetics",
        "Ranitidine 150 mg oral / 50 mg IV, pantoprazole 40 mg, **sodium citrate 0.3 M** 30 mL, metoclopramide 10 mg"],
       ["**A**ttenuation of autonomic reflexes", "\u03b2-blockers, \u03b12-agonists, lignocaine",
        "Esmolol, clonidine, dexmedetomidine"],
       ["**A**ntibiotic prophylaxis", "See \u00a7A6", "Cefazolin 2 g IV within 60 min of incision"],
       ["Also: allergy prophylaxis", "H\u2081 + H\u2082 blockers + steroid",
        "For latex-sensitive or contrast-allergic patients (see \u00a7D3)"]],
      widths=[4.2, 5.0, 9.8])

h2("A9.  Bowel & Bladder Preparation, and ERAS")
bullets([
    "**Mechanical bowel preparation** (polyethylene glycol, sodium phosphate) \u00b1 oral antibiotics: now used "
    "selectively for **colorectal** surgery (oral antibiotics + MBP reduces SSI; MBP alone does not). Causes "
    "dehydration and electrolyte loss \u2014 hence the PDF's emphasis on fluid-volume nursing focus.",
    "**Bladder**: patient voids immediately before transfer; catheterise only if indicated (long procedure, pelvic "
    "surgery, need for hourly output monitoring). Remove the catheter **within 24 h** \u2014 CAUTI prevention.",
    "**Skin/site marking** by the operating surgeon with an indelible marker, patient awake and participating "
    "(\u00a7B1).",
    "**Remove** \u2014 dentures/bridges, contact lenses, jewellery (tape a wedding ring if it cannot be removed), "
    "nail polish & gel/artificial nails (interfere with pulse oximetry and harbour organisms), hairpins, prostheses; "
    "**document the disposition of every valuable** (exactly what Form 1-1 records).",
    "**ERAS (Enhanced Recovery After Surgery)** \u2014 the modern framework the PDF predates: pre-op counselling, "
    "no prolonged fasting + **carbohydrate drink 2 h pre-op**, no routine bowel prep, no routine drains/NG tubes, "
    "short-acting anaesthesia, regional/multimodal opioid-sparing analgesia, active warming, goal-directed fluids, "
    "**early oral feeding and mobilisation within 24 h**. Result: fewer complications and shorter stay.",
])

h2("A10.  Informed Consent \u2014 legal & medico-legal detail")
h3("The five elements of valid consent")
numbered([
    "**Disclosure** \u2014 diagnosis, nature and purpose of the proposed procedure, its risks and benefits, "
    "**alternatives** (including doing nothing) and their risks, and the consequences of refusal.",
    "**Capacity/competence** \u2014 the patient must understand, retain, weigh and communicate a decision.",
    "**Voluntariness** \u2014 free of coercion, fraud, duress or misrepresentation; obtained **before** sedative "
    "premedication (note the PDF's Form 1-2 wording: \u201cI am not under the influence of any preoperative or other "
    "sedating medication\u201d).",
    "**Comprehension** \u2014 in a language the patient understands; use an interpreter, not a family member, "
    "where possible.",
    "**Documentation & authentication** \u2014 signed and dated by patient/legal representative, the surgeon "
    "(affirmation of informed consent) and a witness; filed in the chart **before** the patient enters the surgical "
    "suite.",
])
h3("Types of consent, and who may give it")
table(["Type / situation", "Rule"],
      [["Implied consent", "Patient's conduct implies agreement (extending an arm for venepuncture). Sufficient for "
                           "routine examination only"],
       ["Expressed \u2014 oral or written", "Written required for any invasive/operative procedure, anaesthesia, "
                                            "blood transfusion, and for research"],
       ["Informed consent", "The legal standard for surgery/anaesthesia (elements above)"],
       ["Special/additional consent", "**Sterilization** (patient must state that they understand permanent loss of "
                                      "fertility \u2014 PDF Form 1-3), organ donation, HIV testing, "
                                      "photography/video, participation in teaching, implant use, blood products, "
                                      "high-risk procedures"],
       ["Age", "**18 years and above** in India (Indian Majority Act; **IPC \u00a7 87\u201388**). For a minor or an "
               "incompetent adult, the parent/legal guardian consents (**IPC \u00a7 89** \u2014 child under 12 or "
               "person of unsound mind)"],
       ["Emergency / unconscious patient", "**Doctrine of necessity / implied consent \u2014 IPC \u00a7 92**: "
                                           "life-saving treatment may proceed in good faith without consent when the "
                                           "patient cannot consent and no representative is available. Document the "
                                           "emergency"],
       ["Consent obtained by fear or misconception", "**Invalid \u2014 IPC \u00a7 90**"],
       ["Who must obtain it", "The **operating surgeon** (or a delegate competent to perform the procedure) \u2014 "
                              "**not** the nurse or pharmacist. The perioperative practitioner's duty is to "
                              "**verify** that a valid, signed, dated consent exists and matches the procedure, site "
                              "and side, and to report any discrepancy"],
       ["Withdrawal", "Consent may be withdrawn at **any time** before or during a procedure while the patient is "
                      "competent"],
       ["Refusal of blood", "Honour it; document on the checklist (PDF Form 1-1 \u201crefuses transfusions\u201d) and "
                            "plan cell-saver / haemostatics / tranexamic acid"]],
      widths=[4.6, 14.4])
box("MEDICO-LEGAL DOCTRINES \u2014 one-liners frequently asked",
    ["**Negligence (Bolam test)** \u2014 duty of care, breach of standard of a reasonably competent practitioner, "
     "causation, damage. The 4 D's: **D**uty, **D**ereliction, **D**irect causation, **D**amage.",
     "**Res ipsa loquitur** (\u2018the thing speaks for itself\u2019) \u2014 negligence is presumed; the classic "
     "example is a **retained surgical sponge or instrument**, wrong-site surgery, or a diathermy burn.",
     "**Vicarious liability / respondeat superior** \u2014 the hospital is liable for its employees' acts.",
     "**Captain-of-the-ship doctrine** \u2014 historically made the surgeon liable for everything in the OR; largely "
     "replaced by independent professional accountability of each team member.",
     "**Battery** \u2014 operating without any consent. **Negligence** \u2014 operating with inadequately informed "
     "consent.",
     "**Only what is documented is a legal fact** \u2014 the single most quoted sentence from your PDF (p. 8, p. 48). "
     "An undocumented event \u2018did not happen\u2019 in court."],
    fill=SH_HY, border="D98BA0", title_color=C_ANS)



# ================================================================ PART B
h1("Part B  \u2014  Protection of the Patient in Surgery (Gap-Fill)")

h2("B1.  The WHO Surgical Safety Checklist \u2014 completely absent from your PDF")
para("Your PDF teaches the **Joint Commission Universal Protocol (2004)** and the AORN 2005 addition. The **WHO "
     "Surgical Safety Checklist (2008, \u2018Safe Surgery Saves Lives\u2019)** is the internationally examined "
     "instrument. It has **19 items in 3 phases**, is led by a nominated **checklist coordinator**, and reduced "
     "mortality from 1.5 % to 0.8 % and complications from 11 % to 7 % in the original 8-country study.")
table(["Phase", "When", "Core content"],
      [["**1. SIGN IN**", "Before induction of anaesthesia (with anaesthetist + nurse)",
        "Patient confirms **identity, site, procedure, consent** \u2022 site marked \u2022 anaesthesia safety check "
        "completed \u2022 **pulse oximeter on and working** \u2022 known allergy? \u2022 difficult airway/aspiration "
        "risk? \u2022 risk of blood loss > 500 mL (7 mL/kg in children) \u2192 IV access/fluids/blood planned"],
       ["**2. TIME OUT**", "After induction, **before skin incision** (whole team pauses)",
        "All members **introduce name and role** \u2022 surgeon, anaesthetist and nurse verbally confirm patient, "
        "site, procedure \u2022 **anticipated critical events** (surgeon: critical steps, duration, blood loss; "
        "anaesthetist: patient-specific concerns; nurse: sterility indicator confirmed, equipment issues) \u2022 "
        "**antibiotic prophylaxis given in last 60 min?** \u2022 essential imaging displayed? \u2022 VTE prophylaxis "
        "\u2022 correct position"],
       ["**3. SIGN OUT**", "Before the patient leaves the operating room",
        "Nurse verbally confirms with the team: **name of the procedure recorded** \u2022 **instrument, sponge and "
        "needle counts are correct** \u2022 **specimens labelled** (including patient name, read aloud) \u2022 any "
        "equipment problems to be addressed \u2022 surgeon/anaesthetist/nurse state the **key concerns for recovery "
        "and management** of this patient"]],
      widths=[2.6, 4.4, 12.0])
box("DISTINGUISH THESE THREE \u2014 favourite MCQ",
    ["**Universal Protocol (JC, July 2004)** = pre-procedure **verification** + **site marking** + **\u201ctime "
     "out\u201d** immediately before starting. Aim: prevent wrong-site, wrong-procedure, wrong-person surgery.",
     "**AORN (March 2005)** added confirmation of the **correct patient position** to the time out.",
     "**WHO Checklist (2008)** = the 3-phase, 19-item tool (Sign in / Time out / Sign out) \u2014 broader than site "
     "verification; includes airway, blood loss, antibiotics, imaging, counts and specimen labelling.",
     "**Site marking rules:** by the **operating surgeon**, with an **indelible/permanent marker**, at or near the "
     "incision, **with the awake patient participating**, **before** the patient enters the OR; mark **both** "
     "laterality and level (spine); never mark the non-operative side with an \u2018X\u2019."],
    fill=SH_HY, border="D98BA0", title_color=C_ANS)

h2("B2.  Operating Room Design & Environmental Control \u2014 the numbers")
table(["Parameter", "Standard value", "Rationale / exam point"],
      [["Temperature", "**20\u201323 \u00b0C (68\u201375 \u00b0F)**", "Comfort of gowned team vs prevention of "
                                                                     "**patient hypothermia**; raise to 24\u201326 "
                                                                     "\u00b0C for neonates/burns"],
       ["Relative humidity", "**20\u201360 %** (classically 30\u201360 %)",
        "**Too low \u2192 static electricity/spark risk; too high \u2192 microbial growth & condensation**"],
       ["Air changes per hour", "Minimum **15**, recommended **20\u201325** (\u2265 3 of them outdoor air)",
        "Dilutes airborne contaminants and waste anaesthetic gas"],
       ["Air pressure", "**Positive** with respect to corridors and adjacent areas (\u2265 2.5 Pa / 0.01 in H\u2082O)",
        "Air flows **out** of the OR. (Contrast: TB isolation room = **negative** pressure)"],
       ["Airflow pattern", "Unidirectional, **downward** from ceiling diffusers over the table, exhaust at floor level "
                           "(low wall returns, \u2265 2)", "Keeps the clean air column over the sterile field"],
       ["HEPA filtration", "**High-Efficiency Particulate Air filter = \u2265 99.97 % of particles \u2265 0.3 "
                           "\u00b5m**", "The number most often asked. **ULPA** = 99.999 % at 0.1\u20130.2 \u00b5m "
                                        "(used in smoke evacuators)"],
       ["Laminar (ultraclean) airflow", "**300\u2013600 air changes/h**, unidirectional at 0.3\u20130.5 m/s",
        "Used for **joint arthroplasty and implant surgery**; reduces air bacterial counts to < 10 CFU/m\u00b3"],
       ["Doors", "Kept **closed**; traffic minimised", "Every door opening breaks the positive-pressure gradient; "
                                                      "each additional person adds bacteria-carrying skin scales"],
       ["Wall/floor finish", "Hard, non-porous, seamless, scrubbable; rounded corners; no false ceiling over the OR",
        "Prevents dust retention"],
       ["Lighting", "General 200\u2013300 lux; operative field **20 000\u2013100 000 lux**, shadowless, "
                    "cool (dichroic/LED)", "Heat from lights contributes to fluid evaporation from tissue"],
       ["Noise", "Recommended < 45 dB", "Noise impairs communication and team performance; silence mandatory during "
                                        "**induction** (your PDF, p. 8) and during counts"],
       ["Ideal OR size", "\u2248 400\u2013600 sq ft (37\u201360 m\u00b2); cardiac/orthopaedic larger",
        "Space for equipment without breaching the sterile field"]],
      widths=[3.4, 5.2, 10.4])

h3("Zoning of the operation theatre complex")
table(["Zone", "Also called", "Areas", "Attire"],
      [["**Protective / outer zone**", "Unrestricted", "Reception, waiting, changing rooms, offices, pre-op holding, "
                                                       "PACU", "Street clothes permitted"],
       ["**Clean zone**", "Semi-restricted", "Peripheral corridors, storage, utility, instrument processing, "
                                             "recovery corridor",
        "**Scrub attire + head cover** (all hair covered); mask not required"],
       ["**Sterile / aseptic zone**", "Restricted", "Operating room, scrub area, sterile store/set-up room",
        "Scrub attire + head cover + **mask**; sterile gown & gloves within the sterile field"],
       ["**Disposal / dirty zone**", "Outer/dirty corridor", "Dirty utility, disposal corridor, sluice",
        "Scrub attire; PPE for handling soiled items"]],
      widths=[3.4, 2.8, 8.0, 4.8])

h3("Cleaning schedule (\u201cfumigation\u201d is obsolete \u2014 know why)")
bullets([
    "**Between cases (\u2018turnover\u2019 clean):** all horizontal surfaces, the table, arm boards, positioning "
    "devices, lights, ESU, suction and the floor **within 1\u20131.5 m of the field** are wiped with a hospital-grade "
    "**germicide/disinfectant**. Your PDF's key teaching: **every case is cleaned identically** \u2014 the old "
    "practice of special \u2018ritual\u2019 cleaning only after \u2018dirty\u2019 cases is **obsolete**, because "
    "**all** cases are treated as potentially infectious (Standard Precautions).",
    "**Terminal (daily/end-of-day) clean:** entire room, ceiling, walls, all equipment wheels, cabinets, scrub sinks, "
    "corridors, storage; **wet vacuum** the floor.",
    "**Weekly/monthly:** ducts and filters checked, scrub sinks descaled, cabinets emptied, sterilizer maintenance.",
    "**Formaldehyde fumigation is no longer recommended** (Group 1 carcinogen, mucosal irritant, requires the room to "
    "be sealed and idle for 24\u201348 h, and does not penetrate crevices reliably). Modern alternatives if "
    "no-touch disinfection is needed: **vaporised hydrogen peroxide (VHP), hydrogen peroxide + silver, chlorine "
    "dioxide, UV-C irradiation** \u2014 always as an **adjunct after** physical cleaning, never instead of it.",
    "**Bacteriological surveillance:** settle plates / air sampling only for outbreak investigation or after new "
    "construction \u2014 **routine environmental culturing is not recommended**.",
    "**Blood spill:** absorb, then disinfect with **1 % sodium hypochlorite** (10 000 ppm available chlorine) for "
    "10 min; large spill = 1 % for 30 min. Wear gloves.",
])

h2("B3.  Antiseptics & Disinfectants \u2014 levels and agents")
table(["Level", "What it kills", "Agents", "Typical use"],
      [["**Sterilization**", "**All** microbes **including bacterial spores**",
        "Steam, EO, H\u2082O\u2082 plasma, glutaraldehyde 10 h, peracetic acid, radiation",
        "**Critical** items (\u00a7B12)"],
       ["**High-level disinfection (HLD)**", "All microbes **except large numbers of spores**",
        "Glutaraldehyde 2 % (20\u201345 min), **OPA 0.55 % (12 min)**, H\u2082O\u2082 7.5 % (30 min), peracetic acid, "
        "pasteurisation 75 \u00b0C/30 min", "**Semi-critical** \u2014 flexible endoscopes, laryngoscope blades, "
                                            "respiratory equipment"],
       ["**Intermediate-level**", "Vegetative bacteria, **M. tuberculosis**, most viruses & fungi; **not spores**",
        "Alcohols 60\u201390 %, hypochlorite, iodophors, phenolics",
        "Non-critical surfaces with visible contamination, thermometers"],
       ["**Low-level**", "Most vegetative bacteria, some fungi/viruses; **not TB, not spores**",
        "Quaternary ammonium compounds, low-conc. iodophors",
        "**Non-critical** \u2014 floors, walls, BP cuffs, furniture"]],
      widths=[3.0, 5.0, 6.0, 5.0])
table(["Antiseptic / disinfectant", "Mechanism", "Spectrum & key points"],
      [["**Alcohol** (ethyl 70 %, isopropyl 60\u201390 %)", "Protein denaturation",
        "Rapid, broadest immediate kill; **no residual activity**; **not sporicidal**; **inactivated by organic "
        "matter**; **flammable** \u2014 fire risk; not for open wounds/mucosa"],
       ["**Chlorhexidine gluconate (2\u20134 %)**", "Disrupts cytoplasmic membrane",
        "Excellent Gram-positive, good Gram-negative; **persistent residual activity**, not inactivated by blood; "
        "**ototoxic and corneal-toxic**; poor sporicidal/mycobacterial"],
       ["**Povidone-iodine 10 %** (= 1 % available iodine)", "Iodination/oxidation of proteins",
        "Broad, includes spores with prolonged contact; **needs 2\u20133 min contact & must dry to release free "
        "iodine**; **inactivated by blood/pus**; thyroid absorption in neonates; iodine allergy"],
       ["**Hexachlorophene 3 %**", "Bacteriostatic \u2013 enzyme inhibition",
        "Gram-positive only; cumulative effect; **neurotoxic in neonates (avoid)**"],
       ["**Sodium hypochlorite**", "Oxidising, releases free chlorine",
        "Broad including HIV/HBV & spores; **corrosive to metals**; inactivated by organic matter; "
        "**never mix with acids or ammonia (chlorine gas)**"],
       ["**Glutaraldehyde 2 %**", "Alkylation of protein/nucleic acid",
        "Sterilant/HLD; **does not corrode lenses or metal** \u2192 endoscopes; **irritant, sensitiser, asthma** "
        "\u2014 needs closed system + ventilation"],
       ["**Hydrogen peroxide**", "Free-radical oxidation", "3 % antiseptic, 7.5 % HLD, 58 % vapourised for plasma "
                                                           "sterilization; by-products water + oxygen (eco-friendly)"],
       ["**Ethylene oxide**", "Alkylation of DNA/RNA/protein",
        "Low-temperature sterilant; **toxic, mutagenic, carcinogenic, flammable/explosive**; requires aeration"],
       ["**Formaldehyde**", "Alkylation", "Sterilant/HLD; carcinogen; historically used for fumigation \u2014 "
                                          "now discouraged"],
       ["**Quaternary ammonium compounds**", "Membrane-active cationic surfactant",
        "Low-level only; inactivated by soap, hard water and gauze; environmental surfaces"]],
      widths=[4.0, 3.6, 11.4])

h2("B4.  Surgical Site Infection (SSI) \u2014 classification your PDF never gives")
h3("Wound classification (Altemeier / CDC) and expected infection rate")
table(["Class", "Definition", "Examples", "SSI risk"],
      [["**I \u2013 Clean**", "Uninfected, no inflammation; respiratory, GI, GU and genital tracts **not entered**; "
                              "closed primarily; elective",
        "Thyroidectomy, hernia repair, breast lump, joint replacement", "**< 2 %**"],
       ["**II \u2013 Clean-contaminated**", "Respiratory/GI/GU/genital tract entered under **controlled conditions** "
                                            "without unusual contamination",
        "Elective cholecystectomy, elective bowel resection, hysterectomy, appendicectomy (non-perforated)",
        "**\u2248 3\u201310 %**"],
       ["**III \u2013 Contaminated**", "Open, fresh, accidental wound; **major break in sterile technique**; gross "
                                       "spillage from GI tract; acute non-purulent inflammation",
        "Fresh traumatic wound, enterotomy with spillage, acute appendicitis with inflammation",
        "**\u2248 15\u201320 %**"],
       ["**IV \u2013 Dirty / infected**", "Old traumatic wound with retained devitalised tissue; existing clinical "
                                          "infection or **perforated viscus**",
        "Faecal peritonitis, drainage of abscess, perforated appendix", "**\u2248 30\u201340 %**"]],
      widths=[3.0, 6.2, 6.6, 2.0])
h3("CDC classification of SSI by depth")
bullets([
    "**Superficial incisional** \u2014 skin & subcutaneous tissue, within **30 days** of surgery.",
    "**Deep incisional** \u2014 fascia & muscle layers, within **30 days**, or within **90 days if an implant** is in "
    "place.",
    "**Organ/space** \u2014 any part opened or manipulated during surgery (e.g. intra-abdominal abscess, "
    "mediastinitis, meningitis), within 30/90 days.",
    "Overall, SSI is the **most common healthcare-associated infection in surgical patients** and the "
    "**2nd\u20133rd most common HAI overall**; commonest organism = **Staphylococcus aureus** (then coagulase-negative "
    "staphylococci, Enterococcus, E. coli).",
    "**Sources of contamination:** the **patient's own endogenous flora is the commonest source**; then the "
    "surgical team (shedding, breaks in technique), the environment/air, and instruments.",
    "**Risk factors** \u2014 patient: age, diabetes, obesity, smoking, malnutrition/hypoalbuminaemia, steroids, "
    "immunosuppression, remote infection, colonisation with S. aureus, ASA \u2265 3; operative: duration > 2 h "
    "(or > 75th percentile \u2014 the **T-time**), wound class, hypothermia, hypoxia, blood transfusion, foreign "
    "body/implant, drains, poor haemostasis, dead space, shaving with a razor.",
])

h2("B5.  Aseptic Technique, Scrub, Gowning & the Sterile Field \u2014 hard rules")
h3("Surgical hand antisepsis")
table(["Item", "Standard"],
      [["Agents", "**4 % chlorhexidine gluconate** or **7.5 % povidone-iodine** scrub, **or** an alcohol-based "
                  "surgical hand rub (which has better immediate kill and is now preferred by WHO)"],
       ["Duration", "**First scrub of the day 3\u20135 minutes** (per manufacturer); subsequent scrubs 3 min. "
                    "Older \u2018counted brush-stroke\u2019 method = 30 strokes to nails, 20 per surface"],
       ["Sequence", "**Fingertips/nails \u2192 fingers & interdigital spaces \u2192 hands \u2192 wrists \u2192 "
                    "forearms to 5 cm above the elbow**; always **cleanest (hand) held highest**; rinse from "
                    "fingertips downwards to elbow; **never return to a cleaned area**"],
       ["Before scrubbing", "Remove **all** rings, watches, bracelets; nails short, clean, **no polish, no gel, "
                            "no artificial nails** (AORN prohibits artificial nails \u2014 harbour Gram-negatives & "
                            "fungi); mask and cap adjusted; eyewear in place; **no scrubbing with cuts, abrasions or "
                            "dermatitis**"],
       ["After scrubbing", "Hands held **above the elbows and in front of the body, above waist and below shoulder "
                           "level**; dry with a sterile towel (one end per hand, blot from fingers to elbow, never "
                           "back)"],
       ["WHO \u201c5 Moments\u201d of hand hygiene", "**1** before touching a patient, **2** before a clean/aseptic "
                                                     "procedure, **3** after body-fluid exposure risk, "
                                                     "**4** after touching a patient, **5** after touching patient "
                                                     "surroundings"],
       ["Alcohol rub vs soap & water", "Alcohol rub is preferred routinely; **soap and water is mandatory when hands "
                                       "are visibly soiled and for Clostridioides difficile and spore-formers** "
                                       "(alcohol is not sporicidal)"],
       ["Double gloving", "Recommended for orthopaedic/high-risk cases and known blood-borne infection; "
                          "**perforation indicator** (coloured under-glove) shows breaches; change gloves every "
                          "90\u2013150 min and after every case"]],
      widths=[3.6, 15.4])
h3("Principles of the sterile field (AORN) \u2014 rote-learn these")
numbered([
    "Only **sterile items** are used within a sterile field.",
    "Sterile persons are gowned and gloved; the gown is sterile **in front from chest to the level of the sterile "
    "field, and the sleeves from 5 cm above the elbow to the cuff**. The **back is never considered sterile**.",
    "A **draped table is sterile only at table-top level**; anything hanging below the table edge is unsterile and is "
    "**never brought back up**.",
    "Sterile persons keep hands **at or above waist level and in front of the body**; they pass **back-to-back** or "
    "**face-to-face**.",
    "Unsterile persons stay **at least 30 cm (1 ft) away** from the sterile field, never reach over it, and always "
    "**face** it.",
    "The **edge of any sterile wrapper, package or basin is considered unsterile** \u2014 conventionally a "
    "**2.5 cm (1 inch)** margin.",
    "**Moisture causes contamination by capillary action (strike-through)**; a wet sterile field is a contaminated "
    "field.",
    "Sterile packages are opened **farthest corner first, nearest corner last**, holding the pack away from the body "
    "(exactly your PDF, pp. 35, 42) \u2014 so you never reach over the sterile contents.",
    "Sterile items must be kept **in view**; anything of **doubtful sterility is considered unsterile** "
    "(\u2018when in doubt, throw it out\u2019).",
    "Movement in and around the field is kept minimal; talking is minimised (droplet dispersal); the number of "
    "personnel is kept low.",
])

h2("B6.  Positioning \u2014 complications, angles and the nerve-injury map")
para("Your PDF describes each position beautifully but gives few complication specifics. **Positioning injuries are "
     "the second commonest source of anaesthesia-related claims after airway injury**, and peripheral nerve injury is "
     "the most examined item.")
table(["Position", "Nerve/structure at risk", "Mechanism & prevention", "Physiological effect"],
      [["**Supine**", "**Ulnar nerve (commonest nerve injury overall)**; brachial plexus; radial n. (armboard edge); "
                      "occiput/sacrum/heel pressure",
        "Arms **abducted \u2264 90\u00b0**, **palms up (supination)**, elbows padded; do not let the arm slip off the "
        "armboard; pad occiput, sacrum, heels; legs uncrossed",
        "\u2193 FRC; aorto-caval compression in late pregnancy (\u2192 left lateral tilt 15\u00b0)"],
       ["**Trendelenburg**", "Brachial plexus (shoulder braces), retinal & cerebral congestion, sliding",
        "Padded shoulder braces over the **acromio-clavicular** area, **never** on the soft tissue of the neck; "
        "limit to 30\u201345\u00b0 and to the **shortest time possible**",
        "**\u2191 venous return, \u2191 ICP, \u2191 IOP, \u2191 CVP; \u2193 FRC & compliance (cephalad diaphragm) "
        "\u2192 atelectasis; facial/laryngeal oedema; risk of endotracheal tube migration and regurgitation**"],
       ["**Reverse Trendelenburg / sitting (Fowler's)**", "Sciatic stretch, footdrop, venous pooling; "
                                                          "**venous air embolism** in sitting craniotomy",
        "Padded footboard, antiembolism stockings + **sequential compression device**, slow position change, "
        "secure head in traction device (Crutchfield/Mayfield)",
        "**\u2193 venous return \u2192 hypotension; \u2193 cerebral perfusion pressure (\u2248 2 mmHg per 2.5 cm of "
        "height above heart); air embolism risk**"],
       ["**Lithotomy**", "**Common peroneal nerve (commonest in lithotomy \u2192 foot drop)**; femoral, obturator, "
                         "saphenous, sciatic; **lower-limb compartment syndrome**",
        "Pad the **fibular head** away from the stirrup post; **raise and lower both legs together, slowly, by two "
        "people**; avoid extreme hip abduction/external rotation and hyperflexion; **limit time (risk rises sharply "
        "beyond 2\u20134 h)**; protect fingers when the foot of the table is raised",
        "**\u2191 preload initially; on lowering the legs \u2192 sudden \u2193 venous return \u2192 hypotension** "
        "(exactly why your PDF says lower both legs slowly and simultaneously); \u2193 FRC"],
       ["**Prone**", "Brachial plexus, ulnar, **eye \u2192 corneal abrasion and post-operative visual loss (POVL) "
                     "from ischaemic optic neuropathy or central retinal artery occlusion**; breasts, genitalia, "
                     "iliac crests; neck",
        "**Chest/bolster rolls from clavicle to iliac crest** to free the abdomen and allow diaphragm excursion; "
        "head on a padded horseshoe/donut with the **eyes free of pressure** and taped closed; keep head neutral, "
        "avoid direct pressure on the eyes; check ET tube after turning",
        "\u2191 abdominal & epidural venous pressure \u2192 bleeding; \u2193 cardiac index; airway displacement risk; "
        "improves oxygenation (basal recruitment)"],
       ["**Lateral (kidney / thoracotomy)**", "Dependent brachial plexus & axillary vessels, dependent common "
                                              "peroneal & ear, suprascapular n.",
        "**Axillary (chest) roll placed caudal to the axilla, not in it**; pillow between the legs; dependent leg "
        "flexed; head/neck in line with the spine; tape over benzoin-protected skin",
        "**V/Q mismatch \u2014 dependent lung is better perfused but less ventilated; kidney rest \u2193 venous "
        "return; \u2193 compliance**; post-op \u2018unnatural position\u2019 muscle pain"],
       ["**Jackknife / Kraske**", "Face, eyes, genitalia, knees; venous pooling in dependent legs",
        "Padding at hips, pillow under pelvis and shins, arms supported",
        "Marked \u2193 venous return and \u2193 FRC \u2014 poorly tolerated; ensure gradual positioning"],
       ["**Fracture table / traction**", "**Perineal post \u2192 pudendal nerve injury**, genital/perineal pressure "
                                         "necrosis; contralateral limb",
        "Well-padded perineal post, minimum effective traction, limit traction time, protect the well leg",
        "Blood loss, fat embolism risk in hip surgery"]],
      widths=[2.6, 3.8, 6.4, 5.2])
box("POSITIONING \u2014 GOLDEN RULES + MNEMONIC",
    ["**Never move an anaesthetised patient without the anaesthetist's permission and direction**; the "
     "**anaesthetist guards the head, airway and neck**, the **surgeon protects fractures/deformities**, the "
     "**circulator applies the padding and the safety strap**.",
     "Position is determined by the **surgical approach + patient's physical condition + anaesthetic technique + "
     "surgeon's preference**.",
     "The **safety strap across the mid-thighs, ~3 inches (7.5 cm) above the knees**, is the **single most important "
     "measure to prevent a fall from the table**; two fingers should slide under it.",
     "**PADDING mnemonic** \u2014 **P**ressure points padded, **A**lignment maintained, **D**ignity/exposure "
     "minimised, **D**evices in the room before the patient, **I**V lines & catheters protected, **N**erves free of "
     "stretch and compression, **G**entle simultaneous movement.",
     "Compartment syndrome after lithotomy presents as **pain out of proportion, tense swollen calf, paraesthesia**; "
     "it is a **surgical emergency \u2192 fasciotomy**. Always document the ability to move limbs and any "
     "paraesthesia on arrival in PACU (your PDF's Nursing Focus III note)."],
    fill=SH_MNE, border="6DA97A", title_color=C_H2)

h2("B7.  Safe Patient Transfer & Handling \u2014 numbers to quote")
bullets([
    "Minimum **two persons** to transfer (your PDF); **\u2265 4 persons or a mechanical lifter** for an "
    "anaesthetised, obese or unstable patient; **at least one person on each side** of the table.",
    "**Lock the wheels** of both the gurney and the table; the gurney is placed **against** the table; a "
    "**transfer/roller board (Davis roller), slide sheet, Hoyer pad or air-assisted lateral transfer mattress** "
    "reduces the effective load.",
    "**NIOSH revised lifting equation \u2192 maximum recommended load under ideal conditions = 51 lb (23 kg).** "
    "For **patient handling, the recommended maximum manual lift is 35 lb (16 kg)** \u2014 above this, use "
    "assistive devices. This is the number the PDF's \u2018lift team\u2019 paragraph is really about.",
    "Body mechanics: **bend the knees and hips, keep the back straight, hold the load close, feet apart for a wide "
    "base, pivot with the feet \u2014 never twist the spine, push rather than pull**, lift on a count in unison.",
    "OSHA (2002) requires facilities to reduce ergonomic injury; **back injury is the commonest occupational injury "
    "in health care**.",
    "During transfer protect: **airway/ET tube, IV lines, arterial and CVP lines, epidural catheter, Foley, drains** "
    "\u2014 and the **fingers** when the table sections are raised.",
])

h2("B8.  Electrosurgery \u2014 the physics, circuits and hazards")
h3("Fundamentals")
table(["Concept", "Detail"],
      [["Principle", "**High-frequency alternating current** is concentrated at a small active electrode; tissue "
                     "resistance converts it to heat \u2192 **desiccation, coagulation (fulguration) or cutting "
                     "(vaporisation)**. It is **not** \u2018cautery\u2019 (true cautery = a passively heated wire; "
                     "no current passes through the patient)"],
       ["Frequency", "**300 kHz \u2013 3 MHz (radiofrequency range)**. Frequencies **above ~100 kHz do not stimulate "
                     "nerve or muscle** (the **Faradic effect** is avoided) \u2014 this is why household 50/60 Hz "
                     "current electrocutes but RF current does not"],
       ["**Cut** mode", "**Continuous, low-voltage, high-current** waveform (100 % duty cycle) \u2192 rapid "
                        "vaporisation, minimal haemostasis"],
       ["**Coag** mode", "**Interrupted, high-voltage, low duty cycle (~6 %)** waveform \u2192 coagulum/eschar; "
                         "high voltage means **greater risk of capacitive coupling, insulation breakdown and "
                         "sparking**"],
       ["**Blend**", "Intermediate duty cycle \u2014 cutting with haemostasis"],
       ["**Monopolar** circuit", "Generator \u2192 **active electrode (pencil)** \u2192 patient \u2192 "
                                 "**dispersive/return electrode (\u2018ground\u2019 pad)** \u2192 generator. Current "
                                 "traverses the whole body; **the pad is essential**"],
       ["**Bipolar** circuit", "Both electrodes are the two tines of the **forceps**; current passes only through the "
                               "tissue between them. **No dispersive pad needed**, minimal lateral spread, "
                               "**safest with pacemakers**, usable in fluid, low power \u2014 used in "
                               "neurosurgery, ophthalmic and microsurgery"],
       ["Dispersive pad placement", "Over a **clean, dry, well-vascularised, hair-free muscle mass** (thigh, flank), "
                                    "**as close as practicable to the operative site**, **entire surface in uniform "
                                    "contact**, **away from ECG electrodes**, bony prominences, scar tissue, tattoos, "
                                    "metallic implants and prostheses, and away from pooled fluids"],
       ["Return-electrode monitoring (**REM / contact-quality monitoring**)",
        "A split (dual) pad in which the generator continuously measures impedance between the two halves and "
        "**deactivates itself if contact is inadequate** \u2014 has virtually eliminated pad-site burns. Modern "
        "generators are **isolated** rather than ground-referenced"],
       ["Active-electrode monitoring (**AEM**)", "Shielded laparoscopic instruments that detect and drain stray "
                                                 "current \u2014 prevents insulation-failure and capacitive-coupling "
                                                 "burns"]],
      widths=[4.4, 14.6])
h3("Hazards \u2014 the five classic mechanisms of electrosurgical injury")
numbered([
    "**Return-electrode (pad-site) burn** \u2014 poor/partial contact, hair, gel dried out, pad over bone or scar. "
    "Prevention: correct placement + REM.",
    "**Alternate-site burn** \u2014 the patient's skin contacts a **metal surface** (table edge, stirrup, arm-board "
    "screw, ECG lead) which becomes an unintended return path. Your PDF states this rule twice \u2014 "
    "**no skin-to-metal contact**.",
    "**Direct coupling** \u2014 the activated electrode touches another metal instrument (e.g. a laparoscope or "
    "retractor) which then burns adjacent tissue.",
    "**Capacitive coupling** \u2014 current is transferred through **intact insulation** to a nearby conductor by an "
    "electrostatic field; classic setting = a **metal instrument passing through a plastic (hybrid) trocar cannula** "
    "during laparoscopy, or coiled cords. Worse with **high-voltage coag mode and long, thin instruments**.",
    "**Insulation failure** \u2014 a pin-hole break in the shaft of a laparoscopic instrument delivers current "
    "**out of the surgeon's field of view** \u2192 delayed bowel perforation presenting 3\u201310 days "
    "post-operatively. Inspect insulation before every use.",
])
h3("Other electrosurgery examination points")
bullets([
    "**Repeated requests to increase the power setting** = poor pad contact, a break in a connection, a frayed cord "
    "or a failing generator \u2014 **stop, check connections and pad; do not simply turn it up**. Replace the unit, "
    "label it, and **document the serial numbers and the reason**.",
    "**Keep the pencil in its holster** when not in use (direct-contact burns and drape fires). Clean the eschar off "
    "the tip with a **non-abrasive scratch pad**, not a scalpel.",
    "**Pacemaker / ICD:** use **bipolar** where possible; if monopolar is essential use short bursts at the lowest "
    "power, keep the current path away from the device (pad so that the pad\u2013site line does not cross the "
    "generator), have **magnet, external defibrillator/pacer and the crash cart immediately available**, monitor ECG "
    "and pulse continuously, and consider re-programming/deactivating anti-tachycardia therapy pre-operatively. "
    "**Interrogate the device after surgery.** Also avoid MRI and diathermy near IEDs.",
    "**Do not use ESU or laser in an oxygen-enriched field** (mouth, oropharynx, tracheostomy, or under drapes with "
    "high-flow O\u2082) \u2014 pause the oxygen or reduce FiO\u2082 and wait for washout.",
    "**Advanced energy devices** \u2014 **Ultrasonic (Harmonic) scalpel:** mechanical vibration at "
    "**55.5 kHz**, protein denaturation by cavitation at a **lower temperature (50\u2013100 \u00b0C)**, no electrical "
    "current through the patient, minimal smoke, seals vessels up to 5 mm. **Bipolar vessel sealing (LigaSure):** "
    "pressure + energy seals vessels up to **7 mm** with feedback-controlled cut-off. **Argon-enhanced coagulation:** "
    "argon gas jet carries current, gives a thin uniform eschar (risk of gas embolism). **Radiofrequency ablation, "
    "cryosurgery (\u2013 60 to \u2013 190 \u00b0C), microwave** \u2014 alternatives.",
    "**Surgical smoke (plume):** contains **> 150 chemicals** including benzene, formaldehyde, hydrogen cyanide, "
    "acrolein, plus viable cells, blood fragments and viral DNA (**HPV transmission to surgeons is documented**). "
    "Ablating **1 g of tissue \u2248 the mutagenic exposure of 3\u20136 unfiltered cigarettes**. Control: a "
    "**smoke evacuator with an ULPA filter (99.999 % at 0.1 \u00b5m) held within 2.5\u20135 cm (1\u20132 inches)** of "
    "the plume, plus a **high-filtration surgical mask (or N95)** \u2014 a standard mask is *not* adequate protection.",
])

h2("B9.  Tourniquet, Laser and Powered Equipment \u2014 quick numbers")
table(["Device", "Key parameters & safety"],
      [["**Pneumatic tourniquet**",
        "Cuff ends overlap **3\u20136 inches**; **never at or near the elbow (radial nerve) or the mid-humerus/"
        "fibular neck**. Pressure: **upper limb 250\u2013300 mmHg** (or systolic + 100), **lower limb 300\u2013350 "
        "mmHg** (or 2 \u00d7 systolic); less in children and thin adults. Time: **\u2264 60 min upper limb, "
        "\u2264 90\u2013120 min lower limb**; if longer, **deflate for 10\u201315 min** then re-inflate. Exsanguinate "
        "with an Esmarch bandage from distal to proximal **before** inflation. **Record site, pressure, and inflation "
        "and deflation times** on the perioperative and anaesthesia records. Complications: nerve palsy, "
        "**post-tourniquet syndrome**, rhabdomyolysis, compartment syndrome, **tourniquet pain and hypertension**, "
        "and on release \u2014 **\u2193 BP, transient metabolic acidosis, \u2191 K\u207a, \u2191 EtCO\u2082 and "
        "risk of PE**. Contraindicated in sickle-cell disease, severe PVD, DVT, crush injury"],
       ["**Surgical laser**",
        "**L**ight **A**mplification by **S**timulated **E**mission of **R**adiation. Classes **1\u20134**; all "
        "surgical lasers are **Class 3B or 4** \u2192 mandatory controls. **Wavelength-specific eye protection for "
        "everyone including the patient** (CO\u2082 10 600 nm \u2014 clear glass/plastic suffices and cornea absorbs; "
        "**Nd:YAG 1064 nm** \u2014 penetrates deeply, needs green-tinted goggles; Argon 488/514 nm \u2014 amber; "
        "KTP 532 nm). Controls: **warning signs on all doors, windows covered, restricted entry, key kept by the "
        "designated laser safety officer, \u2018standby\u2019 mode when not firing, foot pedal only to the surgeon, "
        "matte/non-reflective instruments, wet towels and drapes around the site, saline available, special "
        "fire-extinguisher in room, smoke (plume) evacuator, laser-safe cuffed ET tube with saline-filled cuff and "
        "no nitrous oxide, and never use in an oxygen-enriched field.** Hazards: eye (retinal/corneal) injury, skin "
        "burn, **fire and airway fire**, plume, electrical"],
       ["**Powered saws/drills**",
        "Engineering clearance before first use; test before use; ensure attachments are secure and sharp; jigs and "
        "templates size-matched; **let the rotating part come to a complete standstill before setting it down**; "
        "deactivate the power source when not in use for an interval; irrigate to reduce thermal necrosis; "
        "**face shield mandatory** (bone chips, aerosol)"],
       ["**Pulsed lavage**",
        "Single-use; **face shield for all scrubbed persons**; do not submerge the handle; not near an "
        "oxygen-enriched area or flammable anaesthetics; dispose of unit and batteries as **biohazard**"]],
      widths=[3.4, 15.6])

h2("B10.  Counting Procedure \u2014 Retained Surgical Items (RSI)")
para("The PDF gives the 9 counting rules perfectly \u2014 learn them from there. What it lacks is the epidemiology "
     "and terminology.")
bullets([
    "**RSI / RFO (retained foreign object)** is classified by both the JC and CMS as a **\u2018never event\u2019 / "
    "sentinel event**.",
    "**Commonest retained item = a surgical sponge/gauze** (\u2248 50\u201370 %), then instruments (retractors, "
    "malleables) and needles. Commonest site = **abdomen/pelvis** (then vagina and thorax).",
    "A retained gauze surrounded by a foreign-body granuloma is called a **gossypiboma / textiloma / cottonoid "
    "granuloma** \u2014 presents weeks to years later with a mass, pain, sinus, abscess, obstruction or fistula, and "
    "is a classic **res ipsa loquitur** claim.",
    "**Risk factors:** emergency surgery, **unexpected intra-operative change in procedure**, high BMI, more than one "
    "surgical team/procedure, multiple staff changes, incorrect or omitted counts, and long operations.",
    "**Counts are performed:** (1) before the procedure (baseline), (2) when **new items are added**, (3) **before "
    "closure of a cavity within a cavity** (uterus, bladder, bowel), (4) when **wound (fascia) closure begins**, "
    "(5) at **skin closure / end of procedure**, and (6) at any **change of scrub or circulating personnel**.",
    "**Adjunct technologies:** all sponges used in a wound must be **radio-opaque (X-ray detectable)**; "
    "**bar-coded sponge counting systems**; **radiofrequency (RF) tag detection wands/mats** that scan the patient "
    "before closure; intra-operative radiography.",
    "**If the count is incorrect:** repeat the count \u2192 **inform the surgeon immediately** \u2192 surgeon "
    "explores the wound \u2192 if still missing, **intra-operative X-ray before the patient leaves the OR** \u2192 "
    "notify the OR supervisor \u2192 **document the discrepancy, the actions and the outcome, and file an incident "
    "report**. **Nothing (linen, trash, instruments, sponges) leaves the room until the final count is correct.**",
    "**X-ray-detectable sponges are never used as dressings**, and dressing sponges are never used in the wound.",
])

h2("B11.  Operative / Perioperative Records \u2014 beyond the forms")
h3("Specimen handling \u2014 a favourite MCQ that the PDF omits")
table(["Specimen type", "Handling"],
      [["Routine histopathology", "**10 % neutral buffered formalin**, volume **\u2248 10\u00d7 the specimen**; "
                                  "container labelled with patient name, ID, date, site/side and specimen type; "
                                  "requisition matched; **label read aloud during \u2018sign out\u2019**"],
       ["**Frozen section**", "**Fresh, in saline-moistened gauze \u2014 NEVER in formalin** (formalin makes further "
                              "study and immunohistochemistry impossible); send immediately"],
       ["Culture / microbiology", "**Sterile container, no fixative**; anaerobic swab or syringe for anaerobes; "
                                  "transport immediately; **never in formalin**"],
       ["Cytology / fluid", "Fresh or with cytology fixative as per laboratory"],
       ["Stones/calculi", "Dry container, **no fixative**"],
       ["Amputated limb / tissue for the patient", "Per state law, hospital policy and cultural/religious wishes; "
                                                   "documented"],
       ["Lost specimen", "Treated as a **serious incident** \u2014 irreplaceable; incident report mandatory"],
       ["Implants", "Record **type, manufacturer, size, serial/lot number** in the intra-operative record and the "
                    "**implant log** (traceability for recalls)"]],
      widths=[4.2, 14.8])
h3("What makes the record legally sound")
bullets([
    "**JC-mandated minimum content of the operative record:** surgeon's name (and assistants), **pre-operative "
    "diagnosis**, procedure performed, **post-operative diagnosis**, findings, **specimens removed (type and "
    "number)**, estimated blood loss, and the names of all team members providing intra-operative care.",
    "Also document: the **time out**, counts and their outcome, skin condition **before and after** prep and the "
    "prep solution, position and positioning aids, **ESU pad site, serial number and settings**, tourniquet data, "
    "drains/catheters/packing, medications and irrigations (name, strength, dose, route, time, who gave them), "
    "implants, intake/output, transfusion checks, and any **unusual event with the actions taken**.",
    "**Rules of documentation:** contemporaneous, legible, factual and objective (no opinion or blame), signed and "
    "timed; **corrections by a single line through with the author's initials, date and time \u2014 never erase, "
    "overwrite or use correction fluid**; no blank spaces left; late entries labelled as such.",
    "**Incident (occurrence) report** \u2014 an internal risk-management document. Record only the facts; do **not** "
    "photocopy it into the chart or reference it in the clinical notes in many jurisdictions (the PDF's practice of "
    "filing it in the chart follows its own institutional convention \u2014 answer as per your textbook if asked). "
    "File one for: incorrect count, medication error, burn, fall, retained item, equipment failure, needle-stick, "
    "specimen loss, unplanned re-operation and **death in the OR**.",
    "**Retention:** medical records are typically retained for **3 years** for outpatient and **10 years** for "
    "in-patient/medico-legal records in Indian practice (MCI/NMC guidance; longer for minors and medico-legal "
    "cases) \u2014 follow institutional policy.",
    "**Confidentiality** \u2014 patient information is disclosed only on a need-to-know basis; consent needed for "
    "photography, and images in which the patient is identifiable.",
])



h2("B12.  Sterilization & Disinfection \u2014 the full science (highest-yield section)")
h3("Definitions \u2014 get these exactly right")
table(["Term", "Definition"],
      [["**Sterilization**", "Complete **destruction or removal of all forms of microbial life, including bacterial "
                             "and fungal spores**. It is an **absolute** term \u2014 an item is either sterile or it "
                             "is not"],
       ["**Disinfection**", "Destruction of **pathogenic** organisms but **not necessarily spores**; used on "
                            "**inanimate** objects"],
       ["**Antisepsis**", "Destruction/inhibition of organisms on **living tissue** (skin, mucosa)"],
       ["**Sanitisation**", "Reduction of microbial load to a level judged safe by public-health standards"],
       ["**Decontamination**", "Physical/chemical removal or inactivation of contaminants so that an item is safe to "
                               "handle \u2014 **always the FIRST step, before disinfection or sterilization**"],
       ["**Cleaning**", "Physical removal of soil and organic matter (**cannot sterilize a dirty instrument** \u2014 "
                        "organic matter blocks the sterilant)"],
       ["**Asepsis**", "Absence of pathogenic organisms; **medical asepsis** = clean technique (reduce numbers), "
                       "**surgical asepsis** = sterile technique (absence of all organisms)"],
       ["**Bactericidal / -static**", "Kills / merely inhibits multiplication"],
       ["**Sporicidal**", "Kills spores \u2014 the property that distinguishes a **sterilant** from a disinfectant"],
       ["**SAL (Sterility Assurance Level)**", "**10\u207b\u2076** \u2014 probability of \u2264 one viable organism "
                                               "in 10\u2076 sterilized items. The internationally accepted definition "
                                               "of \u2018sterile\u2019"],
       ["**D-value**", "Time (or radiation dose) needed to reduce the microbial population by **1 log\u2081\u2080 "
                       "(90 %)** at a stated condition"],
       ["**Bioburden**", "Number and resistance of microbes on an item before sterilization"]],
      widths=[4.0, 15.0])
h3("Spaulding classification \u2014 how the level of processing is decided")
table(["Category", "Definition", "Examples", "Required processing"],
      [["**Critical**", "Enter **sterile tissue or the vascular system**",
        "Surgical instruments, implants, needles, cardiac & urinary catheters, arthroscopes, laparoscopes",
        "**Sterilization**"],
       ["**Semi-critical**", "Contact **intact mucous membranes or non-intact skin**",
        "Flexible endoscopes (GI, bronchoscope), laryngoscope blades, ET tubes, respiratory & anaesthesia equipment, "
        "vaginal probes, thermometers",
        "**High-level disinfection** (sterilization preferred where feasible)"],
       ["**Non-critical**", "Contact **intact skin only**",
        "BP cuffs, stethoscopes, ECG electrodes, bedpans, OR table, floors, walls, furniture",
        "**Low- or intermediate-level disinfection / cleaning**"]],
      widths=[2.8, 4.6, 7.4, 4.2])
h3("Resistance of organisms to sterilization / disinfection (most \u2192 least resistant)")
para("**Prions > Bacterial spores (C. difficile, Bacillus, Geobacillus) > Coccidia (Cryptosporidium) > "
     "Mycobacteria > Non-lipid/small viruses (polio, norovirus, HAV) > Fungi > Vegetative bacteria > "
     "Lipid/medium-sized viruses (HIV, HBV, herpes, influenza)** \u2014 hence a sporicidal agent covers everything "
     "below spores, and HIV/HBV are among the **easiest** to kill (a common MCQ inversion).")
box("PRIONS \u2014 the exception to every rule (CJD)",
    ["Prions resist routine autoclaving, EO, formaldehyde, alcohol and plasma sterilization.",
     "Recommended: **single-use instruments where possible**, otherwise **1 N NaOH (or 1 % NaOCl) for 1 h + steam "
     "autoclaving at 134 \u00b0C for 18 minutes** (prevacuum), or 121 \u00b0C for 60 min after NaOH immersion. "
     "**Never flash-sterilize; do not reuse in neurosurgery.**"],
    fill=SH_CLIN)

h3("1.  Moist heat / saturated steam under pressure (autoclave) \u2014 the parameters")
para("**Mechanism: denaturation and coagulation of proteins.** Moist heat is far more efficient than dry heat "
     "because **latent heat of vaporisation** is released when steam condenses on the item. It is the "
     "**method of choice \u2014 fastest, cheapest, most reliable, non-toxic** \u2014 for everything that can "
     "tolerate heat and moisture.")
table(["Cycle / type", "Temperature", "Pressure", "Hold time", "Use"],
      [["**Gravity displacement (downward)**", "**121 \u00b0C (250 \u00b0F)**", "**15 psi (\u2248 1 atm / 103 kPa)**",
        "**15\u201320 min** (wrapped packs 30 min)", "Standard general load, glassware, media, rubber"],
       ["**Gravity displacement \u2014 high temp**", "**134 \u00b0C (273 \u00b0F)**", "**30 psi (\u2248 2 atm)**",
        "**3\u20133.5 min**", "Rapid cycles"],
       ["**Prevacuum (high-vacuum, pulsed)**", "**132\u2013135 \u00b0C**", "\u2248 30 psi", "**3\u20134 min** "
        "(+ drying 20\u201330 min)", "**Wrapped instruments, trays, textiles, lumened items** \u2014 air is actively "
        "evacuated so steam penetrates"],
       ["**Flash / IUSS** (Immediate-Use Steam Sterilization)", "**132 \u00b0C (270 \u00b0F)**", "\u2248 27\u201330 psi",
        "**3 min unwrapped** metal only; **10 min** for porous, lumened or complex items",
        "**Only for urgently needed, unwrapped single items** \u2014 e.g. a dropped instrument. **Not** for implants, "
        "not for convenience, not for whole sets"],
       ["**Prions (CJD)**", "134 \u00b0C", "\u2014", "**18 min** prevacuum after NaOH", "Neurosurgical instruments"]],
      widths=[3.8, 2.8, 2.6, 4.4, 5.4])
bullets([
    "**Air is the enemy of the autoclave** \u2014 trapped air lowers the temperature at a given pressure and prevents "
    "steam contact. Hence: **do not overload, do not wrap tightly, place trays with mesh bottoms, position packs on "
    "edge with space between them, and never place a bowl/basin so that it traps air.**",
    "**Bowie-Dick test** = a **daily test of air removal / steam penetration in a PREVACUUM sterilizer**, run in an "
    "**empty chamber** at 134 \u00b0C for 3.5 min; a uniform colour change on the test sheet means no air pocket. "
    "It is a **Class 2 chemical indicator** and it does **not** prove sterility.",
    "**Vacuum leak test** weekly; **biological indicator at least weekly (daily preferred) and with EVERY implant "
    "load** \u2014 implant loads should be quarantined until the BI result is available.",
    "**Cannot be autoclaved:** sharp cutting edges (blunted), oils, greases, powders, petroleum-based ointments, "
    "anhydrous materials, plastics with low melting point, endoscopes with lenses, and heat-labile electronics.",
    "Drying is part of the cycle \u2014 a **wet pack is a contaminated pack (strike-through)**; do not touch or move "
    "packs until cool and dry.",
])
h3("2.  Dry heat (hot air oven / incineration / flaming) \u2014 absent from your PDF")
para("**Mechanism: oxidation and protein denaturation \u2014 needs much higher temperature and longer time than "
     "moist heat.**")
table(["Temperature", "Holding time"],
      [["**160 \u00b0C**", "**2 hours** (the most commonly asked pair)"],
       ["**170 \u00b0C**", "**1 hour**"],
       ["**180 \u00b0C**", "**30 minutes**"],
       ["**150 \u00b0C**", "2.5 hours"]],
      widths=[4.0, 6.0])
bullets([
    "**Uses:** glassware, metal instruments (**sharps \u2014 does not blunt or corrode**), oils, greases, powders, "
    "anhydrous fats, glass syringes, all-glass items.",
    "**Not for:** rubber, plastics, fabrics, most surgical linen.",
    "Load loosely to allow air circulation; **allow to cool before opening** (glass cracks, and hot air rushing in "
    "draws contaminated air).",
    "**Other dry-heat methods:** **incineration** (final disposal of contaminated waste and anatomical parts, "
    "\u2265 1100\u20131200 \u00b0C), **red-heat flaming** (inoculation loops), **glass-bead steriliser** "
    "(250 \u00b0C, 20\u201340 s \u2014 dental).",
    "**Biological indicator: Bacillus atrophaeus** (formerly B. subtilis var. niger).",
])
h3("3.  Ethylene oxide (EO) gas chemical sterilization")
table(["Parameter", "Value"],
      [["Mechanism", "**Alkylation** of DNA, RNA and proteins"],
       ["The 4 critical variables", "**Gas concentration (450\u20131200 mg/L) \u2022 Temperature (37\u201363 \u00b0C) "
                                    "\u2022 Relative humidity (40\u201380 %) \u2022 Exposure time (1\u20136 h; "
                                    "your PDF quotes 3\u20137 h cycles)** \u2014 all four must be met"],
       ["Aeration (\u2018degassing\u2019)", "**Mechanical aerator: 8\u201312 h at 50\u201360 \u00b0C** "
                                            "(122\u2013140 \u00b0F), air changed \u2265 4 times/h. "
                                            "**Ambient room: 7 days at 18\u201322 \u00b0C** \u2014 mandatory, because "
                                            "residual EO causes chemical burns and haemolysis"],
       ["Uses", "**Heat- and moisture-sensitive items** \u2014 flexible endoscopes, cardiac catheters, plastics, "
                "electrical/electronic equipment, prosthetic valves, pacemakers, respiratory equipment, "
                "delicate optics"],
       ["Rule", "**Anything that can withstand steam should NOT be gas-sterilized**"],
       ["Hazards", "**Toxic, irritant, mutagenic, teratogenic, carcinogenic (leukaemia, lymphoma), flammable and "
                   "explosive** (hence mixed with CO\u2082 or HCFC diluents); OSHA **PEL 1 ppm (8-h TWA)**, "
                   "excursion 5 ppm/15 min, action level 0.5 ppm"],
       ["Biological indicator", "**Bacillus atrophaeus**"],
       ["Limitation", "**Cannot penetrate anything that absorbs it**; must be scrupulously dry and clean; long "
                      "turnaround"]],
      widths=[3.4, 15.6])
h3("4.  Low-temperature hydrogen peroxide gas plasma (STERRAD)")
bullets([
    "**58 % hydrogen peroxide (1.8 mL) is vaporised under vacuum, then excited by radiofrequency energy into a "
    "plasma** of free radicals and charged particles \u2192 kills bacteria, fungi, viruses **and spores**.",
    "**Temperature 45\u201350 \u00b0C (\u2248 122 \u00b0F); cycle 28\u201375 min (\u2248 1 h)**; "
    "**by-products = water vapour + oxygen only \u2192 NO aeration or cooling needed, no toxic residue.**",
    "Ideal for: metals and non-metals, **fibre-optic cables, endoscopes, microsurgical and electrical equipment, "
    "glass, ceramics**.",
    "**Cannot be used for cellulose-containing items \u2014 paper, linen, cotton, towels, gauze, dressings \u2014 "
    "or for liquids, powders, or long narrow lumens**, because cellulose absorbs the sterilant. High capital cost.",
    "**Biological indicator: Geobacillus stearothermophilus.**",
])
h3("5.  Liquid chemical sterilization / high-level disinfection")
table(["Agent", "Sterilization", "High-level disinfection", "Notes"],
      [["**Glutaraldehyde 2 % (activated, alkaline \u2014 Cidex)**", "**10 hours** immersion",
        "**20\u201345 min at 20\u201325 \u00b0C** (your textbook quotes 10 min \u2014 quote your textbook if the "
        "question is from it, but CDC says 20\u201345 min)",
        "Non-corrosive to lenses/metal \u2192 endoscopes. **Rinse thoroughly in sterile distilled water.** "
        "Activated solution loses potency \u2192 **test with a test strip before each use**, reuse life 14\u201328 d. "
        "Wear **polyethylene or double gloves + eye protection**; irritant, **asthma/dermatitis**, "
        "**ACGIH ceiling 0.05 ppm**"],
       ["**Ortho-phthalaldehyde (OPA) 0.55 %**", "\u2014", "**12 min at 20 \u00b0C**",
        "Faster, less irritant than glutaraldehyde; stains skin/clothing grey"],
       ["**Peracetic acid 0.2 % (Steris System 1 / Steris 20)**", "**12 min at 50\u201356 \u00b0C**", "\u2014",
        "Single-use, automatic dilution, effective **despite organic soil**, low surface tension \u2192 reaches "
        "lumens; **decomposes to acetic acid + water + oxygen** (environmentally safe); corrosive when concentrated"],
       ["**Hydrogen peroxide 7.5 %**", "6 hours", "**30 min**", "Non-toxic breakdown products"],
       ["**Formaldehyde (8 % + 70 % alcohol)**", "\u2014", "High level", "Carcinogen, pungent \u2014 largely abandoned"]],
      widths=[4.2, 2.8, 4.0, 8.0])
h3("6.  Radiation and filtration sterilization")
table(["Method", "Detail"],
      [["**Ionising (\u03b3) radiation**", "**Cobalt-60** source; dose **25 kGy (2.5 Mrad)**. **Cold, industrial-scale "
                                           "method for pre-packed single-use disposables** \u2014 syringes, needles, "
                                           "gloves, catheters, sutures, blades, prosthetic heart valves, bone grafts. "
                                           "High penetration, no residue. **Not used in hospitals** (cost, shielding). "
                                           "BI = **Bacillus pumilus**"],
       ["**Non-ionising (UV-C, 254 nm)**", "**Surface and air disinfection only \u2014 poor penetration** (does not "
                                           "pass through glass, paper, water films). Used in laminar-flow hoods, "
                                           "operating-room air, biosafety cabinets. **Hazard: keratoconjunctivitis "
                                           "and erythema** \u2014 never expose skin or eyes; lamps must be off when "
                                           "the room is occupied. It is a **disinfectant, not a steriliser**"],
       ["**Filtration**", "**0.22 \u00b5m (bacteria-retaining) membrane filters** for **heat-labile liquids** "
                          "\u2014 serum, antibiotic and vaccine solutions, IV fluids, culture media. **Removes but "
                          "does not kill**; **does not retain viruses or mycoplasma** (needs 0.1 \u00b5m). "
                          "**HEPA filtration** of air = the OR/laminar-flow application (\u00a7B2)"]],
      widths=[3.6, 15.4])
h3("7.  Sterilization monitoring \u2014 the three-indicator system")
table(["Type", "What it is", "What it proves", "Examples"],
      [["**Physical / mechanical**", "Gauges, digital printout, chart recorder, cycle log, time\u2013temperature\u2013"
                                     "pressure graph",
        "That the **machine** achieved the parameters; **read and signed for every cycle**",
        "Sterilizer printout, thermocouple"],
       ["**Chemical (ISO 11140 classes 1\u20136)**", "Heat/chemical-sensitive dyes on tape, strips or cards",
        "**Exposure** to the process \u2014 **NOT sterility**",
        "**Class 1** process indicator (autoclave tape on the outside of the pack \u2014 only distinguishes processed "
        "from unprocessed) \u2022 **Class 2** specific-test indicator (**Bowie-Dick**) \u2022 **Class 3** "
        "single-variable \u2022 **Class 4** multi-variable \u2022 **Class 5 integrating indicator** (responds to all "
        "critical variables, correlates with the BI) \u2022 **Class 6 emulating indicator** (cycle-verification, "
        "most stringent). An **internal** chemical indicator goes in the **centre of every pack**"],
       ["**Biological (BI) \u2014 the GOLD STANDARD**", "A calibrated preparation of **highly resistant bacterial "
                                                        "spores** processed with the load and then incubated",
        "**Actual sterilizing efficacy \u2014 the only indicator that proves the sterilization process killed "
        "spores**",
        "**Steam & H\u2082O\u2082 plasma \u2192 Geobacillus stearothermophilus** (thermophile, incubate 55\u201360 "
        "\u00b0C) \u2022 **Dry heat & ethylene oxide \u2192 Bacillus atrophaeus** (incubate 35\u201337 \u00b0C) \u2022 "
        "**\u03b3-radiation \u2192 Bacillus pumilus**. Frequency: **at least weekly, daily preferred, and with every "
        "implant load**; rapid-readout BIs give results in 1\u20133 h"]],
      widths=[3.6, 4.4, 4.6, 6.4])
box("PACKAGING, STORAGE & SHELF-LIFE \u2014 commonly asked",
    ["Packaging must **permit sterilant penetration, provide a microbial barrier, allow aseptic opening and resist "
     "tearing/moisture**: muslin (double thickness, 140 thread count), non-woven polypropylene, crepe paper, "
     "paper\u2013plastic peel pouches (\u2018see-through\u2019), rigid containers with filters. **Never use closed "
     "metal containers or aluminium foil for steam.**",
     "**Shelf-life is now EVENT-RELATED, not TIME-RELATED**: a pack remains sterile **until an event compromises it** "
     "\u2014 the wrapper becomes wet, torn, punctured or soiled, the seal is broken, it is dropped, or it is handled "
     "excessively. (The old \u201csterile for 30 days\u201d rule is obsolete.)",
     "Store in a **clean, dry, dust-free, closed cupboard**, **20\u201323 \u00b0C, humidity 30\u201360 %**, "
     "**at least 20\u201325 cm (8\u201310 in) above the floor, 45 cm (18 in) below the ceiling/sprinkler and "
     "5 cm (2 in) from outside walls**; stock rotated **FIFO (first-in, first-out)**.",
     "**Sequence of instrument processing:** point-of-use pre-cleaning (enzymatic foam spray in the OR) \u2192 "
     "transport in a closed container \u2192 **decontamination / washer-disinfector** \u2192 **ultrasonic cleaner** "
     "(cavitation removes soil from box locks, serrations, lumens) \u2192 rinse and dry \u2192 inspect, lubricate and "
     "assemble (**box locks and hinged instruments OPEN/unlocked**) \u2192 pack with internal CI \u2192 sterilize "
     "\u2192 cool/dry \u2192 store. **Never sterilize a soiled instrument.**"],
    fill=SH_HY, border="D98BA0", title_color=C_ANS)
box("REUSE OF \u2018SINGLE-USE\u2019 DEVICES \u2014 the PDF's controversy, summarised",
    ["Reprocessing a device labelled **single-use** transfers the **manufacturer's legal liability to the "
     "reprocessor** (the hospital or a third-party reprocessor), who must then meet the regulator's validation "
     "requirements (FDA rule, August 2000).",
     "Risks: **retained bioburden in lumens, endotoxin/pyrogen reaction, material degradation and functional failure, "
     "residual sterilant toxicity, and prion transmission.**",
     "AORN's position: **patient safety is the paramount consideration**; opened-but-unused items are less "
     "problematic than used items. In India, reuse must comply with hospital policy and the "
     "Medical Devices Rules \u2014 always answer that reuse is **acceptable only if validated and permitted by the "
     "manufacturer/regulator**."],
    fill=SH_CLIN)

h2("B13.  Emergencies and Disasters \u2014 protocols, drugs and doses")
h3("Cardiopulmonary arrest in the OR")
table(["Item", "Current (AHA/ILCOR) numbers"],
      [["Sequence", "**C\u2013A\u2013B** (Compressions \u2013 Airway \u2013 Breathing); high-quality CPR is the "
                    "priority"],
       ["Compression rate", "**100\u2013120 per minute**"],
       ["Compression depth", "**\u2265 5 cm (2 in) and \u2264 6 cm** in adults; **\u2153 of chest depth** in "
                             "children/infants (\u2248 5 cm / 4 cm)"],
       ["Compression : ventilation", "**30 : 2** (single/two rescuer adult, unintubated); **continuous compressions "
                                     "with 1 breath every 6 s** once an advanced airway is in place; "
                                     "**15 : 2** for two-rescuer paediatric"],
       ["Chest-compression fraction & interruptions", "Aim **> 60\u201380 %**; interruptions **< 10 s**; "
                                                      "**full recoil** between compressions; change compressor "
                                                      "**every 2 min**"],
       ["Defibrillation", "**Biphasic 120\u2013200 J** (per manufacturer) or **monophasic 360 J**; "
                          "**shockable rhythms = VF and pulseless VT**; "
                          "**PEA and asystole are NOT shockable**"],
       ["Drugs", "**Adrenaline (epinephrine) 1 mg IV/IO every 3\u20135 min** \u2022 **Amiodarone 300 mg** bolus then "
                 "150 mg (or lidocaine 1\u20131.5 mg/kg) for refractory VF/pVT \u2022 treat reversible causes"],
       ["Reversible causes", "**4 H's** \u2013 Hypoxia, Hypovolaemia, Hypo/Hyperkalaemia & metabolic, "
                             "Hypothermia; **4 T's** \u2013 Tension pneumothorax, Tamponade, Toxins, "
                             "Thrombosis (coronary/pulmonary)"],
       ["Monitoring quality", "**Waveform capnography \u2014 EtCO\u2082 < 10 mmHg indicates poor-quality CPR; "
                              "a sudden rise suggests ROSC**"],
       ["Paediatric", "**PALS** competency required for staff caring for children (your PDF); "
                      "defibrillation 2 J/kg, then 4 J/kg; adrenaline 0.01 mg/kg"],
       ["Team logistics in the OR", "**Emergency signal (flashing red light)** operable from each OR \u2022 crash cart "
                                    "brought in by nursing staff \u2022 anaesthetist and surgeon lead \u2022 "
                                    "**scribe documents every drug, dose, route, time and intervention** \u2014 "
                                    "undocumented care is legally non-existent"]],
      widths=[4.0, 15.0])
h3("Crash cart \u2014 minimum contents (asked as \u2018which is NOT present\u2019)")
para("**Defibrillator with paddles/pads + ECG monitor** \u2022 **Airway drawer**: oral/nasal airways, laryngoscope "
     "with blades and spare bulbs/batteries, ET tubes of all sizes with stylet, LMAs, Magill forceps, suction "
     "catheters and a working suction unit, **bag-valve-mask (Ambu) with reservoir**, oxygen cylinder with flowmeter "
     "\u2022 **Drugs (in date)**: adrenaline, atropine, amiodarone, lidocaine, adenosine, sodium bicarbonate, "
     "calcium chloride/gluconate, magnesium sulphate, dopamine/noradrenaline, dextrose 25/50 %, naloxone, "
     "flumazenil, hydrocortisone, chlorpheniramine, furosemide, GTN, midazolam \u2022 IV cannulae, "
     "syringes/needles, IV fluids and giving sets, intraosseous needle \u2022 CVP set, chest-drain set, "
     "cricothyrotomy set \u2022 **backboard, gloves, sharps container, documentation sheet** \u2022 "
     "**checked and re-sealed after every use and at least daily/each shift, with a signed checklist.**")
h3("Malignant hyperthermia (MH) \u2014 full protocol")
table(["Item", "Detail"],
      [["Genetics", "**Autosomal dominant**; mutation of the **ryanodine receptor RYR1 gene on chromosome 19q**; "
                    "uncontrolled Ca\u00b2\u207a release from the sarcoplasmic reticulum \u2192 sustained muscle "
                    "contracture and a **hypermetabolic crisis**"],
       ["Triggers", "**All volatile inhalational agents (halothane, isoflurane, sevoflurane, desflurane, enflurane) "
                    "and the depolarising relaxant SUCCINYLCHOLINE**. **Safe drugs: propofol, thiopentone, "
                    "benzodiazepines, opioids, nitrous oxide, non-depolarising relaxants, all local anaesthetics**"],
       ["Earliest / most sensitive sign", "**Unexplained rise in EtCO\u2082 (with tachypnoea/tachycardia)**; "
                                          "**masseter spasm/trismus** after succinylcholine is the classic early "
                                          "clinical clue. **Hyperthermia is a LATE sign** \u2014 do not wait for it"],
       ["Full picture", "Tachycardia, tachypnoea, **rigidity**, hypercarbia, hypoxaemia, mixed acidosis, "
                        "**hyperkalaemia**, arrhythmias, rising temperature (may rise 1\u20132 \u00b0C every 5 min), "
                        "**myoglobinuria/rhabdomyolysis (cola-coloured urine)**, raised CK, DIC, renal failure"],
       ["**Immediate management**", "**1** Call for help and the MH cart; **2** **STOP all triggering agents**, "
                                    "change to a clean circuit/soda lime, **hyperventilate with 100 % O\u2082 at "
                                    "\u2265 10 L/min**; **3** **DANTROLENE 2.5 mg/kg IV bolus, repeated every "
                                    "5\u201310 min until the crisis resolves \u2014 up to 10 mg/kg (may need "
                                    "30 mg/kg)**; **4** abandon/finish surgery as fast as possible; "
                                    "**5** **active cooling** \u2014 iced IV saline, surface ice packs, gastric/"
                                    "bladder/wound lavage with cold saline; **stop cooling at 38 \u00b0C** to avoid "
                                    "overshoot; **6** treat hyperkalaemia (insulin\u2013dextrose, calcium, "
                                    "bicarbonate), arrhythmias (**avoid calcium-channel blockers with dantrolene** "
                                    "\u2014 risk of hyperkalaemia and cardiac arrest), acidosis; **7** maintain "
                                    "urine output **> 1\u20132 mL/kg/h** (fluids, mannitol, furosemide) to protect "
                                    "the kidneys; **8** monitor core temperature, ABG, K\u207a, CK, coagulation, "
                                    "urine myoglobin; **9** ICU observation **\u2265 24\u201336 h** (recrudescence "
                                    "in up to 25 %); **10** counsel the patient and family, refer for "
                                    "**in-vitro contracture (caffeine-halothane) testing or genetic testing**, and "
                                    "issue an alert card"],
       ["Dantrolene practicalities", "Supplied as **20 mg lyophilised vials reconstituted with 60 mL sterile water** "
                                     "(newer formulation Ryanodex 250 mg/vial); a **36-vial stock** is recommended; "
                                     "acts by blocking the RYR1 channel. **Documentation of all drugs, lines and "
                                     "treatments is mandatory** (your PDF)"],
       ["Mortality", "**> 70\u201380 % untreated \u2192 < 5\u201310 % with prompt dantrolene**"]],
      widths=[3.6, 15.4])
h3("Anaphylaxis in the OR")
bullets([
    "Commonest triggers in the OR: **neuromuscular blocking agents, antibiotics (\u03b2-lactams), chlorhexidine, "
    "LATEX, colloids, contrast media**. **Latex anaphylaxis is typically delayed 20\u201360 min after exposure** "
    "(rather than immediate), because absorption is across mucosa/peritoneum \u2014 a favourite MCQ.",
    "Signs under anaesthesia: **unexplained hypotension and tachycardia (commonest), high airway pressure/"
    "bronchospasm, desaturation, rash/flushing, angio-oedema, cardiac arrest**.",
    "**Management:** stop the trigger \u2192 call for help \u2192 **100 % O\u2082, secure the airway** \u2192 "
    "**ADRENALINE \u2014 0.5 mg (0.5 mL of 1:1000) IM into the antero-lateral thigh, repeat every 5 min**, "
    "or **50 \u00b5g IV boluses titrated (1:10 000, 0.5 mL)** if monitored and experienced, then infusion \u2192 "
    "**rapid IV crystalloid 20 mL/kg**, legs elevated \u2192 second line: **chlorpheniramine 10 mg IV, "
    "hydrocortisone 200 mg IV, salbutamol nebulisation**, glucagon if on \u03b2-blockers.",
    "Take **mast-cell tryptase** samples (immediately, 1\u20132 h and 24 h) and refer for allergy testing; document "
    "and give the patient a written record.",
    "**Latex-safe environment:** first case of the day, latex-free gloves/tubing/catheters/drains, drug ampoules not "
    "vials with rubber bungs, dedicated **latex-free crash cart** (adrenaline, diphenhydramine, steroids, H\u2082 "
    "blockers, salbutamol, glucagon, latex-free airways and tubing), signage on the door. "
    "**Cross-reacting foods: banana, avocado, chestnut, kiwi** (also papaya, passion fruit, tomato, potato). "
    "**Highest-risk patients: spina bifida and children with multiple operations, plus health-care workers.**",
])
h3("Local anaesthetic systemic toxicity (LAST) \u2014 a modern must-know")
para("Progression: perioral tingling, metallic taste, tinnitus, visual disturbance, slurred speech, agitation \u2192 "
     "**seizures** \u2192 CNS depression \u2192 **arrhythmia and cardiac arrest** (bupivacaine is the most "
     "cardiotoxic). **Management: stop injecting, 100 % O\u2082, control seizures with a benzodiazepine (avoid "
     "propofol in cardiovascular instability), and give LIPID EMULSION 20 % \u2014 1.5 mL/kg bolus over 1 min then "
     "0.25 mL/kg/min infusion**; prolonged CPR may be required; **avoid vasopressin, calcium-channel blockers, "
     "\u03b2-blockers and further local anaesthetic**.")
h3("Fire in the operating room")
table(["Element", "Detail"],
      [["**Fire triangle**", "**FUEL + OXIDISER + IGNITION (heat) source** \u2014 all three are abundant in an OR "
                             "(some texts add a 4th: the chemical chain reaction \u2192 \u2018fire tetrahedron\u2019)"],
       ["Commonest **ignition** source", "**Electrosurgical unit (\u2248 70\u201390 % of OR fires)**, then the laser, "
                                         "then fibre-optic light cables and defibrillators"],
       ["Commonest **fuel**", "**Alcohol-based skin preparation solutions** (pooled or not dry), then drapes, gowns, "
                              "gauze, hair, plastic/rubber, ET tubes, GI bowel gas"],
       ["**Oxidiser**", "**Oxygen and nitrous oxide** \u2014 the OR is an \u2018oxygen-enriched atmosphere\u2019; "
                        "O\u2082 pools under drapes near the head (the classic **head-and-neck surgery fire**)"],
       ["Prevention", "Let alcohol prep **dry completely (\u2265 3 min)**, never allow pooling, blot puddles, remove "
                      "wet linen; keep FiO\u2082 **< 30 %** for head/neck cases and use a sealed airway; "
                      "**pause oxygen before activating the ESU/laser in the airway**; moisten sponges, drapes and "
                      "hair near the site; keep the pencil holstered; use laser-safe ET tubes with saline-filled "
                      "cuffs; scavenge oxygen from under the drapes"],
       ["**Response \u2014 R A C E**", "**R**escue/**R**emove the patient from danger \u2022 **A**larm (activate the "
                                       "alarm and call the code, e.g. **CODE RED**) \u2022 **C**onfine/**C**ontain "
                                       "(close doors, shut off medical gases) \u2022 **E**xtinguish or **E**vacuate"],
       ["**Extinguisher use \u2014 P A S S**", "**P**ull the pin \u2022 **A**im at the base of the fire \u2022 "
                                               "**S**queeze the handle \u2022 **S**weep side to side"],
       ["Extinguisher classes", "**A** = ordinary combustibles (paper, cloth) \u2013 water/foam \u2022 **B** = "
                                "flammable liquids \u2013 CO\u2082/dry chemical \u2022 **C** = **electrical** \u2013 "
                                "**CO\u2082 or dry chemical (never water)** \u2022 **D** = combustible metals \u2022 "
                                "**K** = kitchen fats. **The OR normally stocks a CO\u2082 or ABC dry-chemical "
                                "extinguisher; CO\u2082 is preferred for electrical/equipment fires because it "
                                "leaves no residue.** (Halon 1301 was the historical OR agent)"],
       ["**Airway fire drill**", "**1** Stop ventilation and **disconnect the oxygen/circuit immediately**; "
                                 "**2** **remove the endotracheal tube** and any burning material; "
                                 "**3** pour **saline** into the airway; **4** re-establish ventilation on air, "
                                 "**5** re-intubate and **bronchoscope** to assess thermal injury, "
                                 "**6** steroids/ICU care. Then preserve the equipment, file an incident report, and "
                                 "report the device to the regulator"],
       ["Drills", "Fire drills **at least twice a year, unannounced**; every employee must know the code name, the "
                  "location of the alarm, extinguishers, medical-gas shut-off valves and the evacuation route. "
                  "**Drapes and gowns are fire-RETARDANT, not fire-PROOF** \u2014 your PDF's repeated caution"]],
      widths=[3.8, 15.2])
h3("Mass-casualty / disaster preparedness \u2014 absent from your PDF")
table(["Concept", "Content"],
      [["Types", "**Internal** (fire, power failure, gas failure, structural, IT/cyber, violent intruder) vs "
                 "**External** (earthquake, flood, road/rail crash, bomb blast, chemical/radiation leak, epidemic)"],
       ["Framework", "**Hospital Emergency Incident Command System (HEICS/HICS)** with a written "
                     "**Hospital Disaster Management Plan** based on a **Hazard Vulnerability Analysis (HVA)**; "
                     "defined roles, call-tree, surge capacity, alternate site, mock drills at least annually"],
       ["**Triage** (Fr. __trier__, \u2018to sort\u2019)", "**START = Simple Triage And Rapid Treatment** "
                                                           "(\u2264 60 s per casualty), assessing **RPM \u2014 "
                                                           "Respiration, Perfusion, Mental status**: "
                                                           "walking wounded \u2192 **GREEN**; RR > 30, absent radial "
                                                           "pulse/capillary refill > 2 s, or not obeying commands "
                                                           "\u2192 **RED**; otherwise \u2192 **YELLOW**; "
                                                           "not breathing after airway opening \u2192 **BLACK**"],
       ["Triage colour codes", "**RED = Priority I \u2013 immediate** (life-threatening but salvageable: airway "
                               "obstruction, tension pneumothorax, major haemorrhage, shock) \u2022 "
                               "**YELLOW = Priority II \u2013 delayed** (can wait 1\u20134 h: long-bone fracture, "
                               "abdominal injury without shock) \u2022 **GREEN = Priority III \u2013 minor/ambulatory** "
                               "\u2022 **BLACK = Priority 0 \u2013 dead / expectant** (injuries incompatible with "
                               "survival given available resources)"],
       ["Key principle", "**In a mass casualty the aim shifts from \u2018the best for each individual\u2019 to "
                         "\u2018the greatest good for the greatest number\u2019** \u2014 hence expectant "
                         "categorisation, and **re-triage at every stage** (triage is dynamic)"],
       ["Pharmacist's / medical assistant's disaster role", "Maintain and rotate **disaster drug stock and "
                                                           "antidote kits** (atropine\u2013pralidoxime for "
                                                           "organophosphate, potassium iodide for radiation, "
                                                           "hydroxocobalamin/nitrites for cyanide, naloxone, "
                                                           "antivenom, tetanus toxoid/Ig), prepare IV fluids and "
                                                           "analgesics in bulk, control narcotics documentation, "
                                                           "advise on drug\u2013drug interactions and dosing in "
                                                           "renal failure/crush injury, oversee cold chain, and "
                                                           "manage donated-medicine screening"],
       ["Power failure", "Automatic **emergency generator** back-up (must start within 10 s and carry the OR, PACU, "
                         "ICU and lifts); **UPS** for monitors, ventilators and the anaesthesia machine; "
                         "**every team member must know where the high-intensity torches are kept**; know the "
                         "location of **manual (self-inflating) ventilation bags and the reserve oxygen cylinder** "
                         "for medical-gas failure"],
       ["Code colours (institution-specific \u2014 typical Indian hospital set)",
        "**Blue** = cardiac arrest \u2022 **Red** = fire \u2022 **Pink** = infant/child abduction \u2022 "
        "**Orange** = mass casualty / external emergency \u2022 **Yellow** = internal emergency / bomb threat (varies)"
        " \u2022 **Black** = bomb threat (varies) \u2022 **Brown** = evacuation \u2022 **Grey** = violent person / "
        "security \u2022 **Purple** = hostage. **Always answer per your own institution's published list.**"]],
      widths=[3.8, 15.2])



# ================================================================ PART C
h1("Part C  \u2014  POST-OPERATIVE PATIENT CARE  (the PDF's biggest omission)")
box("WHY THIS PART MATTERS",
    ["Your syllabus heading is \u201cPreparation of Patient for Operation, **Pre & Post Operative Patient Care**\u201d. "
     "Chapter 1 of the PDF ends at the sentence \u201cthe patient's care is transferred to the perioperative RN in the "
     "PACU, who now assumes the role of the patient's advocate\u201d \u2014 and the post-operative record is only "
     "reproduced as a blank form. **Everything in Part C is therefore new material, and it is where the majority of "
     "clinical MCQs from this unit are set.**"],
    fill=SH_GAP, border="9B6FC4", title_color=C_H3)

h2("C1.  Phases of Post-Operative Care & the PACU")
table(["Phase", "Location", "Focus"],
      [["**Phase I \u2014 immediate recovery**", "PACU / post-anaesthesia care unit (recovery room), "
                                                 "**1 nurse : 1\u20132 patients**",
        "Emergence from anaesthesia; airway, breathing, circulation, consciousness, temperature, pain, "
        "haemorrhage; continuous monitoring"],
       ["**Phase II \u2014 intermediate**", "Step-down/ambulatory unit or ward",
        "Preparation for discharge home or to the ward; oral intake, ambulation, teaching"],
       ["**Phase III \u2014 extended / convalescent**", "Ward, day-care lounge, home",
        "Rehabilitation, wound care, complication surveillance, follow-up"]],
      widths=[4.2, 5.6, 9.2])
h3("PACU hand-over \u2014 what must be communicated (SBAR / \u2018ISBAR\u2019)")
para("Patient's name, age and language \u2022 procedure performed and findings \u2022 surgeon and anaesthetist "
     "\u2022 **type of anaesthesia and reversal given** \u2022 airway (still intubated? LMA removed?) \u2022 "
     "**allergies** \u2022 relevant co-morbidity and medications \u2022 intra-operative course including vital-sign "
     "trends and any events \u2022 **estimated blood loss, fluids and blood products given, urine output** \u2022 "
     "drains, catheters, packs, dressings \u2022 analgesia and antiemetics given (drug, dose, time) \u2022 "
     "**laterality and position of the limb/dressing** \u2022 post-operative orders, expected concerns, and whom to "
     "call. **The anaesthetist does not leave until the PACU nurse accepts the patient and the hand-over is "
     "documented.**")
h3("PACU monitoring \u2014 what and how often")
bullets([
    "**Vital signs every 5\u201315 min** initially, then every 15 min until stable, then every 30 min: "
    "**RR/pattern, SpO\u2082, BP, pulse, temperature, ECG, EtCO\u2082 if intubated, pain score, level of "
    "consciousness**, plus **urine output and drain output hourly**.",
    "Position: **lateral or head-up 15\u201330\u00b0** if not contraindicated \u2014 keeps the tongue forward and "
    "allows secretions to drain (\u2018recovery position\u2019); **supine flat** after spinal anaesthesia or "
    "hypotension.",
    "**Oxygen** by mask/nasal prongs until awake and SpO\u2082 is adequate on air; **side rails up, never leave the "
    "patient unattended**, keep suction and an Ambu bag at the bedside.",
    "Check: **dressing (ooze/soakage), surgical site, drain patency and character, IV site, peripheral pulses and "
    "capillary refill distal to any cast/tourniquet, ability to move limbs and any paraesthesia (rule out positioning "
    "injury and assess spinal regression), bladder distension.**",
    "**Warming** \u2014 forced-air warming blanket, warmed IV fluids, raised room temperature; aim core temperature "
    "**> 36 \u00b0C** before discharge from PACU.",
])
h3("Aldrete Score (Post-Anaesthesia Recovery Score) \u2014 memorise this table")
table(["Parameter", "2 points", "1 point", "0 points"],
      [["**A**ctivity (voluntary movement)", "Moves **all four limbs** on command", "Moves **two** extremities",
        "Unable to move"],
       ["**R**espiration", "Breathes deeply and **coughs freely**", "Dyspnoeic, shallow or limited breathing",
        "Apnoeic / needs ventilation"],
       ["**C**irculation (BP)", "BP **\u00b1 20 %** of pre-anaesthetic level", "BP \u00b1 20\u201349 %",
        "BP \u00b1 \u2265 50 %"],
       ["**C**onsciousness", "**Fully awake**", "Rousable on calling", "Not responding"],
       ["**O**xygen saturation", "**SpO\u2082 > 92 % on room air**", "Needs O\u2082 to keep SpO\u2082 > 90 %",
        "SpO\u2082 < 90 % even with O\u2082"]],
      widths=[4.6, 5.4, 4.6, 4.4])
box("ALDRETE \u2014 EXAM POINTS",
    ["**Maximum score = 10.** Discharge from PACU requires **\u2265 9** (with no single parameter scoring 0), "
     "**and** stable vital signs, controlled pain and PONV, no active bleeding, and normothermia.",
    "**Mnemonic: A-R-C-C-O** (Activity, Respiration, Circulation, Consciousness, Oxygenation). "
     "The **original 1970 Aldrete score used skin COLOUR** as the fifth criterion; the **Modified Aldrete score "
     "replaced colour with pulse-oximetry SpO\u2082**.",
     "**Pain and nausea are NOT part of the Aldrete score** \u2014 they belong to the **PADSS**. This inversion is a "
     "classic MCQ.",
     "**PADSS (Post-Anaesthesia Discharge Scoring System, for day-care/ambulatory discharge HOME)** scores 0\u20132 "
     "each for: **vital signs \u2022 ambulation \u2022 nausea/vomiting \u2022 pain \u2022 surgical bleeding**; "
     "**\u2265 9 required for discharge home**. Modern \u2018fast-tracking\u2019 also requires: an escort, written "
     "instructions, ability to tolerate oral fluids, and \u2014 in most protocols \u2014 the **ability to void is no "
     "longer mandatory** except after spinal/anorectal/hernia surgery.",
     "After spinal/epidural anaesthesia additionally document **return of motor power, sensation and "
     "proprioception**, and the **sensory level**."],
    fill=SH_HY, border="D98BA0", title_color=C_ANS)

h2("C2.  Immediate Post-Operative Complications \u2014 recognise and act")
h3("A.  Respiratory")
table(["Complication", "Cause", "Recognition", "Management"],
      [["**Airway obstruction (commonest immediate PACU emergency)**",
        "**Tongue falling back against the posterior pharynx** in the sedated supine patient; secretions, blood, "
        "laryngeal oedema, foreign body, dentures",
        "**Snoring/stridor, see-saw (paradoxical) chest movement, tracheal tug, no air movement**, desaturation",
        "**Head tilt\u2013chin lift / jaw thrust**, lateral position, suction, oral/nasal airway, 100 % O\u2082; "
        "re-intubate if needed"],
       ["**Laryngospasm**", "Stimulation of the cords during light anaesthesia, secretions, blood",
        "**Inspiratory stridor progressing to complete silence**, paradoxical effort, desaturation",
        "Remove the stimulus, suction, **100 % O\u2082 with CPAP and jaw thrust**; if it persists \u2014 "
        "**propofol, then succinylcholine 0.1\u20131 mg/kg** and intubate"],
       ["**Hypoventilation / respiratory depression**",
        "**Residual anaesthetic, opioids, incomplete reversal of neuromuscular blockade**, splinting from pain, "
        "obesity, COPD, hypothermia",
        "Slow shallow breathing, \u2191 EtCO\u2082, \u2193 SpO\u2082, drowsiness, pin-point pupils (opioid)",
        "Stimulate, O\u2082, support ventilation; **naloxone 40\u201380 \u00b5g IV increments** for opioids; "
        "**neostigmine/sugammadex** for residual relaxant; **flumazenil** for benzodiazepines"],
       ["**Hypoxaemia**", "Atelectasis, hypoventilation, V/Q mismatch, diffusion hypoxia, aspiration, "
                          "pulmonary oedema, pneumothorax",
        "SpO\u2082 < 90 %, restlessness (an **early** sign), tachycardia, cyanosis (**late**)",
        "O\u2082, sit up, deep breathing/incentive spirometry, treat cause, chest X-ray"],
       ["**Aspiration**", "Full stomach, obtunded reflexes", "Cough, wheeze, tachypnoea, desaturation, "
                                                             "right-lower-lobe infiltrate",
        "Head-down lateral, suction, O\u2082, bronchoscopy if particulate; **steroids and prophylactic antibiotics "
        "are NOT routinely indicated**"],
       ["**Atelectasis**", "**Commonest post-operative pulmonary complication**; shallow breathing, retained "
                           "secretions, upper-abdominal/thoracic incision",
        "**Low-grade fever POD 1\u20132**, decreased breath sounds, dull percussion at the bases, tachypnoea",
        "**Deep breathing + incentive spirometry (10 breaths/hour), early mobilisation, adequate analgesia, "
        "chest physiotherapy, humidified O\u2082**"]],
      widths=[3.2, 4.4, 5.0, 6.4])
h3("B.  Cardiovascular")
table(["Problem", "Common causes", "Action"],
      [["**Hypotension**", "**Hypovolaemia (commonest \u2014 blood loss, third-space loss, inadequate replacement)**, "
                           "residual anaesthetic/vasodilatation, sympathetic block from spinal/epidural, "
                           "opioids, arrhythmia, myocardial ischaemia, tension pneumothorax, tamponade, sepsis, "
                           "anaphylaxis",
        "**Assess the surgical site and drains for bleeding first**; head-down/legs up, O\u2082, "
        "**rapid crystalloid bolus 250\u2013500 mL (10\u201320 mL/kg)**, vasopressor if needed, ECG, Hb, "
        "call surgeon/anaesthetist; **re-exploration** for surgical bleeding"],
       ["**Hypertension**", "**Pain (commonest), bladder distension**, hypoxia/hypercarbia, hypothermia with "
                            "shivering, pre-existing hypertension with omitted drugs, fluid overload, emergence "
                            "excitement",
        "**Treat the cause first \u2014 analgesia, catheterise/scan the bladder, oxygenate, warm** \u2014 before "
        "reaching for antihypertensives (labetalol, GTN, hydralazine)"],
       ["**Arrhythmia**", "**Sinus tachycardia** (pain, hypovolaemia, fever, anaemia, anxiety, anticholinergics) "
                          "\u2022 **bradycardia** (opioids, neostigmine, vagal traction, hypoxia, "
                          "\u03b2-blockade, spinal to T1\u2013T4) \u2022 **AF/VT** (electrolyte imbalance, ischaemia, "
                          "hypoxia, catecholamines)",
        "Correct **hypoxia, hypovolaemia, K\u207a/Mg\u00b2\u207a, acidosis and pain**; ECG; atropine for symptomatic "
        "bradycardia; treat per ACLS if unstable"],
       ["**Myocardial ischaemia/infarction**", "Peak incidence **POD 0\u20133**; often **silent or atypical** in the "
                                               "post-op period because of analgesia",
        "**Any unexplained hypotension, arrhythmia, dyspnoea or confusion after surgery \u2192 ECG + troponin**; "
        "O\u2082, analgesia, cardiology referral"],
       ["**Haemorrhage / shock**", "**Reactionary haemorrhage** (within 24 h \u2014 slipped ligature, "
                                   "dislodged clot as BP rises and vasoconstriction wears off) \u2022 "
                                   "**Secondary haemorrhage** (POD **7\u201314 \u2014 infection eroding a vessel**) "
                                   "\u2022 primary = during surgery",
        "**Class I\u2013IV haemorrhagic shock**: tachycardia and narrowed pulse pressure precede hypotension; "
        "**tachycardia with a normal BP in a young patient is compensated shock \u2014 do not be reassured**. "
        "Two large-bore cannulae, crystalloid, **cross-matched blood / massive transfusion protocol (1:1:1 RBC:FFP:"
        "platelets)**, tranexamic acid, keep warm, **surgical control of bleeding is definitive**"]],
      widths=[3.4, 7.0, 8.6])
h3("C.  Neurological & thermoregulatory")
table(["Problem", "Key facts"],
      [["**Delayed emergence**", "Residual anaesthetic/opioid/relaxant \u2022 **hypoglycaemia, hypoxia, hypercarbia, "
                                 "hyponatraemia, hypothermia (< 33 \u00b0C prevents awakening)** \u2022 raised ICP, "
                                 "stroke, seizure \u2022 central anticholinergic syndrome. **Check glucose, ABG, "
                                 "electrolytes and pupils; reverse opioids/benzodiazepines; think stroke if focal "
                                 "signs**"],
       ["**Emergence delirium / agitation**", "Commonest in **children and the elderly**; also hypoxia (**always "
                                              "exclude hypoxia first**), pain, full bladder, ketamine, "
                                              "anticholinergics, alcohol/benzodiazepine withdrawal. Reassure, "
                                              "orientate, involve parents, treat pain; avoid physical restraint"],
       ["**Post-operative cognitive dysfunction (POCD) / post-op delirium**",
        "Elderly, pre-existing cognitive impairment, major/cardiac surgery, sepsis, sleep disruption, "
        "anticholinergic burden. Prevent: orientation, glasses/hearing aids, sleep hygiene, avoid benzodiazepines and "
        "anticholinergics, treat pain and constipation, early mobilisation"],
       ["**Hypothermia (core < 36 \u00b0C)**", "Causes: cold OR, anaesthetic-induced vasodilatation and loss of "
                                               "thermoregulation, cold fluids/irrigation, large exposed cavity, "
                                               "elderly/neonates, prolonged surgery. **Consequences: shivering "
                                               "(\u2191 O\u2082 consumption up to 400 %), coagulopathy and \u2191 "
                                               "blood loss, \u2191 SSI, arrhythmia and cardiac events, delayed drug "
                                               "metabolism and delayed emergence, patient discomfort.** "
                                               "**Prevent/treat: forced-air warming, warmed IV fluids and irrigation, "
                                               "warmed blankets, raised room temperature, heat-moisture exchanger, "
                                               "limit exposure**"],
       ["**Shivering**", "Very common; treat by **warming** plus **pethidine (meperidine) 12.5\u201325 mg IV** "
                         "(most effective), clonidine, dexmedetomidine, tramadol, ondansetron; give oxygen"],
       ["**Malignant hyperthermia**", "Can present **in the PACU** as well as intra-operatively \u2014 see \u00a7B13"],
       ["**Nerve injury / paralysis**", "Examine and document limb movement, power and sensation on arrival; "
                                        "**paraesthesia or inability to move a limb \u2192 exclude positioning "
                                        "injury, epidural haematoma (back pain + bilateral weakness + urinary "
                                        "retention = EMERGENCY, needs urgent MRI and decompression within 8 h) and "
                                        "compartment syndrome**"]],
      widths=[4.2, 14.8])

h2("C3.  Post-Operative Pain Management")
table(["Element", "Detail"],
      [["Assessment", "**Self-report is the gold standard**: Numeric Rating Scale 0\u201310, Visual Analogue Scale, "
                      "Wong-Baker **FACES** scale (children \u2265 3 y), **FLACC** (Face, Legs, Activity, Cry, "
                      "Consolability \u2014 pre-verbal children), **FPS-R**; behavioural/physiological cues in the "
                      "unconscious. **Assess at rest and on movement, and re-assess after each intervention**"],
       ["Consequences of untreated pain", "Splinting \u2192 **atelectasis and pneumonia**; immobility \u2192 **DVT**; "
                                          "sympathetic surge \u2192 tachycardia, hypertension, ischaemia; ileus; "
                                          "urinary retention; delirium; poor sleep; **chronic post-surgical pain**"],
       ["**Multimodal analgesia** (the modern standard)",
        "Combining drugs with different mechanisms lowers each dose and its side-effects: "
        "**paracetamol 1 g q6h + an NSAID (diclofenac/ketorolac) + local/regional block \u00b1 opioid \u00b1 "
        "adjuvants (gabapentinoid, ketamine, dexmedetomidine, dexamethasone, IV lignocaine, magnesium)**"],
       ["WHO analgesic ladder", "**Step 1** non-opioid \u00b1 adjuvant \u2192 **Step 2** weak opioid (tramadol, "
                                "codeine) \u00b1 non-opioid \u2192 **Step 3** strong opioid (morphine, fentanyl). "
                                "For acute post-operative pain the ladder is used **downwards** (start strong, step "
                                "down)"],
       ["Routes", "IV titration in PACU \u2192 **PCA (patient-controlled analgesia)** \u2192 oral once tolerating. "
                  "**IM injections are discouraged** (painful, erratic absorption). Avoid the oral route until bowel "
                  "function returns after major GI surgery"],
       ["**PCA**", "Typical morphine PCA: **bolus 1 mg, lockout 5\u201310 min, no background infusion** in "
                   "opioid-naive adults; the **patient must be awake to press the button \u2014 the built-in safety "
                   "feature**; **never allow a relative to press it**; monitor sedation score and RR"],
       ["Regional techniques", "**Epidural** (thoracic for laparotomy/thoracotomy \u2014 best dynamic pain relief, "
                               "reduces pulmonary complications and ileus; watch for hypotension, motor block, "
                               "urinary retention, **epidural haematoma/abscess**), spinal opioids, "
                               "**TAP block, rectus-sheath, paravertebral, brachial plexus, femoral/adductor canal, "
                               "erector-spinae** blocks, wound infiltration and catheters"],
       ["Opioid adverse effects (\u2018N-CROP\u2019)", "**N**ausea/vomiting \u2022 **C**onstipation (does not develop "
                                                       "tolerance \u2192 always prescribe a laxative) \u2022 "
                                                       "**R**espiratory depression (**sedation precedes it \u2014 "
                                                       "monitor the sedation score, not just RR**) \u2022 "
                                                       "**O**bstruction/ileus and urinary retention \u2022 "
                                                       "**P**ruritus, plus miosis, hypotension, tolerance"],
       ["NSAID cautions", "Renal impairment, hypovolaemia, GI bleeding/ulcer, asthma, elderly, bleeding risk, "
                          "**avoid after bariatric anastomoses and in some bone-graft/fusion surgery**; "
                          "ketorolac max 5 days"],
       ["Paracetamol", "**Maximum 4 g/24 h (3 g in the frail, low weight, liver disease or alcohol excess)** \u2014 "
                       "watch for duplication in combination products, a classic pharmacist intervention"]],
      widths=[3.6, 15.4])

h2("C4.  Post-Operative Nausea & Vomiting (PONV)")
table(["Item", "Detail"],
      [["Incidence", "**20\u201330 % overall; up to 70\u201380 % in high-risk patients.** Commonest reason for "
                     "unanticipated admission after day surgery, and the outcome patients fear most after pain"],
       ["**Apfel simplified risk score** (adults)", "**1 point each: female sex \u2022 non-smoker \u2022 previous "
                                                    "PONV or motion sickness \u2022 expected post-operative "
                                                    "opioid use.** Risk with 0/1/2/3/4 factors \u2248 "
                                                    "**10 / 20 / 40 / 60 / 80 %**"],
       ["Other risk factors", "Younger age, general (vs regional) anaesthesia, **volatile agents and nitrous oxide**, "
                              "duration of anaesthesia, neostigmine in high dose, type of surgery "
                              "(**gynaecological, laparoscopic, ENT/middle-ear, squint, breast, bariatric**), "
                              "gastric distension, hypotension, dehydration, pain, early mobilisation, opioids"],
       ["Drug classes & receptors", "**5-HT\u2083 antagonists** \u2014 **ondansetron 4\u20138 mg IV** (give at the "
                                    "**end** of surgery; SE: headache, constipation, **QT prolongation**), "
                                    "palonosetron \u2022 **Corticosteroid \u2014 dexamethasone 4\u20138 mg IV** "
                                    "(give at induction; also improves analgesia; watch glucose) \u2022 "
                                    "**D\u2082 antagonists** \u2014 **metoclopramide** (prokinetic; extrapyramidal "
                                    "effects, avoid in obstruction), droperidol (QT), domperidone \u2022 "
                                    "**NK\u2081 antagonist \u2014 aprepitant** \u2022 **Antihistamine \u2014 "
                                    "cyclizine, promethazine** \u2022 **Anticholinergic \u2014 hyoscine (scopolamine) "
                                    "patch** \u2022 **Others** \u2014 propofol TIVA (antiemetic), "
                                    "acupressure at **P6 (Neiguan)**"],
       ["Strategy", "**Risk-adapted: 1\u20132 risk factors \u2192 2 prophylactic agents from DIFFERENT classes; "
                    "\u2265 3 \u2192 3\u20134 agents + TIVA with propofol, avoid nitrous oxide and volatile agents, "
                    "minimise opioids (regional/multimodal), ensure adequate hydration.** For breakthrough PONV, "
                    "use a drug from a **class not already given**"],
       ["Nursing measures", "Nil by mouth until nausea settles then sips, avoid sudden movement, oral hygiene, "
                            "**decompress the stomach with an NG tube if distended**, treat pain and hypotension, "
                            "correct electrolytes, IV fluids, cool cloth, position to avoid aspiration"]],
      widths=[3.6, 15.4])

h2("C5.  Post-Operative Fever \u2014 the \u201c6 W's\u201d timeline (very high yield)")
table(["Onset", "\u201cW\u201d", "Cause", "Clues & action"],
      [["**Immediate \u2014 during or within hours**", "**W**onder drugs / **W**hat did we do?",
        "Drug or transfusion reaction, **malignant hyperthermia**, pre-existing infection, trauma of surgery "
        "(cytokine response), thyroid storm",
        "MH \u2192 rigidity + \u2191 EtCO\u2082 \u2192 dantrolene. Stop the transfusion for any reaction"],
       ["**POD 0\u20132 (24\u201348 h)**", "**W**ind", "**ATELECTASIS (commonest cause of early post-op fever)**; "
                                                       "aspiration, early pneumonia",
        "Low-grade fever, \u2193 basal breath sounds. **Incentive spirometry, deep breathing, ambulation, analgesia, "
        "physiotherapy** \u2014 not antibiotics"],
       ["**POD 3\u20135**", "**W**ater", "**Urinary tract infection** (catheter-associated)",
        "Dysuria/suprapubic pain, urinalysis + culture; **remove or change the catheter**, antibiotics"],
       ["**POD 4\u20136**", "**W**alking (\u2018Wein\u2019)", "**Deep-vein thrombosis / pulmonary embolism**",
        "Calf pain/swelling, unilateral oedema, sudden dyspnoea, pleuritic pain, tachycardia, "
        "hypoxia. Doppler/CTPA; anticoagulate"],
       ["**POD 5\u20137**", "**W**ound", "**Surgical site infection**; also anastomotic leak (POD 5\u20137, "
                                         "with tachycardia and ileus)",
        "Inspect \u2014 erythema, tenderness, induration, discharge; **open and drain, culture, antibiotics**; "
        "suspect a leak if the patient is 'not doing well' with tachycardia"],
       ["**POD > 7**", "**W**onder drugs & **W**ounder-things",
        "**Drug fever** (antibiotics, heparin), **abscess/collection**, C. difficile colitis, thrombophlebitis at the "
        "IV site, central-line infection, sinusitis, acalculous cholecystitis, transfusion-related",
        "Review the drug chart, culture blood/urine/line tips, CT for a collection"]],
      widths=[3.4, 2.8, 5.8, 7.0])
box("MNEMONIC \u2014 \u201c**W**ind, **W**ater, **W**ound, **W**alking, **W**onder drugs (+ **W**hat we did / "
    "**W**onder about lines)\u201d",
    ["Sequence in time: **Wind (1\u20132) \u2192 Water (3\u20135) \u2192 Walking/Wein (4\u20136) \u2192 "
     "Wound (5\u20137) \u2192 Wonder drugs (> 7)**.",
     "**A fever in the first 24 h is usually NOT infective** \u2014 it is the inflammatory/cytokine response to "
     "surgery, atelectasis, a drug or transfusion reaction, or malignant hyperthermia.",
     "**Fever after POD 5 with wound signs = infection until proved otherwise.**"],
    fill=SH_MNE, border="6DA97A", title_color=C_H2)

h2("C6.  Wound Healing, Wound Care and Wound Complications")
h3("Phases of wound healing")
table(["Phase", "Timing", "Events", "Dominant cell"],
      [["**Haemostasis**", "Immediate (minutes\u2013hours)", "Vasoconstriction, platelet plug, fibrin clot",
        "Platelets"],
       ["**Inflammatory (lag/substrate) phase**", "**Day 0\u20134/5**", "Vasodilatation, exudate, phagocytosis, "
                                                                       "debridement. Clinically **rubor, calor, "
                                                                       "tumor, dolor**; wound held only by the clot "
                                                                       "and sutures",
        "**Neutrophils (24\u201348 h) \u2192 Macrophages (48\u201372 h \u2014 the key orchestrating cell)**"],
       ["**Proliferative (fibroplasia / granulation) phase**", "**Day 4\u201321**", "Fibroblast migration, "
                                                                                   "**collagen III deposition**, "
                                                                                   "angiogenesis, granulation tissue, "
                                                                                   "epithelialisation, wound "
                                                                                   "contraction",
        "**Fibroblasts** (+ endothelial cells, myofibroblasts)"],
       ["**Maturation / remodelling phase**", "**Day 21 \u2192 up to 1\u20132 years**", "**Collagen III replaced by "
                                                                                       "collagen I**, cross-linking, "
                                                                                       "scar becomes pale and flat, "
                                                                                       "tensile strength rises",
        "Fibroblasts, collagenase (MMPs)"]],
      widths=[4.2, 3.0, 7.4, 4.4])
bullets([
    "**Tensile strength: \u2248 5\u201310 % at 1 week, 30\u201350 % at 1 month, and only ever reaches "
    "70\u201380 % of unwounded skin.** This is why the wound is at greatest risk of dehiscence in the **first week** "
    "and why heavy lifting is restricted for 6 weeks.",
    "**Types of healing** \u2014 **Primary (first) intention:** clean wound, apposed edges, minimal granulation, fine "
    "scar. **Secondary intention:** wound left open, heals by granulation, contraction and epithelialisation "
    "(contaminated/infected wounds, abscess cavities) \u2192 broad scar. **Tertiary / delayed primary closure:** "
    "left open 3\u20135 days then closed once contamination is controlled.",
    "**Factors delaying healing** \u2014 local: **infection, ischaemia/poor perfusion, haematoma/seroma, foreign "
    "body, tension, dead space, radiation, movement, poor apposition**; systemic: **age, diabetes, malnutrition "
    "(protein, vitamin C \u2192 collagen cross-linking, vitamin A, zinc), anaemia and hypoxia, obesity, smoking, "
    "steroids and immunosuppressants, chemotherapy, jaundice, uraemia, malignancy**.",
    "**Dressing:** the first dressing over a primarily closed wound is usually left undisturbed for "
    "**24\u201348 h** (epithelialisation is complete by then and it forms a bacterial barrier); thereafter change "
    "using **aseptic (sterile) technique**, inspect for erythema/discharge, and document. Change immediately if "
    "soaked (strike-through).",
    "**Suture removal times** \u2014 face/neck **3\u20135 days**; scalp **7\u201310**; trunk/abdomen "
    "**7\u201310 (up to 12)**; upper limb **10\u201314**; lower limb **10\u201314**; over joints, back, palms and "
    "soles **14 days**; children heal faster (shorter times), the elderly/steroid users need longer.",
])
h3("Wound complications")
table(["Complication", "Timing & features", "Management"],
      [["**Haematoma**", "Within 24\u201348 h; swelling, discoloration, pain, may compress "
                         "(**neck haematoma after thyroidectomy = airway emergency**)",
        "Small \u2192 observe; expanding/compressive \u2192 **evacuate, open the wound (bedside for neck), "
        "secure haemostasis**"],
       ["**Seroma**", "Days 3\u201310, especially after **mastectomy, axillary/groin dissection, hernia mesh**; "
                      "fluctuant, painless swelling",
        "Aspirate under aseptic technique, compression, drains; may recur"],
       ["**Surgical site infection**", "**POD 5\u20137**; pain, erythema, warmth, induration, purulent discharge, "
                                       "fever, wound edge separation. (Very early, POD 1\u20132, spreading "
                                       "cellulitis/crepitus \u2192 suspect **\u03b2-haemolytic streptococcus or "
                                       "clostridial/necrotising infection \u2014 surgical emergency**)",
        "**Open, drain and culture; debride; pack for healing by secondary intention; antibiotics guided by culture** "
        "\u2014 \u2018the treatment of pus is drainage\u2019"],
       ["**Wound dehiscence**", "**POD 5\u201310 (classically day 7\u20138)**; the herald sign is a "
                                "**sudden discharge of serosanguineous (\u2018salmon-pink\u2019) fluid** from the "
                                "wound and a palpable gap; patient may report a \u2018giving-way\u2019 sensation on "
                                "coughing",
        "Reinforce with an abdominal binder, avoid straining, **inform the surgeon \u2192 usually re-suturing in "
        "theatre** with tension sutures"],
       ["**Evisceration (burst abdomen)**", "Protrusion of viscera through the wound; often follows a cough or "
                                            "straining",
        "**EMERGENCY: stay with the patient, place them supine with knees flexed (low Fowler's), COVER the viscera "
        "with sterile gauze soaked in warm sterile normal saline (never dry gauze, never attempt to push the bowel "
        "back), keep NBM, IV access and fluids, monitor for shock, analgesia, notify the surgeon and prepare for "
        "immediate theatre.**"],
       ["**Incisional hernia**", "Weeks\u2013years later, a late consequence of a subclinical dehiscence or infection",
        "Elective repair, usually with mesh"],
       ["**Hypertrophic scar vs keloid**", "**Hypertrophic** \u2014 raised but **stays within** the original wound "
                                           "margins, tends to regress. **Keloid** \u2014 **extends beyond** the "
                                           "original margins, does not regress, recurs after excision; commoner in "
                                           "darker skin, over the sternum, shoulder and ear lobe",
        "Silicone sheeting, pressure, intralesional steroid, excision + adjuvant therapy for keloid"],
       ["**Wound sinus / fistula**", "Persistent discharge; a retained suture, foreign body, necrotic tissue or a "
                                     "communication with a viscus (**or a retained sponge \u2014 gossypiboma**)",
        "Investigate (imaging, sinogram), excise the tract, remove the cause"]],
      widths=[3.2, 8.2, 7.6])

h2("C7.  Drains, Tubes and Catheters")
table(["Device", "Type", "Points"],
      [["**Penrose**", "**Passive, open**", "Soft latex; drains by gravity/capillary action; higher infection risk"],
       ["**Corrugated / sheet drain**", "Passive, open", "Retro-peritoneal and abscess cavities"],
       ["**Jackson-Pratt (JP)**", "**Closed, active (bulb suction)**", "Bulb compressed to create low negative "
                                                                      "pressure; measure and record output; "
                                                                      "**re-compress the bulb after emptying** "
                                                                      "(a fully expanded bulb means the suction is "
                                                                      "lost)"],
       ["**Hemovac / Redivac**", "**Closed, active (spring or high vacuum)**", "Used after orthopaedic, breast and "
                                                                              "neck surgery; Redivac is high-vacuum"],
       ["**Sump / double-lumen**", "Active", "Air vent prevents tissue being sucked into the drain"],
       ["**Chest (intercostal) drain**", "**Closed, underwater seal**", "**Keep the bottle below the level of the "
                                                                       "chest at all times; never clamp a bubbling "
                                                                       "drain; \u2018swinging\u2019 with respiration "
                                                                       "= patent; continuous bubbling = air leak.** "
                                                                       "If it disconnects \u2192 place the end in "
                                                                       "sterile water; if it falls out \u2192 cover "
                                                                       "with an occlusive dressing taped on 3 sides"],
       ["**Nasogastric tube**", "Decompression / feeding", "Confirm position (**pH < 5.5 on aspirate, or X-ray \u2014 "
                                                           "the auscultation \u2018whoosh\u2019 test is unreliable "
                                                           "and unsafe**); record aspirate volume and character; "
                                                           "routine NG use after laparotomy is **no longer "
                                                           "recommended (ERAS)**"],
       ["**Urinary (Foley) catheter**", "Closed drainage", "Aseptic insertion, **closed system, bag always below "
                                                           "bladder level and off the floor, no dependent loops, "
                                                           "no routine bladder irrigation, remove within 24 h**. "
                                                           "**CAUTI is the commonest healthcare-associated "
                                                           "infection.** Report output **< 0.5 mL/kg/h "
                                                           "(< 30 mL/h in an adult)**"]],
      widths=[3.4, 3.6, 12.0])
bullets([
    "**Purposes of a drain:** evacuate an existing collection, prevent accumulation of blood/serum/pus/bile/urine in "
    "a dead space, warn of a leak (\u2018sentinel\u2019 drain), and allow a controlled tract to form.",
    "**Nursing responsibilities:** secure to prevent traction, keep patent and unkinked, measure and chart the "
    "**volume, colour and character** each shift, aseptic dressing at the exit site, observe for bleeding or "
    "enteric/bile content, and **never advance a drain back in**. Remove when output is minimal (typically "
    "< 25\u201350 mL/24 h) or as ordered.",
    "**Bright red brisk drain output \u2192 haemorrhage; faecal or bilious output \u2192 anastomotic leak; "
    "cloudy/foul \u2192 infection** \u2014 report immediately.",
])

h2("C8.  Fluids, Electrolytes, Nutrition and the GI Tract")
bullets([
    "**Maintenance fluid (4-2-1 rule):** 4 mL/kg/h for the first 10 kg + 2 mL/kg/h for the next 10 kg + "
    "1 mL/kg/h for each kg above 20 \u2014 or roughly **25\u201330 mL/kg/day** with Na\u207a 1 mmol/kg/day and "
    "K\u207a 1 mmol/kg/day; add replacement of measured losses (drains, NG aspirate, urine) and third-space losses.",
    "**Monitor:** hourly urine output (**target > 0.5 mL/kg/h**), pulse, BP, capillary refill, JVP, daily weight, "
    "fluid balance chart, electrolytes, urea/creatinine, lactate.",
    "**Post-operative fluid retention (the \u2018stress response\u2019):** ADH and aldosterone rise \u2192 salt and "
    "water retention, hyponatraemia and oliguria in the first 24\u201348 h \u2014 **oliguria after surgery is more "
    "often hypovolaemia than renal failure, but beware of over-transfusing hypotonic fluid \u2192 dilutional "
    "hyponatraemia (SIADH)**. Also expect **hyperglycaemia, hypokalaemia, negative nitrogen balance and a "
    "catabolic state**.",
    "**Post-operative ileus:** physiological after laparotomy; recovery order = **small intestine within "
    "\u2248 24 h \u2192 stomach 24\u201348 h \u2192 COLON 48\u201372 h (last to recover)**. Diagnose return of "
    "function by **passage of flatus/stool** (bowel sounds are unreliable). **Paralytic ileus** beyond 3\u20135 days, "
    "or with vomiting and distension, is pathological \u2014 look for **hypokalaemia, hypomagnesaemia, opioids, "
    "anticholinergics, peritonitis, anastomotic leak, retroperitoneal haematoma**. Treat: NBM/NG decompression, "
    "correct electrolytes, minimise opioids, mobilise, chewing gum and early feeding help.",
    "**Distinguish ileus from mechanical obstruction:** ileus \u2014 **absent** bowel sounds, generalised "
    "distension, gas throughout including the colon; obstruction \u2014 **colicky pain, high-pitched/tinkling bowel "
    "sounds, no gas beyond the obstruction**, often POD > 7 (adhesions).",
    "**Nutrition (ERAS):** **early oral feeding within 24 h** is safe after most surgery and reduces complications; "
    "enteral is always preferred to parenteral (\u2018**if the gut works, use it**\u2019); TPN only when the gut "
    "cannot be used for > 7 days. Watch for **refeeding syndrome** (\u2193 PO\u2084, \u2193 K\u207a, \u2193 Mg\u00b2\u207a) "
    "in the malnourished.",
    "**Constipation:** almost universal with opioids \u2014 prescribe a **stimulant + softener laxative "
    "prophylactically**, encourage fluids, fibre and mobilisation.",
    "**Urinary retention:** commonest after **anorectal, hernia, pelvic and spinal-anaesthetic** surgery, in older "
    "men and with opioids/anticholinergics. Recognise by suprapubic fullness, restlessness, overflow dribbling, "
    "**bladder scan > 500\u2013600 mL**; try privacy, standing/sitting position, running water, warm compress, "
    "analgesia \u2192 then **in-and-out catheterisation**.",
])

h2("C9.  Post-Operative VTE, Mobilisation and Respiratory Care")
bullets([
    "**DVT:** peak POD **5\u201310**; often **clinically silent**; calf pain/tenderness, unilateral swelling "
    "(> 3 cm difference), warmth, pitting oedema, dilated superficial veins. (Homans' sign is unreliable and should "
    "not be elicited.) Diagnose with **compression ultrasound + D-dimer/Wells score**.",
    "**PE:** sudden dyspnoea, pleuritic chest pain, tachycardia, tachypnoea, hypoxia, haemoptysis, syncope, "
    "unexplained hypotension. **Massive PE is a leading cause of preventable post-operative death.** Investigate with "
    "**CTPA**; treat with anticoagulation \u00b1 thrombolysis.",
    "**Prevention** = mechanical (early ambulation, ankle exercises, TED stockings, intermittent pneumatic "
    "compression) **+** pharmacological (LMWH), plus **adequate hydration and analgesia so the patient can move** "
    "\u2014 see \u00a7A7.",
    "**Respiratory care bundle:** semi-Fowler's position, **deep breathing and incentive spirometry 10 breaths "
    "hourly while awake**, **splinted (pillow-supported) coughing**, adequate analgesia (a patient who cannot cough "
    "because of pain will get atelectasis), early ambulation **within 24 h**, chest physiotherapy, oral hygiene "
    "(reduces ventilator-associated and post-op pneumonia), smoking cessation, humidified oxygen, nebulisers.",
])

h2("C10.  Discharge, Counselling and the Pharmacist's Role")
bullets([
    "**Discharge criteria (day-care/ambulatory):** stable vital signs for \u2265 1 h, oriented, "
    "**PADSS \u2265 9**, pain and PONV controlled on oral medication, able to take oral fluids, minimal bleeding, "
    "able to ambulate/dress with minimal help, **a responsible escort and transport**, and **written** instructions "
    "with an emergency contact number.",
    "**Counsel the patient not to drive, operate machinery, sign legal documents, drink alcohol or care for a child "
    "unaided for at least 24 hours** after general anaesthesia or sedation.",
    "**Written discharge information must cover:** wound and dressing care, when and how to shower/bathe, suture or "
    "staple removal date, activity and lifting restrictions, driving and return-to-work advice, diet, "
    "**red-flag symptoms requiring urgent review \u2014 fever, increasing pain, spreading redness, purulent or "
    "foul discharge, wound gaping, bleeding, calf pain or swelling, chest pain or breathlessness, persistent "
    "vomiting, inability to pass urine, no bowel movement** \u2014 follow-up appointment, and drain/catheter care if "
    "going home with one.",
    "**Medication reconciliation (a core pharmacist deliverable):** compare the pre-admission list with the discharge "
    "list; restart the drugs that were withheld (**metformin, ACE-I/ARB, anticoagulants \u2014 with a clear restart "
    "date and dose**); stop the drugs that were only for the operation (**prophylactic antibiotics, PPIs, "
    "anti-emetics**); check for **therapeutic duplication (paracetamol in combination analgesics), interactions "
    "(NSAID + anticoagulant, tramadol + SSRI \u2192 serotonin syndrome), renal-dose adjustments and allergies**; "
    "counsel on **opioid tapering, laxative co-prescription, the finite duration of the analgesic course, and safe "
    "disposal of leftover opioids**; provide a **VTE-prophylaxis continuation plan** and injection technique training "
    "for LMWH; **document the counselling**.",
    "**Also the pharmacist's OT-related duties:** maintaining the OT drug tray and crash cart (stock, expiry, "
    "temperature and cold-chain logs, seal checks), **narcotic/controlled-substance accounting and register "
    "(NDPS)**, sterile preparation and dilution of injectables, ensuring **all syringes and containers on the sterile "
    "field are labelled with drug, strength and concentration**, segregating **look-alike/sound-alike and high-alert "
    "drugs** (concentrated potassium chloride, insulin, heparin, neuromuscular blockers), reporting and analysing "
    "medication errors and **adverse drug reactions (pharmacovigilance)**, and in-service education on new drugs and "
    "devices.",
])
box("MEDICATION SAFETY IN THE OR \u2014 the pharmacist's exam angle",
    ["Extend the PDF's \u2018five rights\u2019 to the **modern eight rights**: right **patient, drug, dose, route, "
     "time**, plus right **documentation, reason/indication and response**; add the right to **refuse** and patient "
     "**education**.",
     "**Every medication and solution on the sterile field must be labelled immediately** with name, strength and "
     "concentration \u2014 **unlabelled medications are ALWAYS discarded**; **all vials, syringes and containers stay "
     "in the room until the end of the case**; medications and containers are shown to the **relief person** at "
     "hand-over; calculations are **independently double-checked by a second licensed professional**.",
     "**High-alert medications in the OR:** concentrated **potassium chloride** (never as a bolus), insulin, heparin, "
     "neuromuscular blockers (**\u2018paralysing agent \u2014 warning\u2019 auxiliary label mandatory**), "
     "concentrated adrenaline, local anaesthetics, oxytocin, and any **look-alike/sound-alike pair** "
     "(e.g. ephedrine/epinephrine, heparin/insulin).",
     "**Constraints/forcing functions** recommended to reduce error: **minimise the number of calculations required, "
     "stock ready-to-use dose-specific concentrations, standardise concentrations, use pre-filled syringes, "
     "bar-code verification, and regular competency assessment.**"],
    fill=SH_HY, border="D98BA0", title_color=C_ANS)



# ================================================================ PART D
h1("Part D  \u2014  OT Safety for the Medical Assistant / Pharmacist (Gap-Fill)")

h2("D1.  In-Service Education, Ergonomics & Fatigue \u2014 only the missing numbers")
bullets([
    "**Orientation vs in-service education vs continuing education:** *orientation* = initial familiarisation with "
    "policies, procedures and the physical environment (AORN recommends a **mentor / buddy system**); "
    "*in-service* = **ongoing, employer-provided, job-specific** education (perioperative departments hold it "
    "**monthly**); *continuing education* = professional development beyond the employer, usually credit-based.",
    "**Mandatory annual training topics:** fire safety and evacuation, **fire drills at least twice a year "
    "(unannounced)**, disaster drill (at least annually), infection control and hand hygiene, biomedical waste "
    "handling, blood-borne pathogen/needle-stick protocol, radiation safety, laser safety, electrical safety, "
    "hazard communication (SDS), CPR/BLS recertification (**every 2 years**), lifting/ergonomics, and "
    "**competency validation for every new device before use**.",
    "**NIOSH lifting limit = 51 lb (23 kg)** under ideal conditions; **recommended maximum manual patient lift = "
    "35 lb (16 kg)** \u2014 beyond that, use a mechanical aid or a lift team.",
    "**Fatigue:** AORN recommends **no more than 12 consecutive hours** of direct patient care and "
    "**\u2264 60 hours per week including on-call**, with **8 hours of uninterrupted sleep**. Being awake "
    "**17\u201319 hours** impairs performance about as much as a blood alcohol of 0.05 %, and 24 h \u2248 0.10 %. "
    "**Working when judgement is impaired is an ethical breach, not just a personal risk.**",
    "**Other occupational hazards worth a line:** standing for long periods (varicose veins \u2014 use "
    "anti-fatigue mats and compression stockings), repetitive strain and awkward posture during laparoscopy and "
    "microsurgery, **noise > 85 dB** from saws/drills/suction (hearing protection), sharps injury, latex, "
    "psychological stress and burnout, and **second-victim** support after an adverse event.",
])

h2("D2.  Radiation Safety \u2014 the quantitative detail your PDF lacks")
h3("Units")
table(["Quantity", "SI unit", "Old unit", "Conversion / meaning"],
      [["**Exposure**", "Coulomb/kg", "Roentgen (R)", "Ionisation produced in air"],
       ["**Absorbed dose**", "**Gray (Gy) = 1 J/kg**", "rad", "**1 Gy = 100 rad** \u2014 energy deposited in tissue"],
       ["**Equivalent dose** (absorbed dose \u00d7 radiation weighting factor)", "**Sievert (Sv)**", "rem",
        "**1 Sv = 100 rem; 1 mSv = 100 mrem** \u2014 the unit used for **all dose limits and personnel monitoring**"],
       ["**Effective dose**", "Sievert (Sv)", "rem", "Equivalent dose weighted for tissue sensitivity \u2014 "
                                                    "whole-body risk"],
       ["**Activity**", "**Becquerel (Bq) = 1 disintegration/s**", "Curie (Ci)", "1 Ci = 3.7 \u00d7 10\u00b9\u2070 Bq"]],
      widths=[5.0, 4.4, 2.6, 7.0])
h3("Dose limits (ICRP / AERB \u2014 India)")
table(["Category", "Limit"],
      [["**Occupational \u2014 whole body**", "**20 mSv per year averaged over 5 consecutive years "
                                             "(100 mSv/5 y), with a maximum of 50 mSv in any single year**"],
       ["Occupational \u2014 lens of the eye", "**20 mSv/year** (ICRP 2011 revision; formerly 150 mSv)"],
       ["Occupational \u2014 skin, hands and feet (extremities)", "**500 mSv/year**"],
       ["**Pregnant worker (after declaration of pregnancy)**", "**\u2264 1 mSv to the fetus/abdominal surface for the "
                                                               "remainder of the pregnancy** (US NCRP: "
                                                               "**0.5 mSv per month**)"],
       ["**Members of the public**", "**1 mSv per year**"],
       ["Trainees / apprentices aged 16\u201318", "6 mSv per year"]],
      widths=[6.4, 12.6])
h3("Protection \u2014 the three cardinal principles + practical rules")
table(["Principle", "Application"],
      [["**1. Time**", "Minimise fluoroscopy screening time; use **pulsed fluoroscopy, last-image hold and "
                       "collimation**; the operator, not the assistant, should control the pedal; record total "
                       "screening time on the perioperative record"],
       ["**2. Distance \u2014 the INVERSE SQUARE LAW**", "**Intensity \u221d 1/(distance)\u00b2.** "
                                                        "**Doubling the distance reduces exposure to one quarter "
                                                        "(25 %); tripling it to one ninth.** Practical rule: stand "
                                                        "**at least 6 feet (2 m)** from the source, or leave the room. "
                                                        "In lateral fluoroscopy stand on the **image-intensifier "
                                                        "side, away from the X-ray tube**, because scatter is highest "
                                                        "on the tube side"],
       ["**3. Shielding**", "**Lead apron 0.25\u20130.5 mm lead equivalence** (0.5 mm attenuates \u2248 95\u201399 % "
                           "of the primary beam at 100 kVp); **thyroid collar, lead glasses (0.75 mm Pb "
                           "equivalent lenses), lead gloves for holding cassettes, mobile lead screens and "
                           "ceiling-suspended shields**. Aprons must be **hung, never folded**, and "
                           "**radiographically screened annually for cracks**. Wrap-around aprons for procedures "
                           "where you turn your back to the source"],
       ["**ALARA**", "**A**s **L**ow **A**s **R**easonably **A**chievable \u2014 the governing philosophy; "
                     "**justification, optimisation and dose limitation** (the three ICRP pillars)"],
       ["**Monitoring**", "**Personal dosimeter worn at chest/collar level OUTSIDE the lead apron** (a second one "
                          "inside for pregnant staff at waist level). Types: **TLD (thermoluminescent, LiF \u2014 "
                          "most common, read monthly/quarterly), film badge, OSL, and a pocket dosimeter/electronic "
                          "personal dosimeter for real-time reading**. Records are kept for the worker's lifetime"],
       ["**Personnel rules**", "**All non-essential staff leave the room or stand behind a lead screen**; "
                               "**scrubbed staff wear a lead apron donned BEFORE scrubbing**; "
                               "**never hand-hold the cassette or the patient if it can be avoided \u2014 use "
                               "lead gloves and a holder if unavoidable**; announce \u2018X-ray!\u2019 before every "
                               "exposure; **pregnant staff should be excluded from fluoroscopic cases**"],
       ["**Patient rules**", "**Ask every woman of child-bearing age whether she is or could be pregnant** "
                             "(the \u201810-day rule\u2019 for non-urgent radiography); shield the "
                             "**gonads/reproductive organs and thyroid** when the beam cannot be collimated; "
                             "collimate to the area of interest; use the lowest diagnostic exposure; document "
                             "fluoroscopy time and dose"],
       ["**Radioactive implants (brachytherapy)**", "Handle sources with **long forceps, never with fingers**; "
                                                   "keep a lead pot/container in the room; limit staff time in the "
                                                   "room and rotate staff; restrict visitors and exclude pregnant "
                                                   "staff and children; survey the room and linen with a "
                                                   "**Geiger-M\u00fcller counter** after removal; specific disposal "
                                                   "and incident procedures for a lost source"]],
      widths=[3.6, 15.4])
h3("Biological effects \u2014 the distinction examiners test")
table(["Effect", "Nature", "Examples"],
      [["**Deterministic (non-stochastic)**", "Has a **threshold dose**; severity **increases with dose**",
        "Skin erythema (\u2248 2 Gy) and desquamation, epilation, **cataract**, sterility, bone-marrow suppression, "
        "acute radiation syndrome, fetal malformation"],
       ["**Stochastic**", "**No threshold**; the **probability** (not the severity) increases with dose; assumed "
                          "**linear-no-threshold**", "**Carcinogenesis (leukaemia, thyroid, breast, lung) and "
                                                     "heritable genetic mutation**"]],
      widths=[4.0, 6.0, 9.0])
para("**Most radiosensitive tissues** (rapidly dividing): **bone marrow/lymphoid tissue, gonads (spermatogonia), "
     "intestinal crypt epithelium, skin basal layer, lens, fetus** \u2014 the **fetus is most sensitive during "
     "organogenesis, weeks 2\u20138**. **Least sensitive:** nerve, muscle and bone in the adult.")

h2("D3.  Infection Control \u2014 precautions, exposure management, and waste")
h3("Standard vs transmission-based precautions")
table(["Category", "Applies to", "Requirements", "Examples"],
      [["**Standard Precautions** (CDC 1996; incorporate the earlier **Universal Precautions** and "
        "**Body Substance Isolation**)",
        "**EVERY patient, every time, regardless of diagnosis** \u2014 all blood, all body fluids, secretions and "
        "excretions **except sweat**, non-intact skin and mucous membranes",
        "**Hand hygiene \u2022 gloves \u2022 gown \u2022 mask/eye protection or face shield when splash is likely "
        "\u2022 safe injection practices and sharps handling \u2022 respiratory hygiene/cough etiquette \u2022 "
        "safe handling of contaminated equipment, linen and waste \u2022 environmental cleaning \u2022 "
        "patient placement**",
        "Universal application in the OR"],
       ["**Contact Precautions**", "Direct or indirect contact transmission",
        "Standard + **gown and gloves for all contact**, dedicated equipment, single room preferred",
        "**MRSA, VRE, ESBL/CRE, C. difficile (soap and water \u2014 alcohol does not kill spores; hypochlorite for "
        "the environment), scabies, RSV, hepatitis A**"],
       ["**Droplet Precautions**", "Large droplets (> 5 \u00b5m) that travel **\u2248 3\u20136 feet (1\u20132 m)**",
        "Standard + **surgical mask within 1\u20132 m**, single room, mask on the patient during transport",
        "**Influenza, pertussis, mumps, rubella, Neisseria meningitidis, diphtheria, group A streptococcus, "
        "adenovirus**"],
       ["**Airborne Precautions**", "Small droplet nuclei (\u2264 5 \u00b5m) that remain suspended and travel long "
                                    "distances",
        "Standard + **N95 (or higher) respirator, fit-tested**, and an **Airborne Infection Isolation Room \u2014 "
        "NEGATIVE pressure with \u2265 12 air changes/hour, exhaust to outside or HEPA-filtered**, door kept closed",
        "**Pulmonary/laryngeal tuberculosis, measles (rubeola), varicella/disseminated zoster, SARS-CoV-2 during "
        "aerosol-generating procedures**"]],
      widths=[4.0, 4.0, 6.0, 5.0])
h3("PPE sequence \u2014 rote-learn the order")
table(["Donning (put on)", "Doffing (take off)"],
      [["**1** Hand hygiene \u2192 **2** **Gown** \u2192 **3** **Mask or respirator** (fit-check the N95) \u2192 "
        "**4** **Goggles / face shield** \u2192 **5** **Gloves**",
        "**1** **Gloves** \u2192 **2** **Goggles/face shield** \u2192 **3** **Gown** \u2192 **4** **Mask/respirator** "
        "(remove **outside** the room; touch only the ties/elastic, never the front) \u2192 **5** **Hand hygiene**"]],
      widths=[9.5, 9.5])
para("**The most contaminated items come off first (gloves), and the respirator comes off last** because it protects "
     "you until you have left the room. **Hand hygiene both before donning and immediately after doffing**, and "
     "**at any point where PPE is breached.**")
h3("Sharps injury and post-exposure management")
table(["Item", "Detail"],
      [["Transmission risk from a single percutaneous needle-stick from an infected source",
        "**Hepatitis B \u2248 6\u201330 %** (up to 30 % if HBeAg-positive; the **highest** risk) \u2022 "
        "**Hepatitis C \u2248 1.8 % (range 0\u20137 %)** \u2022 **HIV \u2248 0.3 % (1 in 300)**; "
        "**mucous-membrane exposure to HIV \u2248 0.09 %**"],
       ["**Immediate first aid**", "**Do NOT squeeze or suck the wound and do NOT use caustic agents "
                                   "(bleach, alcohol, iodine) inside the wound.** "
                                   "**Wash with soap and running water** (irrigate mucous membranes/eye copiously "
                                   "with water or saline for 10\u201315 min, remove contact lenses), allow it to bleed "
                                   "freely, dry and cover"],
       ["Then", "**Report immediately** to the designated officer/occupational health \u2192 assess the exposure "
                "(device, depth, volume, source status) \u2192 **baseline serology of the exposed person and, with "
                "consent, the source** \u2192 **start PEP** \u2192 counselling \u2192 follow-up serology at "
                "**6 weeks, 3 months and 6 months** \u2192 **document and file an incident report**"],
       ["**HIV PEP**", "Start **as early as possible \u2014 ideally within 2 hours, and no later than 72 hours**; "
                       "**3-drug regimen (e.g. tenofovir + lamivudine/emtricitabine + dolutegravir) for 28 days**; "
                       "the exposed person should use barrier contraception and not donate blood during follow-up"],
       ["**Hepatitis B PEP**", "If the exposed person is **unvaccinated or a non-responder** \u2192 "
                              "**hepatitis B immunoglobulin (HBIG) 0.06 mL/kg IM within 24 hours (up to 7 days) "
                              "PLUS start the vaccine series**. A **vaccinated responder (anti-HBs \u2265 10 mIU/mL) "
                              "needs no prophylaxis.** **Hepatitis B vaccination is mandatory for all "
                              "health-care workers \u2014 3 doses at 0, 1 and 6 months, with anti-HBs titre checked "
                              "1\u20132 months after the last dose**"],
       ["**Hepatitis C**", "**No vaccine and no PEP** \u2014 monitor with HCV RNA at 3\u20136 weeks and antibody at "
                          "4\u20136 months; treat early infection with direct-acting antivirals"],
       ["Prevention (hierarchy of controls)", "**Elimination/substitution** (needleless systems) \u2192 "
                                              "**engineering controls** (retractable and sheathed needles, blunt "
                                              "suture needles, **puncture-proof sharps container filled to only "
                                              "\u00be**, scalpel-blade removers) \u2192 **administrative/work-practice "
                                              "controls** (**never recap by hand \u2014 if unavoidable, one-handed "
                                              "\u2018scoop\u2019 technique**; never pass a sharp hand-to-hand \u2014 "
                                              "use a **neutral zone / hands-free \u2018no-touch\u2019 technique with "
                                              "a magnetic pad or basin**; mount the blade with a needle holder; "
                                              "announce sharps) \u2192 **PPE** (**double gloving reduces inner-glove "
                                              "perforation by \u2248 70\u201380 %**, blunt needles, face shield)"]],
      widths=[4.4, 14.6])
h3("Biomedical waste segregation \u2014 Biomedical Waste Management Rules 2016 (India)")
table(["Colour", "Category", "Contents", "Treatment / disposal"],
      [["**YELLOW**", "Human & animal anatomical waste; soiled waste; expired/discarded medicines; chemical waste; "
                      "chemical liquid waste; discarded linen contaminated with blood/body fluids; microbiology, "
                      "biotechnology and other clinical laboratory waste",
        "Body parts, placenta, tissue, organs, dressings, cotton, blood-soaked items, plaster casts, expired drugs, "
        "**cytotoxic drugs and vials (in a separate yellow \u2018cytotoxic\u2019 bag/container)**",
        "**Incineration / plasma pyrolysis / deep burial** (deep burial only in towns < 500 000 population). "
        "**Cytotoxic waste \u2192 incineration at \u2265 1200 \u00b0C** or return to the manufacturer"],
       ["**RED**", "Contaminated **recyclable** waste (plastics)",
        "IV tubing and sets, catheters, urine bags, syringes **without needles**, vacutainers without needles, "
        "gloves",
        "**Autoclaving / microwaving / hydroclaving, then shredding \u2192 sent to a registered recycler.** "
        "**No chemical pre-treatment before recycling; no landfilling**"],
       ["**WHITE (translucent, puncture-proof, leak-proof)**", "**Waste SHARPS including metals**",
        "Needles, syringes **with fixed needles**, needles from needle-tip cutters, scalpels, blades, "
        "any contaminated sharp object that may cause puncture",
        "**Autoclaving/dry-heat sterilisation followed by shredding or mutilation; final disposal by encapsulation "
        "in a concrete/metal sharps pit, or sent to an iron foundry / sanitary landfill**"],
       ["**BLUE (cardboard box with blue marking / puncture-proof container)**",
        "**Glassware and metallic body implants**",
        "Broken or discarded and contaminated glass including medicine vials and ampoules (**except those "
        "contaminated with cytotoxic drugs \u2014 those go in yellow**), metallic implants",
        "**Disinfection (1 % hypochlorite) or autoclaving/microwaving/hydroclaving, then sent for recycling**"]],
      widths=[3.0, 4.0, 6.0, 6.0])
bullets([
    "**Golden rule: segregate at the point of generation, by the person generating the waste \u2014 never sort waste "
    "later.** Bags are filled to **\u00be capacity**, tied (never stapled), labelled with the "
    "**biohazard/cytotoxic symbol** and the department/date, and stored no longer than **48 hours**.",
    "**Never recap, bend or break needles by hand; use a needle-tip cutter or destroyer at the point of use "
    "(the needle goes to the WHITE container).**",
    "**Liquid waste** (laboratory, washing, cleaning): pre-treat with disinfectant, then discharge to a sewer/ETP as "
    "per the Effluent standards.",
    "**Mercury** (broken thermometer/sphygmomanometer): **do not vacuum**; collect droplets with a syringe/adhesive "
    "tape or a mercury spill kit, place in a sealed container of water, dispose as hazardous waste \u2014 mercury is a "
    "neurotoxin. Mercury devices should be phased out.",
    "**Records, training and reporting:** annual report to the Pollution Control Board, maintenance of a waste "
    "register, **annual health check-up and immunisation (hepatitis B, tetanus) of waste handlers**, provision of PPE, "
    "and reporting of any **major accident** within 24 h.",
    "**Radioactive waste** is governed separately by the **Atomic Energy Regulatory Board (AERB)** \u2014 decay-in-"
    "storage in a shielded room until background level, then normal disposal.",
])

h2("D4.  Chemical & Physical Hazards \u2014 exposure limits and control")
h3("Occupational exposure limits worth memorising")
table(["Agent", "Limit", "Health effects & control"],
      [["**Nitrous oxide**", "**NIOSH REL 25 ppm (TWA over the procedure)**",
        "Bone-marrow/megaloblastic change (methionine synthase inhibition), neuropathy, **\u2191 spontaneous "
        "abortion and reduced fertility**, addiction. Control: **scavenging system, no leaks, tight-fitting masks, "
        "\u2265 15\u201325 air changes/h**"],
       ["**Halogenated volatile agents** (halothane, isoflurane, sevoflurane)", "**NIOSH REL 2 ppm** "
                                                                              "(0.5 ppm if used with N\u2082O)",
        "Headache, fatigue, impaired **mental performance, audiovisual ability and manual dexterity**; hepatotoxicity; "
        "teratogenic concern"],
       ["**Ethylene oxide**", "**OSHA PEL 1 ppm (8-h TWA)**; excursion **5 ppm/15 min**; action level 0.5 ppm",
        "Mucosal irritation, **mutagen, carcinogen (leukaemia/lymphoma)**, reproductive toxicity, "
        "flammable/explosive. Control: dedicated ventilated room, gas monitors and alarms, aeration cabinet, "
        "personal badges"],
       ["**Formaldehyde**", "**OSHA PEL 0.75 ppm TWA; STEL 2 ppm**", "Eye/airway irritation, asthma, dermatitis, "
                                                                    "**Group 1 human carcinogen (nasopharyngeal)** "
                                                                    "\u2014 this is why routine OT fumigation is no "
                                                                    "longer advised. Handle specimens in a "
                                                                    "**ventilated hood** with gloves and eye "
                                                                    "protection"],
       ["**Glutaraldehyde**", "**ACGIH ceiling 0.05 ppm** (no safe TWA)", "**Occupational asthma, rhinitis, "
                                                                         "conjunctivitis, contact dermatitis, "
                                                                         "epistaxis, headache**. Control: "
                                                                         "**covered/closed soaking basin with local "
                                                                         "exhaust, nitrile or butyl gloves (NOT latex "
                                                                         "\u2014 it permeates), goggles, apron; "
                                                                         "thorough rinsing of instruments**"],
       ["**Methyl methacrylate** (bone cement)", "**OSHA PEL 100 ppm**", "Pungent, irritant, dermatitis, headache; "
                                                                        "**cement implantation syndrome in the "
                                                                        "patient \u2014 hypotension, hypoxia, "
                                                                        "arrhythmia and cardiac arrest during "
                                                                        "cementing**. Control: **mix in a closed "
                                                                        "vacuum system with local exhaust**, nitrile "
                                                                        "gloves, no contact with soft contact lenses "
                                                                        "(absorbs vapour)"],
       ["**Peracetic acid / hydrogen peroxide**", "Manufacturer/ACGIH limits", "Irritant, corrosive at high "
                                                                              "concentration; largely safe "
                                                                              "breakdown products"],
       ["**Surgical smoke**", "No formal PEL; treated as an aerosol hazard", "See \u00a7B8 \u2014 smoke evacuator "
                                                                            "with ULPA filter + high-filtration "
                                                                            "mask/N95"],
       ["**Noise**", "OSHA action level **85 dB (8-h TWA)**", "Saws, drills, suction, alarms, music \u2014 "
                                                             "hearing protection and noise reduction; also impairs "
                                                             "team communication"],
       ["**Compressed gases**", "\u2014", "**Secure every cylinder upright in a stand or chained; never drop a "
                                          "cylinder or let one strike another; transport with the valve cap on; "
                                          "never oil/grease an oxygen valve; store away from heat; segregate full and "
                                          "empty cylinders.** A cylinder that falls can become a missile"]],
      widths=[3.6, 3.8, 11.6])
h3("Hazard communication, cytotoxics and spills")
bullets([
    "**Hazard Communication Standard / \u2018Right to Know\u2019:** every hazardous chemical must have a "
    "**Safety Data Sheet (SDS \u2014 formerly MSDS), in 16 standardised sections**, accessible to every employee at "
    "all times, plus **GHS labelling** with the signal words \u2018**Danger**\u2019 or \u2018**Warning**\u2019 and "
    "the nine pictograms (flame, flame over circle/oxidiser, exploding bomb, corrosion, gas cylinder, skull and "
    "crossbones, exclamation mark, health hazard, environment).",
    "**Hierarchy of controls (in order):** **elimination \u2192 substitution \u2192 engineering controls "
    "(ventilation, closed systems, scavenging) \u2192 administrative/work-practice controls \u2192 PPE (the LAST and "
    "weakest line of defence)** \u2014 a very commonly asked ranking.",
    "**Cytotoxic / hazardous drug handling (USP <800> principles):** prepare only in a **Class II Type B biological "
    "safety cabinet or a compounding aseptic containment isolator, inside a negatively-pressured, externally-vented "
    "room**; use **closed-system transfer devices** and Luer-lock syringes; PPE = **two pairs of chemotherapy-tested "
    "gloves (changed every 30\u201360 min), an impermeable back-closing gown, eye protection and an N95/respirator**; "
    "**never crush or split cytotoxic tablets**; label with a **\u2018Cytotoxic \u2014 handle with care\u2019** "
    "warning; **pregnant, breast-feeding and trying-to-conceive staff should not handle them**; dispose in the "
    "**yellow cytotoxic stream**; keep a **spill kit** and document exposure.",
    "**Chemical spill response \u2014 evacuate and restrict access \u2192 identify the agent from the SDS \u2192 "
    "don appropriate PPE \u2192 contain and absorb from the outside inwards with the spill-kit absorbent \u2192 "
    "neutralise/decontaminate \u2192 place waste in a labelled hazardous-waste bag \u2192 ventilate \u2192 report and "
    "document.** For a **skin/eye splash: irrigate at the eyewash station or shower for a minimum of 15 minutes** "
    "and seek occupational-health review.",
])



# ================================================================ PART E
h1("Part E  \u2014  The Numbers Sheet (one page, learn cold)")
table(["Topic", "Number to remember"],
      [["Fasting", "**2 h clear fluids \u2022 4 h breast milk \u2022 6 h formula/light meal \u2022 8 h fatty meal**"],
       ["Aspiration risk", "Gastric volume **> 25 mL (0.4 mL/kg)** and pH **< 2.5**"],
       ["Antibiotic prophylaxis", "**Within 60 min of incision** (120 min for vancomycin); redose after **2 half-lives** "
                                  "or **> 1500 mL** blood loss; **stop within 24 h**"],
       ["Patient identifiers", "**At least 2**; WHO checklist = **3 phases, 19 items**"],
       ["Safety strap", "Across mid-thighs, **\u2248 3 inches (7.5 cm) above the knees**"],
       ["Arm abduction", "**\u2264 90\u00b0**; palms up; Trendelenburg tilt **30\u201345\u00b0**"],
       ["Sterile field margins", "Wrapper edge unsterile = **2.5 cm (1 inch)**; unsterile person keeps "
                                 "**30 cm (1 ft)** away; gown sterile **chest to field level**, sleeves to **5 cm "
                                 "above the elbow**"],
       ["Surgical scrub", "First scrub of the day **3\u20135 min**"],
       ["OR environment", "**Temp 20\u201323 \u00b0C \u2022 RH 20\u201360 % \u2022 \u2265 15 (ideally 20\u201325) air "
                          "changes/h \u2022 POSITIVE pressure \u2022 HEPA 99.97 % at 0.3 \u00b5m \u2022 laminar flow "
                          "300\u2013600 ACH**"],
       ["Autoclave", "**121 \u00b0C / 15 psi / 15\u201320 min**  \u2022  **134 \u00b0C / 30 psi / 3\u20133.5 min** "
                     "\u2022 prevacuum **132\u2013135 \u00b0C / 3\u20134 min** \u2022 flash/IUSS **132 \u00b0C / 3 min "
                     "unwrapped (10 min lumened)** \u2022 prions **134 \u00b0C / 18 min**"],
       ["Hot air oven", "**160 \u00b0C / 2 h \u2022 170 \u00b0C / 1 h \u2022 180 \u00b0C / 30 min**"],
       ["Bowie-Dick", "**134 \u00b0C for 3.5 min in an EMPTY prevacuum chamber, daily** \u2014 tests **air removal**"],
       ["Ethylene oxide", "**450\u20131200 mg/L \u2022 37\u201363 \u00b0C \u2022 RH 40\u201380 % \u2022 1\u20136 h**; "
                          "aeration **8\u201312 h at 50\u201360 \u00b0C** or **7 days** at room temperature; "
                          "**PEL 1 ppm**"],
       ["H\u2082O\u2082 plasma", "**58 % H\u2082O\u2082, 45\u201350 \u00b0C, \u2248 1 h**; **no cellulose, no "
                                "liquids**"],
       ["Glutaraldehyde 2 %", "**Sterilization 10 h \u2022 HLD 20\u201345 min** (textbook: 10 min); ceiling "
                              "**0.05 ppm**"],
       ["Other chemicals", "**OPA 0.55 % \u2014 12 min \u2022 peracetic acid \u2014 12 min at 50\u201356 \u00b0C \u2022 "
                           "H\u2082O\u2082 7.5 % \u2014 30 min HLD**"],
       ["Sterility", "**SAL = 10\u207b\u2076** \u2022 \u03b3-radiation dose **25 kGy** \u2022 filter "
                     "**0.22 \u00b5m**"],
       ["Biological indicators", "**Steam & plasma \u2192 Geobacillus stearothermophilus \u2022 dry heat & EO \u2192 "
                                 "Bacillus atrophaeus \u2022 radiation \u2192 Bacillus pumilus**"],
       ["Electrosurgery", "**300 kHz \u2013 3 MHz**; nerve stimulation avoided **above ~100 kHz**; "
                          "ultrasonic scalpel **55.5 kHz**; LigaSure seals up to **7 mm** vessels"],
       ["Surgical smoke", "1 g of tissue \u2248 **3\u20136 cigarettes**; evacuator within **2.5\u20135 cm "
                          "(1\u20132 in)**; **ULPA 99.999 % at 0.1 \u00b5m**"],
       ["Tourniquet", "**Upper limb 250\u2013300 mmHg, \u2264 60 min \u2022 lower limb 300\u2013350 mmHg, "
                      "\u2264 90\u2013120 min**; overlap **3\u20136 inches**; deflate **10\u201315 min** if longer"],
       ["CPR", "Rate **100\u2013120/min**, depth **5\u20136 cm**, ratio **30:2**, biphasic shock **120\u2013200 J**, "
               "adrenaline **1 mg q3\u20135 min**, amiodarone **300 mg**"],
       ["Anaphylaxis", "**Adrenaline 0.5 mg IM (0.5 mL of 1:1000)**, repeat q5 min; IV **50 \u00b5g** boluses; "
                       "fluid **20 mL/kg**"],
       ["Malignant hyperthermia", "**Dantrolene 2.5 mg/kg IV, up to 10 (\u201330) mg/kg**; O\u2082 at **10 L/min**; "
                                  "cool to **38 \u00b0C**; urine output **> 1\u20132 mL/kg/h**; gene **RYR1, "
                                  "chromosome 19**"],
       ["LAST", "**20 % lipid emulsion 1.5 mL/kg bolus, then 0.25 mL/kg/min**"],
       ["Aldrete", "**5 parameters \u00d7 0\u20132 = max 10; discharge \u2265 9**; PADSS also **\u2265 9**"],
       ["Apfel PONV", "**4 factors** \u2192 10 / 20 / 40 / 60 / **80 %**"],
       ["Post-op fever", "**Wind 1\u20132 \u2022 Water 3\u20135 \u2022 Walking 4\u20136 \u2022 Wound 5\u20137 \u2022 "
                         "Wonder drugs > 7 days**"],
       ["Wound healing", "Inflammatory **0\u20134 d** \u2022 proliferative **4\u201321 d** \u2022 remodelling "
                         "**21 d\u20131 y**; tensile strength **5\u201310 % at 1 wk, 30\u201350 % at 1 month, "
                         "max 70\u201380 %**"],
       ["Dehiscence", "**POD 5\u201310**; herald sign = **serosanguineous \u2018salmon-pink\u2019 discharge**"],
       ["Suture removal", "Face **3\u20135 d** \u2022 scalp/trunk **7\u201310 d** \u2022 limbs **10\u201314 d** \u2022 "
                          "joints/palms/soles **14 d**"],
       ["Wound class SSI risk", "**Clean < 2 % \u2022 clean-contaminated 3\u201310 % \u2022 contaminated "
                                "15\u201320 % \u2022 dirty 30\u201340 %**"],
       ["SSI definition window", "**30 days** (**90 days** if an implant is present)"],
       ["Ileus recovery", "**Small bowel 24 h \u2192 stomach 24\u201348 h \u2192 colon 48\u201372 h**"],
       ["Urine output", "**> 0.5 mL/kg/h (\u2248 30 mL/h)**; maintenance fluid **4-2-1 rule / 25\u201330 mL/kg/day**"],
       ["Radiation", "**1 Gy = 100 rad \u2022 1 Sv = 100 rem**; occupational **20 mSv/y (max 50 in one year)**; "
                     "public **1 mSv/y**; eye **20 mSv/y**; extremities **500 mSv/y**; fetus **1 mSv** "
                     "(0.5 mSv/month); apron **0.25\u20130.5 mm Pb**; stand **\u2265 2 m (6 ft)**; doubling distance "
                     "\u2192 **\u00bc** exposure"],
       ["Needle-stick risk", "**HBV 6\u201330 % \u2022 HCV 1.8 % \u2022 HIV 0.3 %**; HIV PEP within **2 h "
                             "(max 72 h) for 28 days**; HBIG **0.06 mL/kg within 24 h**"],
       ["Exposure limits", "**EO 1 ppm \u2022 formaldehyde 0.75 ppm \u2022 glutaraldehyde 0.05 ppm ceiling \u2022 "
                           "N\u2082O 25 ppm \u2022 volatiles 2 ppm \u2022 methyl methacrylate 100 ppm \u2022 "
                           "noise 85 dB**"],
       ["Lifting", "**NIOSH 51 lb (23 kg)**; patient handling **35 lb (16 kg)**"],
       ["Waste bins", "**Yellow** anatomical/soiled/expired drugs/cytotoxic \u2022 **Red** recyclable plastics \u2022 "
                      "**White** sharps \u2022 **Blue** glass & metallic implants; bags filled to **\u00be**; "
                      "cytotoxic incineration **\u2265 1200 \u00b0C**"],
       ["Fatigue", "**\u2264 12 consecutive hours** direct care; **\u2264 60 h/week** including call; "
                   "**8 h** uninterrupted sleep"],
       ["Drills", "Fire drill **\u2265 twice a year**; BLS recertification **every 2 years**"]],
      widths=[4.4, 14.6])

# ================================================================ PART F  MCQ
h1("Part F  \u2014  Exam-Style MCQ Bank with Explanations")
box("ABOUT THESE QUESTIONS",
    ["These **88 questions are modelled on the pattern and difficulty of previous-year questions** in Indian "
     "competitive and university papers that examine this unit (GPAT/NIPER-style pharmacy papers, nursing and "
     "OT-technology papers, and PG entrance patterns). They are **written for this document from the standard sources "
     "listed on page 1** \u2014 they are **not** verbatim reproductions of any single paper, so **use them as "
     "coverage checks, not as a leaked question bank.**",
     "Deliberate weighting, matched to where the marks fall: **post-operative care 18 \u2022 sterilization & "
     "counting 15 \u2022 asepsis/environment/SSI 11 \u2022 electrosurgery & equipment 9 \u2022 positioning 9 \u2022 "
     "emergencies 9 \u2022 radiation/chemical/waste 8 \u2022 identification & records 5 \u2022 pre-op & consent 4.**",
     "**Every question is answered with the reasoning and the distractor logic**, so a wrong attempt still teaches "
     "you the three facts around it."],
    fill=SH_GAP, border="9B6FC4", title_color=C_H3)

h2("Section 1 \u2014 Pre-Operative Preparation and Consent")
mcq("As per ASA guidelines, the minimum fasting period for clear liquids before elective surgery in a healthy adult is",
    ["1 hour", "2 hours", "4 hours", "6 hours"], 1,
    "Remember **2-4-6-8**: clear fluids 2 h, breast milk 4 h, formula/light meal 6 h, fatty meal 8 h. Prolonged "
    "fasting is harmful and does not reduce gastric volume.")
mcq("The preferred method of hair removal before surgery is",
    ["Shaving with a razor the night before", "Shaving with a razor in the OR",
     "Electric clippers immediately before surgery", "No removal is ever permitted"], 2,
    "Ideally hair is **not removed at all**; if it interferes with the procedure, use **clippers immediately before "
    "surgery, outside the OR**. Razor shaving \u2014 especially the night before \u2014 causes micro-abrasions and "
    "carries the **highest** SSI risk.")
mcq("Consent for a surgical procedure on a 10-year-old child must be obtained from the guardian under which section "
    "of the IPC?",
    ["Section 87", "Section 88", "Section 89", "Section 92"], 2,
    "**\u00a7 89** covers acts done in good faith for the benefit of a **child under 12 years** or a person of "
    "unsound mind, by or with the consent of the guardian. \u00a7 87\u201388 relate to consent by persons above 18; "
    "**\u00a7 92** covers acts done in good faith **without consent** in an emergency.")
mcq("A patient with well-controlled hypertension and no functional limitation is classified as",
    ["ASA I", "ASA II", "ASA III", "ASA IV"], 1,
    "**ASA II** = mild systemic disease **without** functional limitation. Pregnancy, current smoking and BMI "
    "30\u201340 are also ASA II. Functional limitation moves the patient to ASA III.")

h2("Section 2 \u2014 Patient Identification, Safety Checklist and Records")
mcq("The WHO Surgical Safety Checklist consists of how many phases?",
    ["Two", "Three", "Four", "Five"], 1,
    "**Three: Sign In** (before induction), **Time Out** (after induction, before skin incision) and **Sign Out** "
    "(before the patient leaves the OR) \u2014 19 items in total.")
mcq("The surgical \u201ctime out\u201d must be performed",
    ["On admission to the ward", "In the holding area before transfer",
     "After induction of anaesthesia but before skin incision", "Immediately after wound closure"], 2,
    "The time out is the whole-team verbal pause **immediately before incision** (before the scalpel is passed). "
    "The **Universal Protocol** was mandated by the Joint Commission in **July 2004**; AORN added confirmation of "
    "**correct position** in March 2005.")
mcq("Which of the following is confirmed during the \u201cSign Out\u201d phase of the WHO checklist?",
    ["Pulse oximeter functioning", "Site marking", "Correctness of instrument, sponge and needle counts",
     "Anticipated blood loss"], 2,
    "Sign Out covers the **name of the procedure recorded, correctness of the counts, correct labelling of specimens, "
    "equipment problems, and key concerns for recovery**. Pulse oximetry and site marking belong to Sign In; "
    "anticipated blood loss is discussed at Sign In/Time Out.")
mcq("A specimen sent for frozen section should be transported",
    ["In 10 % formalin", "Fresh, in saline-moistened gauze", "In absolute alcohol", "In glutaraldehyde"], 1,
    "**Frozen sections must be sent FRESH \u2014 never in formalin**, which precludes further processing and "
    "immunohistochemistry. Routine histopathology goes into **10 % neutral buffered formalin at about 10\u00d7 the "
    "specimen volume**; specimens for culture go into a sterile container with no fixative.")
mcq("The legal responsibility for recording the surgical count rests with the",
    ["Surgeon", "Scrub person", "Circulator", "Anaesthesia provider"], 2,
    "Both the scrub person and the circulator **count aloud together**, but keeping the **written record of the count "
    "is the legal responsibility of the circulator**, who also initials any additions.")

h2("Section 3 \u2014 Positioning and Transfer")
mcq("The single most important measure to prevent a patient falling from the operating table is",
    ["Keeping the side rails up", "A safety (restraint) strap firmly applied across the mid-thighs",
     "Padded armboards", "Two staff at the bedside"], 1,
    "The safety strap across the **mid-thighs, about 3 inches (7.5 cm) above the knees**, applied firmly but not so "
    "tightly as to impair venous return, is the key measure. The patient is also **never left unattended** in the OR.")
mcq("Which nerve is most commonly injured in the lithotomy position, resulting in foot drop?",
    ["Femoral nerve", "Sciatic nerve", "Common peroneal nerve", "Obturator nerve"], 2,
    "The **common peroneal nerve** is compressed against the **fibular head** by the stirrup post \u2192 foot drop and "
    "loss of sensation over the dorsum of the foot. Adequate padding at the fibular neck prevents it.")
mcq("The commonest peripheral nerve injury associated with the supine position is that of the",
    ["Radial nerve", "Ulnar nerve", "Median nerve", "Axillary nerve"], 1,
    "**Ulnar neuropathy** is the commonest peri-operative nerve injury overall. Prevention: arms abducted "
    "**\u2264 90\u00b0**, elbows padded, **palms supinated (up)**, no pressure at the medial epicondyle.")
mcq("In the Trendelenburg position, all of the following occur EXCEPT",
    ["Increased intracranial pressure", "Increased venous return", "Increased functional residual capacity",
     "Cephalad displacement of the diaphragm"], 2,
    "Trendelenburg **decreases** FRC and compliance because the abdominal contents push the diaphragm cephalad "
    "\u2192 atelectasis. It increases venous return, CVP, ICP and IOP, and may cause facial/laryngeal oedema and "
    "tube migration.")
mcq("Reverse Trendelenburg position is typically used for",
    ["Abdominal hysterectomy", "Thyroidectomy", "Haemorrhoidectomy", "Nephrectomy"], 1,
    "Reverse Trendelenburg (head up) is used for **neck procedures \u2014 thyroidectomy, parathyroidectomy, scalene "
    "node biopsy \u2014 and laparoscopic cholecystectomy**. Trendelenburg is used for pelvic surgery such as "
    "abdominal hysterectomy; lithotomy/jackknife for anorectal surgery; lateral kidney for nephrectomy.")
mcq("The Kraske (jackknife) position is a modification of the",
    ["Supine position", "Prone position", "Lateral position", "Fowler's position"], 1,
    "**Kraske/jackknife is a prone modification** with the table flexed at the hips to expose the sacrococcygeal and "
    "anorectal area. Trendelenburg, reverse Trendelenburg, Fowler's and lithotomy are all modifications of "
    "**supine**.")
mcq("When taking a patient out of the lithotomy position, both legs must be lowered slowly and simultaneously "
    "chiefly to prevent",
    ["Nerve traction injury", "A sudden fall in blood pressure", "Wound dehiscence", "Deep vein thrombosis"], 1,
    "Lowering the legs abruptly shifts blood into the lower limbs \u2192 **sudden reduction in venous return and a "
    "precipitous drop in blood pressure**. Simultaneous, slow lowering by two persons also protects the spine and "
    "hips.")
mcq("In the lateral position, the axillary (chest) roll is placed",
    ["Directly in the axilla", "Just caudal to the axilla, under the upper chest wall", "Under the iliac crest",
     "Between the knees"], 1,
    "It must be placed **caudal to (below) the axilla** to lift the chest wall and relieve pressure on the "
    "**dependent brachial plexus and axillary vessels**. Placing it **in** the axilla causes the very injury it is "
    "meant to prevent.")
mcq("Post-operative visual loss due to ischaemic optic neuropathy is most associated with which position?",
    ["Supine", "Lithotomy", "Prone", "Sitting"], 2,
    "The **prone position** (especially long spinal surgery with hypotension, anaemia and venous congestion) carries "
    "the highest risk of **POVL**, along with corneal abrasion. The head must be kept neutral with the eyes free of "
    "all pressure.")

h2("Section 4 \u2014 Environment, Asepsis and Surgical Site Infection")
mcq("A HEPA filter used in operating-room air conditioning removes",
    ["99.97 % of particles \u2265 0.3 \u00b5m", "99.9 % of particles \u2265 1 \u00b5m",
     "95 % of particles \u2265 0.5 \u00b5m", "99.999 % of particles \u2265 0.1 \u00b5m"], 0,
    "**HEPA = 99.97 % of particles of 0.3 \u00b5m or larger.** The 99.999 % at 0.1 \u00b5m figure describes an "
    "**ULPA** filter, used in surgical-smoke evacuators.")
mcq("The operating room should be maintained at which pressure relative to the adjacent corridor?",
    ["Negative pressure", "Positive pressure", "Equal pressure", "Alternating pressure"], 1,
    "**Positive** pressure ensures air flows **out** of the OR, keeping contaminants out. **Negative** pressure "
    "rooms are used for **airborne isolation (tuberculosis, measles, varicella)** with \u2265 12 air changes/hour.")
mcq("The recommended relative humidity in an operating room is",
    ["Below 20 %", "20\u201360 %", "70\u201380 %", "Above 80 %"], 1,
    "**20\u201360 %** (classically 30\u201360 %). **Too low favours static electricity and sparks; too high favours "
    "microbial growth, condensation and pack strike-through.** Temperature is kept at 20\u201323 \u00b0C with "
    "\u2265 15 (ideally 20\u201325) air changes per hour.")
mcq("The single most effective measure for preventing healthcare-associated infection is",
    ["Prophylactic antibiotics", "Hand hygiene", "Ultraviolet irradiation of the OR", "Routine fumigation"], 1,
    "**Hand hygiene with an antimicrobial agent is the single greatest factor** in preventing nosocomial infection "
    "(stated explicitly in your PDF, p. 34). Routine formaldehyde fumigation is obsolete, and UV is only a surface/air "
    "adjunct.")
mcq("A draped instrument table is considered sterile",
    ["Down to 30 cm below the table top", "Only at and above the level of the table top",
     "Down to the floor", "Only in the central one third"], 1,
    "**Anything below the table-top level is unsterile** and is never brought back up. Related rules: the wrapper "
    "edge (**2.5 cm/1 inch**) is unsterile, the gown back is unsterile, and **moisture causes contamination by "
    "strike-through**.")
mcq("An elective cholecystectomy without spillage is classified as which type of wound?",
    ["Clean", "Clean-contaminated", "Contaminated", "Dirty"], 1,
    "Entry into the biliary/GI tract under controlled conditions without unusual contamination = "
    "**clean-contaminated (Class II), expected SSI 3\u201310 %**. Clean < 2 %, contaminated 15\u201320 %, "
    "dirty/infected 30\u201340 %.")
mcq("A surgical site infection occurring in a patient with an implant is defined by the CDC as an infection within",
    ["7 days", "30 days", "90 days", "1 year"], 2,
    "Superficial incisional SSI = within **30 days**; deep incisional and organ/space SSI = within **30 days, or "
    "90 days when an implant is in place**.")
mcq("Which organism is the commonest cause of surgical site infection?",
    ["Escherichia coli", "Pseudomonas aeruginosa", "Staphylococcus aureus", "Clostridium perfringens"], 2,
    "**Staphylococcus aureus** (including MRSA) is the commonest, followed by coagulase-negative staphylococci, "
    "Enterococcus and E. coli. The commonest **source** of contamination is the **patient's own endogenous flora**.")
mcq("Alcohol-based hand rub is INEFFECTIVE and soap and water must be used when dealing with",
    ["Methicillin-resistant Staphylococcus aureus", "Clostridioides difficile", "Influenza virus",
     "Pseudomonas aeruginosa"], 1,
    "**Alcohol is not sporicidal**, so **C. difficile spores require mechanical removal by soap and water** (plus "
    "contact precautions and hypochlorite for the environment). Hands that are **visibly soiled** also require soap "
    "and water.")
mcq("Artificial fingernails are prohibited for operating-room personnel because they",
    ["Tear surgical gloves", "Harbour Gram-negative bacteria and fungi", "Interfere with pulse oximetry",
     "Absorb chlorhexidine"], 1,
    "AORN prohibits artificial nails and chipped polish because the **subungual space harbours Gram-negative bacilli "
    "and fungi** and cannot be adequately decontaminated. (Nail polish does also interfere with pulse oximetry \u2014 "
    "but the infection-control reason is the one asked.)")
mcq("Which of the following correctly describes the sequence of DOFFING personal protective equipment?",
    ["Gown \u2192 gloves \u2192 mask \u2192 goggles", "Gloves \u2192 goggles \u2192 gown \u2192 mask",
     "Mask \u2192 gown \u2192 gloves \u2192 goggles", "Goggles \u2192 gloves \u2192 mask \u2192 gown"], 1,
    "**Doffing: gloves \u2192 goggles/face shield \u2192 gown \u2192 mask/respirator (outside the room) \u2192 hand "
    "hygiene.** The most contaminated item (gloves) comes off first and the respirator last. Donning is the reverse "
    "order: **gown \u2192 mask \u2192 goggles \u2192 gloves**.")

h2("Section 5 \u2014 Electrosurgery and Equipment")
mcq("Electrosurgical units operate in which frequency range?",
    ["50\u201360 Hz", "1\u201310 kHz", "300 kHz \u2013 3 MHz", "10\u2013100 MHz"], 2,
    "**Radiofrequency, 300 kHz\u20133 MHz.** Above about **100 kHz** the current no longer stimulates nerve and "
    "muscle (the **Faradic effect** is avoided), which is why RF current cuts and coagulates without electrocuting "
    "the patient.")
mcq("Which of the following is true of a bipolar electrosurgical unit?",
    ["It requires a dispersive return electrode pad", "Current passes through the whole patient",
     "No dispersive pad is required", "It cannot be used near a pacemaker"], 2,
    "In bipolar electrosurgery both electrodes are the **two tines of the forceps**, so current passes only through "
    "the tissue between them. **No pad is needed**, lateral spread is minimal, and it is the **safest mode in "
    "patients with pacemakers**.")
mcq("The coagulation waveform of an electrosurgical generator is characterised by",
    ["Continuous low-voltage current", "Interrupted high-voltage current", "Direct current",
     "Continuous high-frequency current at 50 Hz"], 1,
    "**Coag = interrupted, high-voltage, low duty cycle (\u2248 6 %)**; **cut = continuous, low-voltage, "
    "high-current (100 % duty cycle)**. The high voltage of coag mode is why it carries the greatest risk of "
    "**capacitive coupling and insulation breakdown**.")
mcq("Capacitive coupling during laparoscopic electrosurgery is MOST likely when",
    ["A bipolar forceps is used", "A metal instrument passes through a plastic (hybrid) trocar cannula",
     "The dispersive pad is placed on the thigh", "The generator is in cut mode at low power"], 1,
    "A **hybrid metal\u2013plastic trocar system** allows current to be induced through intact insulation into "
    "surrounding tissue. Use **all-metal or all-plastic** systems, lowest effective power, cut rather than coag mode, "
    "short activations, and **active-electrode monitoring (AEM)**.")
mcq("A pin-hole break in the insulation of a laparoscopic electrosurgical instrument classically presents as",
    ["An immediate visible spark", "A return-electrode burn", "Delayed bowel perforation several days later",
     "Failure of the generator to activate"], 2,
    "**Insulation failure** discharges current **outside the surgeon's field of view**, causing a full-thickness bowel "
    "burn that perforates **3\u201310 days post-operatively** with peritonitis. Inspect insulation before every use "
    "and use AEM instruments.")
mcq("A return-electrode contact quality monitoring (REM) system on a modern ESU prevents",
    ["Alternate-site burns from metal contact", "Pad-site burns from inadequate pad contact", "Capacitive coupling",
     "Surgical smoke exposure"], 1,
    "A **split/dual pad** continuously measures impedance between its two halves and **shuts the generator down if "
    "contact becomes inadequate**, virtually eliminating **return-electrode (pad-site) burns**.")
mcq("The surgeon repeatedly asks the circulator to increase the power on the electrosurgical unit. The correct action "
    "is to",
    ["Increase the power as requested", "Change to bipolar mode",
     "Stop, check all connections and the dispersive pad, and replace the unit if faulty",
     "Reposition the ECG electrodes"], 2,
    "Repeated requests for more power indicate a **fault \u2014 poor pad contact, a loose connection, a frayed cord "
    "or a failing generator**. Check the circuit; if no cause is found, **turn off and unplug the unit, label it, "
    "obtain another, and document the occurrence and the serial numbers** in the operative record.")
mcq("Which statement about surgical smoke (plume) is correct?",
    ["A standard surgical mask provides adequate protection",
     "Ablating 1 g of tissue is roughly equivalent to smoking 3\u20136 unfiltered cigarettes",
     "It is sterile and harmless", "A HEPA filter is required in the evacuator, not ULPA"], 1,
    "Surgical smoke contains over 150 chemicals plus viable cells and viral DNA (**HPV transmission to surgeons is "
    "documented**). **1 g of ablated tissue \u2248 3\u20136 unfiltered cigarettes.** Control = **smoke evacuator with "
    "an ULPA filter held within 1\u20132 inches**, plus a **high-filtration mask or N95** \u2014 a standard mask is "
    "not enough.")
mcq("For an average adult, the maximum recommended continuous pneumatic tourniquet time and pressure for the LOWER "
    "limb are approximately",
    ["30 min at 200 mmHg", "60 min at 250 mmHg", "90\u2013120 min at 300\u2013350 mmHg", "3 hours at 400 mmHg"], 2,
    "**Lower limb: 300\u2013350 mmHg for \u2264 90\u2013120 min; upper limb: 250\u2013300 mmHg for \u2264 60 min.** "
    "If more time is needed, **deflate for 10\u201315 minutes** and re-inflate. Record site, pressure and both "
    "inflation and deflation times.")

h2("Section 6 \u2014 Counting, Sterilization and Disinfection")
mcq("The commonest retained surgical item is",
    ["A needle", "A surgical sponge/gauze", "A retractor", "A guidewire"], 1,
    "**Sponges/gauze account for roughly half to two-thirds** of retained items, most often in the "
    "**abdomen/pelvis**. A retained gauze encased in granulation tissue is a **gossypiboma (textiloma)**, and is a "
    "classic **res ipsa loquitur** claim. Risk factors: emergency surgery, unexpected change of procedure and high "
    "BMI.")
mcq("Standard gravity-displacement autoclave conditions are",
    ["100 \u00b0C at atmospheric pressure for 30 min", "121 \u00b0C at 15 psi for 15\u201320 min",
     "134 \u00b0C at 15 psi for 30 min", "160 \u00b0C for 2 hours"], 1,
    "**121 \u00b0C (250 \u00b0F) at 15 psi for 15\u201320 minutes**, or **134 \u00b0C at 30 psi for 3\u20133.5 "
    "minutes**. 160 \u00b0C for 2 h is a **hot air oven (dry heat)** cycle.")
mcq("Flash (immediate-use) steam sterilization of an unwrapped metal instrument is carried out at",
    ["121 \u00b0C for 20 min", "132 \u00b0C for 3 min", "132 \u00b0C for 30 min", "160 \u00b0C for 1 hour"], 1,
    "**132 \u00b0C (270 \u00b0F) for 3 minutes** for unwrapped non-porous metal items; **10 minutes** for porous, "
    "lumened or complex items. IUSS is **only for an urgently needed single item** \u2014 never for implants, whole "
    "sets, or for convenience.")
mcq("The Bowie-Dick test is performed to detect",
    ["Sterilizer air leaks and inadequate air removal in a prevacuum sterilizer",
     "Adequacy of the ethylene oxide concentration", "Spore killing efficacy", "Steam superheating only"], 0,
    "The Bowie-Dick is a **Class 2 chemical indicator** run **daily in an EMPTY prevacuum chamber at 134 \u00b0C for "
    "3.5 min**, and it tests **air removal / steam penetration**. It does **not** prove sterility \u2014 only a "
    "**biological indicator** does that.")
mcq("The biological indicator used to monitor steam sterilization is",
    ["Bacillus atrophaeus", "Geobacillus stearothermophilus", "Bacillus pumilus", "Clostridium sporogenes"], 1,
    "**Geobacillus stearothermophilus** (a thermophile, incubated at 55\u201360 \u00b0C) monitors **steam and "
    "hydrogen peroxide plasma**. **Bacillus atrophaeus** monitors **dry heat and ethylene oxide**; "
    "**Bacillus pumilus** monitors **gamma radiation**.")
mcq("Sterilization by hot air oven requires",
    ["121 \u00b0C for 15 min", "134 \u00b0C for 3 min", "160 \u00b0C for 2 hours", "180 \u00b0C for 2 hours"], 2,
    "Dry heat: **160 \u00b0C for 2 h**, 170 \u00b0C for 1 h, or 180 \u00b0C for 30 min. It is the method of choice "
    "for **glassware, oils, powders, greases and sharp cutting instruments** (does not blunt or corrode them), but "
    "not for rubber, plastics or linen.")
mcq("Which item CANNOT be sterilized in a hydrogen peroxide gas plasma (STERRAD) sterilizer?",
    ["A fibre-optic cable", "A stainless-steel instrument", "Cotton gauze and paper drapes",
     "A microsurgical instrument"], 2,
    "**Cellulose-containing materials \u2014 paper, cotton, linen, gauze, dressings \u2014 absorb the sterilant** and "
    "make the cycle ineffective; liquids, powders and long narrow lumens are also unsuitable. Its advantages are a "
    "**low temperature (45\u201350 \u00b0C), a ~1 h cycle, and no aeration** because the by-products are only water "
    "and oxygen.")
mcq("Sterilization with 2 % activated glutaraldehyde requires an immersion time of",
    ["10 minutes", "30 minutes", "3 hours", "10 hours"], 3,
    "**10 hours for sterilization**; **high-level disinfection** is achieved in **20\u201345 minutes** (some texts, "
    "including your PDF, quote 10 minutes). Items must be **rinsed thoroughly in sterile distilled water**, and the "
    "activated solution must be **tested for potency with a test strip** before each use.")
mcq("After use, peracetic acid decomposes into",
    ["Formaldehyde and water", "Acetic acid, water and oxygen", "Chlorine and water", "Ethylene glycol"], 1,
    "Peracetic acid breaks down into **acetic acid (vinegar), water and oxygen**, making it environmentally safe and "
    "allowing disposal down an ordinary drain. It works **even in the presence of organic soil** and has a low surface "
    "tension, so it reaches lumens.")
mcq("According to the Spaulding classification, a flexible gastrointestinal endoscope is a",
    ["Critical item requiring sterilization", "Semi-critical item requiring high-level disinfection",
     "Non-critical item requiring low-level disinfection", "Non-critical item requiring only cleaning"], 1,
    "Flexible endoscopes contact **intact mucous membranes** \u2192 **semi-critical \u2192 high-level disinfection** "
    "(sterilization preferred where feasible). **Critical** items enter sterile tissue or the vascular system and must "
    "be **sterilized**; **non-critical** items touch only intact skin.")
mcq("The internationally accepted Sterility Assurance Level (SAL) is",
    ["10\u207b\u00b3", "10\u207b\u2074", "10\u207b\u2076", "10\u207b\u00b9\u00b2"], 2,
    "**SAL = 10\u207b\u2076** \u2014 a probability of no more than one viable microorganism in one million sterilized "
    "items. This is the quantitative definition of \u2018sterile\u2019.")
mcq("Which of the following organisms is MOST resistant to sterilization and disinfection?",
    ["HIV", "Mycobacterium tuberculosis", "Bacterial spores", "Prions"], 3,
    "Order of decreasing resistance: **prions > bacterial spores > coccidia > mycobacteria > small non-lipid viruses "
    "> fungi > vegetative bacteria > lipid-enveloped viruses (HIV, HBV \u2014 the easiest to kill)**. Prions require "
    "**1 N NaOH plus 134 \u00b0C for 18 min**.")
mcq("The shelf-life of a sterilized wrapped pack is currently determined by",
    ["A fixed 7-day limit", "A fixed 30-day limit", "Event-related sterility", "The type of chemical indicator used"],
    2,
    "**Event-related sterility** \u2014 the pack stays sterile until an event compromises it (wetting, tearing, "
    "puncture, broken seal, dropping, excessive handling). The old time-based \u201csterile for 30 days\u201d rule is "
    "obsolete.")
mcq("Gamma radiation sterilization is typically carried out with a cobalt-60 source at a dose of",
    ["2.5 kGy", "25 kGy", "250 kGy", "2500 kGy"], 1,
    "**25 kGy (2.5 Mrad)**. It is a **cold, industrial-scale** method for pre-packed single-use disposables "
    "(syringes, gloves, catheters, sutures, blades, heart valves, bone grafts) and is **not used within hospitals** "
    "because of cost and shielding requirements.")
mcq("Membrane filtration used to sterilize a heat-labile solution employs a pore size of",
    ["0.45 \u00b5m", "0.22 \u00b5m", "1.2 \u00b5m", "5 \u00b5m"], 1,
    "**0.22 \u00b5m** membranes retain bacteria and are used for heat-labile liquids (sera, antibiotic solutions, "
    "vaccines, media). Note that filtration **removes but does not kill**, and **does not retain viruses or "
    "mycoplasma** (a 0.1 \u00b5m filter is needed for those).")

h2("Section 7 \u2014 Emergencies and Disasters")
mcq("Malignant hyperthermia is triggered by",
    ["Propofol and nitrous oxide", "Volatile inhalational agents and succinylcholine",
     "Non-depolarising muscle relaxants", "All local anaesthetics"], 1,
    "Triggers are **all volatile agents (halothane, isoflurane, sevoflurane, desflurane) and the depolarising "
    "relaxant succinylcholine**. **Safe drugs: propofol, thiopentone, benzodiazepines, opioids, nitrous oxide, "
    "non-depolarising relaxants and all local anaesthetics.**")
mcq("The initial dose of dantrolene in malignant hyperthermia is",
    ["0.5 mg/kg", "1 mg/kg", "2.5 mg/kg", "10 mg/kg immediately"], 2,
    "**2.5 mg/kg IV as a bolus, repeated every 5\u201310 minutes up to 10 mg/kg** (occasionally 30 mg/kg). "
    "Simultaneously: stop the triggers, hyperventilate with 100 % O\u2082 at \u2265 10 L/min, cool actively "
    "(**stop at 38 \u00b0C**), treat hyperkalaemia and arrhythmias, and maintain urine output > 1\u20132 mL/kg/h.")
mcq("The earliest and most sensitive sign of malignant hyperthermia under general anaesthesia is",
    ["Hyperthermia", "An unexplained rise in end-tidal CO\u2082", "Myoglobinuria", "Hypotension"], 1,
    "**Rising EtCO\u2082 with tachypnoea and tachycardia** is the earliest sign; **masseter spasm** after "
    "succinylcholine is the classic clinical clue. **Hyperthermia is a LATE sign** \u2014 waiting for it costs lives. "
    "The defect is in the **RYR1 ryanodine receptor gene on chromosome 19**.")
mcq("The recommended adult intramuscular dose of adrenaline in anaphylaxis is",
    ["0.1 mg of 1:10 000", "0.5 mg of 1:1000 (0.5 mL) into the antero-lateral thigh", "1 mg of 1:1000 subcutaneously",
     "5 mg nebulised"], 1,
    "**0.5 mg IM (0.5 mL of 1:1000) into the antero-lateral thigh, repeated every 5 minutes as needed**; "
    "**50 \u00b5g IV boluses** may be titrated by an experienced practitioner with monitoring. Then "
    "**rapid IV crystalloid 20 mL/kg**, chlorpheniramine, hydrocortisone and salbutamol.")
mcq("Latex anaphylaxis under anaesthesia characteristically differs from drug anaphylaxis in that it",
    ["Never causes hypotension", "Is typically delayed 20\u201360 minutes after exposure",
     "Occurs only in adults", "Responds only to steroids"], 1,
    "Because absorption is across **mucosa or peritoneum**, latex reactions are typically **delayed 20\u201360 "
    "minutes** rather than immediate. Highest-risk groups: **spina bifida, children with multiple operations, and "
    "health-care workers**. Cross-reacting foods: **banana, avocado, chestnut, kiwi**.")
mcq("The components of the fire triangle are",
    ["Fuel, oxidiser and an ignition source", "Fuel, water and heat", "Oxygen, nitrogen and heat",
     "Fuel, carbon dioxide and a spark"], 0,
    "**Fuel + oxidiser + ignition (heat).** In the OR the commonest **ignition** source is the **electrosurgical unit "
    "(\u2248 70\u201390 % of OR fires)**, the commonest **fuel** is **alcohol-based skin prep** (or drapes), and the "
    "**oxidiser** is oxygen or nitrous oxide.")
mcq("For an electrical fire in the operating room, the extinguisher of choice is",
    ["Water", "Foam", "Carbon dioxide (Class C)", "Class D dry powder"], 2,
    "**CO\u2082 (or ABC dry chemical) for Class C electrical fires \u2014 never water.** CO\u2082 is preferred around "
    "equipment because it leaves no residue. Classes: **A** ordinary combustibles, **B** flammable liquids, "
    "**C** electrical, **D** combustible metals, **K** kitchen fats.")
mcq("The acronym RACE, used in the event of a fire, stands for",
    ["Rescue, Alarm, Confine, Extinguish/Evacuate", "Report, Assess, Contain, Escape",
     "Remove, Alert, Cover, Exit", "Rescue, Assess, Call, Evacuate"], 0,
    "**R**escue \u2013 **A**larm \u2013 **C**onfine \u2013 **E**xtinguish/**E**vacuate. To use an extinguisher: "
    "**PASS \u2014 P**ull the pin, **A**im at the base of the fire, **S**queeze, **S**weep.")
mcq("In START triage, a casualty who is not breathing even after the airway has been opened is tagged",
    ["Red", "Yellow", "Green", "Black"], 3,
    "**Black = deceased/expectant.** **Red** = immediate (RR > 30, absent radial pulse or capillary refill > 2 s, or "
    "not obeying commands); **Yellow** = delayed; **Green** = minor/walking wounded. START assesses **RPM \u2014 "
    "Respiration, Perfusion, Mental status** in under a minute per casualty.")

h2("Section 8 \u2014 Post-Operative Care")
mcq("Which of the following is NOT a component of the modified Aldrete score?",
    ["Respiration", "Circulation", "Pain", "Oxygen saturation"], 2,
    "The modified Aldrete score assesses **Activity, Respiration, Circulation, Consciousness and Oxygen saturation "
    "(A-R-C-C-O)**, each 0\u20132, maximum 10, **discharge \u2265 9**. **Pain and nausea/vomiting belong to the "
    "PADSS**, used for discharge home after day surgery.")
mcq("The original Aldrete score used which parameter that the modified score replaced with pulse oximetry?",
    ["Temperature", "Skin colour", "Urine output", "Bowel sounds"], 1,
    "The 1970 Aldrete score used **skin colour**; the modified score substituted **SpO\u2082 measured by pulse "
    "oximetry**, which is objective and more sensitive.")
mcq("The commonest cause of upper airway obstruction in the immediate post-anaesthesia period is",
    ["Laryngospasm", "The tongue falling back against the posterior pharyngeal wall", "Bronchospasm",
     "Vocal-cord paralysis"], 1,
    "Loss of pharyngeal muscle tone in the sedated supine patient allows the **tongue to fall backwards** \u2014 "
    "recognised by snoring and **see-saw (paradoxical) chest movement**. Treat with **head tilt\u2013chin lift or jaw "
    "thrust, lateral position, suction and an oral/nasal airway**.")
mcq("The commonest post-operative pulmonary complication is",
    ["Pneumonia", "Atelectasis", "Pulmonary embolism", "Pneumothorax"], 1,
    "**Atelectasis** \u2014 typically causing low-grade fever on **POD 1\u20132** with reduced basal breath sounds. "
    "Treatment is **deep breathing, incentive spirometry (10 breaths hourly), early ambulation, adequate analgesia "
    "and chest physiotherapy** \u2014 not antibiotics.")
mcq("Fever on the second post-operative day is most likely due to",
    ["Wound infection", "Atelectasis", "Deep vein thrombosis", "Urinary tract infection"], 1,
    "Follow the **W's in time order: Wind (atelectasis, POD 1\u20132) \u2192 Water (UTI, POD 3\u20135) \u2192 "
    "Walking (DVT, POD 4\u20136) \u2192 Wound (SSI, POD 5\u20137) \u2192 Wonder drugs (> POD 7)**. Fever in the first "
    "24 h is usually the inflammatory response, a drug/transfusion reaction, or malignant hyperthermia.")
mcq("Fever appearing on the sixth post-operative day with wound erythema and purulent discharge indicates",
    ["Atelectasis", "Drug fever", "Surgical site infection", "Malignant hyperthermia"], 2,
    "**SSI classically declares itself on POD 5\u20137.** Management is to **open and drain the wound, send a culture, "
    "debride, pack for healing by secondary intention** and give culture-guided antibiotics \u2014 "
    "\u2018the treatment of pus is drainage\u2019.")
mcq("Wound dehiscence most commonly occurs on which post-operative day, and its herald sign is",
    ["POD 1\u20132; bright red bleeding", "POD 5\u201310; serosanguineous \u2018salmon-pink\u2019 discharge",
     "POD 14\u201321; purulent discharge", "POD 30; a palpable lump"], 1,
    "Dehiscence peaks around **POD 5\u201310 (classically day 7\u20138)**, when the wound has only 5\u201310 % of its "
    "final tensile strength; the warning sign is a **sudden serosanguineous (salmon-pink) discharge** with a palpable "
    "gap.")
mcq("The correct immediate nursing action for post-operative wound evisceration is to",
    ["Push the viscera back and apply a dry dressing",
     "Cover the viscera with sterile gauze moistened with warm sterile normal saline, place the patient in low "
     "Fowler's with knees flexed, keep NBM and call the surgeon",
     "Sit the patient upright and give oral fluids", "Apply a tight abdominal binder and discharge"], 1,
    "**Never attempt to replace the viscera and never use dry gauze.** Cover with **warm sterile saline-moistened "
    "gauze**, position supine/low Fowler's with knees flexed to reduce tension, keep the patient **nil by mouth**, "
    "secure IV access, monitor for shock, give analgesia and **prepare for immediate theatre**.")
mcq("Sutures on the face are typically removed on day",
    ["3\u20135", "7\u201310", "10\u201314", "14\u201321"], 0,
    "**Face/neck 3\u20135 days** (excellent vascularity, cosmetic priority); scalp and trunk **7\u201310 days**; "
    "limbs **10\u201314 days**; over joints, back, palms and soles **14 days**.")
mcq("The tensile strength of a surgical wound at one week is approximately what percentage of normal skin?",
    ["5\u201310 %", "30\u201350 %", "70\u201380 %", "100 %"], 0,
    "Only **5\u201310 % at 1 week**, rising to **30\u201350 % at 1 month**, and reaching a maximum of only "
    "**70\u201380 %** of unwounded skin \u2014 the wound never regains full strength.")
mcq("Which cell is the key orchestrating cell of the inflammatory phase of wound healing?",
    ["Neutrophil", "Macrophage", "Fibroblast", "Platelet"], 1,
    "**Neutrophils** arrive first (24\u201348 h), but the **macrophage (48\u201372 h) is the pivotal cell** \u2014 it "
    "debrides and releases the growth factors that recruit fibroblasts. **Fibroblasts** dominate the proliferative "
    "phase (day 4\u201321), depositing **collagen III**, which is later replaced by **collagen I** during remodelling.")
mcq("Which of the following is a CLOSED active suction drain?",
    ["Penrose drain", "Corrugated drain", "Jackson-Pratt drain", "Sump drain with an open vent"], 2,
    "The **Jackson-Pratt (bulb)** and **Hemovac/Redivac (spring/high vacuum)** drains are **closed active** systems "
    "with a lower infection risk. **Penrose and corrugated drains are passive and open.** A fully re-expanded JP bulb "
    "means suction has been lost \u2014 re-compress it after emptying.")
mcq("Post-operative return of bowel function occurs in which order?",
    ["Colon \u2192 stomach \u2192 small bowel", "Small bowel \u2192 stomach \u2192 colon",
     "Stomach \u2192 colon \u2192 small bowel", "All simultaneously at 72 hours"], 1,
    "**Small intestine within \u2248 24 h \u2192 stomach 24\u201348 h \u2192 COLON 48\u201372 h (last to recover).** "
    "The clinical marker of resolution is the **passage of flatus or stool** \u2014 bowel sounds are unreliable.")
mcq("According to the Apfel score, which of the following is NOT a risk factor for PONV?",
    ["Female sex", "Non-smoking status", "Smoking", "Expected post-operative opioid use"], 2,
    "The four Apfel factors are **female sex, non-smoker, previous PONV or motion sickness, and expected "
    "post-operative opioid use** \u2014 giving risks of roughly **10/20/40/60/80 %** for 0\u20134 factors. "
    "**Smoking is actually protective.**")
mcq("The drug of choice given at the END of surgery for PONV prophylaxis is",
    ["Metoclopramide", "Ondansetron", "Dexamethasone", "Atropine"], 1,
    "**Ondansetron 4\u20138 mg IV**, a 5-HT\u2083 antagonist, is most effective given **at the end of surgery** "
    "(watch for **QT prolongation** and headache). **Dexamethasone 4\u20138 mg is given at induction.** "
    "High-risk patients need **agents from 2\u20134 different classes plus propofol TIVA and opioid sparing**.")
mcq("The minimum acceptable post-operative urine output in an adult is",
    ["0.1 mL/kg/h", "0.5 mL/kg/h", "2 mL/kg/h", "5 mL/kg/h"], 1,
    "**> 0.5 mL/kg/h (roughly 30 mL/h in an adult).** Post-operative oliguria is **most often hypovolaemia** \u2014 "
    "assess perfusion and give a fluid challenge \u2014 but beware of causing **dilutional hyponatraemia** with "
    "hypotonic fluid in the ADH-driven post-operative state.")
mcq("Post-operative shivering is most effectively treated with warming plus",
    ["Diazepam", "Pethidine (meperidine)", "Furosemide", "Naloxone"], 1,
    "**Pethidine 12.5\u201325 mg IV** is the most effective drug (via \u03ba-opioid and \u03b12 effects); clonidine, "
    "dexmedetomidine, tramadol and ondansetron are alternatives. **Always give oxygen**, because shivering can "
    "increase oxygen consumption by up to 400 %.")
mcq("Secondary haemorrhage after surgery classically occurs on which post-operative day and is due to",
    ["Day 0; a slipped ligature", "Within 24 h; a dislodged clot as blood pressure rises",
     "POD 7\u201314; infection eroding a vessel wall", "POD 30; a false aneurysm"], 2,
    "**Primary** haemorrhage = during surgery; **reactionary** = within 24 h (slipped ligature or a clot dislodged as "
    "BP rises and vasoconstriction wears off); **secondary = POD 7\u201314 from infection eroding the vessel wall**.")



h2("Section 9 \u2014 Radiation, Chemical Hazards and Biomedical Waste")
mcq("The SI unit of equivalent dose of radiation is the",
    ["Gray", "Sievert", "Becquerel", "Roentgen"], 1,
    "**Sievert (Sv)** is the unit of **equivalent/effective dose** and is used for all dose limits and personnel "
    "monitoring (**1 Sv = 100 rem**). The **Gray** is **absorbed dose** (1 Gy = 100 rad) and the **Becquerel** is "
    "**activity**.")
mcq("The recommended annual occupational whole-body radiation dose limit is",
    ["1 mSv", "20 mSv averaged over 5 years", "50 mSv every year without restriction", "150 mSv"], 1,
    "**20 mSv per year averaged over 5 consecutive years (100 mSv in 5 years), with a maximum of 50 mSv in any single "
    "year.** Public limit **1 mSv/year**; lens of the eye **20 mSv/year**; extremities **500 mSv/year**; and after a "
    "pregnancy is declared, **\u2264 1 mSv to the fetus** (0.5 mSv per month).")
mcq("If the distance from a radiation source is doubled, the exposure received is reduced to",
    ["One half", "One quarter", "One eighth", "One sixteenth"], 1,
    "**Inverse square law: intensity \u221d 1/d\u00b2**, so doubling the distance gives **one quarter** the exposure "
    "and tripling it gives one ninth. Practically, stand **at least 2 m (6 ft)** from the source, or leave the room. "
    "**Time, distance and shielding** are the three cardinal principles, governed by **ALARA**.")
mcq("A personal radiation dosimeter should be worn",
    ["Inside the lead apron at waist level", "Outside the lead apron at chest/collar level",
     "On the wrist of the dominant hand", "Attached to the X-ray machine"], 1,
    "The primary dosimeter is worn **outside the apron at chest or collar level** so that it records the dose to the "
    "unshielded neck, head and lens. A **pregnant** worker wears a **second badge under the apron at waist level**. "
    "**TLD** badges are the commonest type.")
mcq("Carcinogenesis following radiation exposure is an example of a",
    ["Deterministic effect with a threshold", "Stochastic effect with no threshold",
     "Acute deterministic effect", "Non-radiation effect"], 1,
    "**Stochastic** effects (**cancer and heritable mutation**) have **no threshold** \u2014 the **probability** rises "
    "with dose, not the severity. **Deterministic** effects (erythema, epilation, **cataract**, sterility, marrow "
    "suppression) have a **threshold**, and their severity increases with dose.")
mcq("The OSHA permissible exposure limit (8-hour TWA) for ethylene oxide is",
    ["0.05 ppm", "1 ppm", "25 ppm", "100 ppm"], 1,
    "**EO = 1 ppm (8-h TWA)** with a 15-min excursion limit of 5 ppm. Other limits: **glutaraldehyde ceiling "
    "0.05 ppm**, **formaldehyde 0.75 ppm**, **nitrous oxide 25 ppm (NIOSH)**, **halogenated volatiles 2 ppm**, "
    "**methyl methacrylate 100 ppm**.")
mcq("Under the Biomedical Waste Management Rules 2016, a used syringe with a fixed needle should be discarded into "
    "which colour container?",
    ["Yellow bag", "Red bag", "White translucent puncture-proof container", "Blue container"], 2,
    "**Sharps, including needles and syringes with fixed needles, go into the WHITE translucent, puncture-proof, "
    "leak-proof container**, then are autoclaved and shredded/mutilated before encapsulation or foundry disposal. "
    "**Red** = contaminated recyclable plastics (tubing, syringes **without** needles); **Yellow** = anatomical, "
    "soiled, expired drugs and cytotoxics; **Blue** = glassware and metallic implants.")
mcq("Cytotoxic drug waste must be incinerated at a temperature of at least",
    ["800 \u00b0C", "1000 \u00b0C", "1200 \u00b0C", "600 \u00b0C"], 2,
    "**\u2265 1200 \u00b0C**, or the waste is returned to the manufacturer. Cytotoxic waste is collected in a "
    "**separate yellow container labelled \u2018cytotoxic\u2019** \u2014 including glass vials contaminated with "
    "cytotoxics, which therefore do **not** go into the blue stream.")

# ================================================================ PART G
h1("Part G  \u2014  Rapid Revision: Mnemonics and Last-Minute Sheet")

h2("Mnemonics")
table(["Mnemonic", "Stands for"],
      [["**2 \u2013 4 \u2013 6 \u2013 8**", "Pre-op fasting: clear fluids 2 h, breast milk 4 h, formula/light meal "
                                           "6 h, fatty meal 8 h"],
       ["**7 A's**", "Premedication aims: **A**nxiolysis, **A**mnesia, **A**nalgesia, **A**ntisialagogue, "
                     "**A**nti-emetic, **A**nti-acid, **A**utonomic attenuation (+ **A**ntibiotic prophylaxis)"],
       ["**4 G's + E**", "Herbals to stop 2 weeks pre-op: **G**arlic, **G**inger, **G**inkgo, **G**inseng, "
                         "**E**phedra (+ St John's wort, vitamin E, fish oil)"],
       ["**Sign In \u2013 Time Out \u2013 Sign Out**", "The 3 phases of the WHO Surgical Safety Checklist"],
       ["**A-R-C-C-O**", "Aldrete score: **A**ctivity, **R**espiration, **C**irculation, **C**onsciousness, "
                         "**O**xygen saturation \u2014 max 10, discharge \u2265 9"],
       ["**Wind \u2013 Water \u2013 Walking \u2013 Wound \u2013 Wonder drugs**",
        "Post-op fever by day: 1\u20132 \u2192 3\u20135 \u2192 4\u20136 \u2192 5\u20137 \u2192 > 7"],
       ["**Virchow's triad**", "**S**tasis + **E**ndothelial injury + **H**ypercoagulability \u2192 VTE"],
       ["**R \u2013 A \u2013 C \u2013 E**", "Fire: **R**escue, **A**larm, **C**onfine, **E**xtinguish/Evacuate"],
       ["**P \u2013 A \u2013 S \u2013 S**", "Extinguisher: **P**ull, **A**im at the base, **S**queeze, **S**weep"],
       ["**R \u2013 P \u2013 M**", "START triage: **R**espiration, **P**erfusion, **M**ental status"],
       ["**4 H's & 4 T's**", "Reversible causes of arrest: Hypoxia, Hypovolaemia, Hypo/Hyperkalaemia-metabolic, "
                             "Hypothermia; Tension pneumothorax, Tamponade, Toxins, Thrombosis"],
       ["**4 D's**", "Negligence: **D**uty, **D**ereliction, **D**irect causation, **D**amage"],
       ["**Critical \u2013 Semi-critical \u2013 Non-critical**", "Spaulding: sterilize \u2013 high-level disinfect "
                                                                "\u2013 low-level disinfect"],
       ["**Geo = Steam** ; **Atro = Air & EO**", "**Geo**bacillus stearothermophilus \u2192 **steam** & plasma; "
                                                 "Bacillus **atro**phaeus \u2192 dry heat (air) & **E**thylene oxide"],
       ["**\u2018Yellow burns, Red recycles, White is sharp, Blue is glass\u2019**",
        "Biomedical waste colour coding (BMW Rules 2016)"],
       ["**Time \u2013 Distance \u2013 Shielding**", "The three cardinal principles of radiation protection (ALARA)"],
       ["**\u2018If the gut works, use it\u2019**", "Enteral is always preferred to parenteral nutrition"],
       ["**\u2018The treatment of pus is drainage\u2019**", "Surgical site infection management"],
       ["**\u2018When in doubt, throw it out\u2019**", "An item of doubtful sterility is considered unsterile"],
       ["**\u2018Only what is documented is a legal fact\u2019**", "Perioperative documentation"]],
      widths=[5.4, 13.6])

h2("The 30 statements most likely to appear as a one-mark question")
numbered([
    "The **circulator** keeps the legal record of the count and acts as the **patient's advocate**.",
    "**At least two identifiers** are used; the patient **states** their own name rather than being asked "
    "\u201cAre you Mary Jones?\u201d",
    "The site is marked by the **operating surgeon** with an **indelible marker**, with the patient participating.",
    "The **\u2018time out\u2019** occurs **after induction and before skin incision**, and every team member must "
    "participate.",
    "**Hearing is the last sense to be lost** \u2014 silence is mandatory during induction.",
    "The **safety strap across the mid-thighs** is the single most important measure against falls from the table.",
    "**Ulnar nerve** injury is the commonest peri-operative nerve injury; **common peroneal** in lithotomy.",
    "**Trendelenburg raises ICP and IOP and lowers FRC**; do not maintain it longer than necessary.",
    "**Never move an anaesthetised patient without the anaesthetist's permission.**",
    "The **patient's skin must never contact metal** when the ESU is used (alternate-site burn).",
    "The **electrosurgical pencil is kept in its holster** when not in use.",
    "ESU works at **300 kHz\u20133 MHz**; **bipolar needs no dispersive pad** and is safest with pacemakers.",
    "**Autoclave 121 \u00b0C / 15 psi / 15\u201320 min**; **hot air oven 160 \u00b0C / 2 h**.",
    "The **Bowie-Dick test checks air removal** in a prevacuum steriliser; the **biological indicator is the only "
    "proof of sterilizing efficacy**.",
    "**Ethylene oxide requires aeration**; **hydrogen peroxide plasma cannot be used for cellulose**.",
    "**SAL = 10\u207b\u2076**; sterility is now **event-related, not time-related**.",
    "**HEPA = 99.97 % at 0.3 \u00b5m**; the OR is under **positive** pressure with **\u2265 15 air changes/hour**.",
    "**Hand hygiene is the single most effective measure** against healthcare-associated infection.",
    "A **draped table is sterile only at table-top level**; the wrapper edge (**1 inch**) is unsterile; "
    "**moisture = contamination**.",
    "**X-ray-detectable sponges are never used as dressings**, and **nothing leaves the OR until the final count is "
    "correct**.",
    "**Malignant hyperthermia** \u2014 triggers are **volatile agents and succinylcholine**; treat with "
    "**dantrolene 2.5 mg/kg**; the earliest sign is a **rising EtCO\u2082**.",
    "**Anaphylaxis \u2014 adrenaline 0.5 mg IM (1:1000)**; **latex reactions are typically delayed 20\u201360 min**.",
    "**CPR: 100\u2013120/min, 5\u20136 cm deep, 30:2; VF and pulseless VT are the only shockable rhythms.**",
    "The commonest **ignition** source of OR fires is the **ESU**; the commonest **fuel** is **alcohol-based prep**.",
    "**Aldrete \u2265 9** for PACU discharge; **PADSS \u2265 9** for discharge home; **pain is in PADSS, not "
    "Aldrete**.",
    "The commonest immediate PACU problem is **airway obstruction from the tongue**; the commonest post-op pulmonary "
    "complication is **atelectasis**.",
    "**Post-op fever: Wind \u2192 Water \u2192 Walking \u2192 Wound \u2192 Wonder drugs.** Fever in the first 24 h "
    "is usually **not** infective.",
    "**Dehiscence POD 5\u201310**, heralded by **serosanguineous discharge**; **evisceration \u2192 cover with warm "
    "sterile saline-soaked gauze and call the surgeon**.",
    "The **colon is the last part of the gut to recover** (48\u201372 h); return of function is judged by **passage "
    "of flatus**.",
    "**Radiation: 20 mSv/year occupational; doubling the distance quarters the exposure; the dosimeter is worn "
    "OUTSIDE the apron.**",
])

divider()
fp = doc.add_paragraph()
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.space_before = Pt(3)
r = fp.add_run("End of supplementary notes  \u2022  Revise the PDF for Chapters 1\u20133 core text, "
               "and this document for every gap listed in the coverage audit on page 1.")
r.italic = True
r.font.size = Pt(9)
r.font.color.rgb = C_H1

# footer
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer_p.add_run("Supplementary Notes \u2014 Perioperative Patient Care & OT Safety  |  "
                      "gap-fill for Chapters 1\u20133  |  page ")
fr.font.size = Pt(7.5)
fr.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer_p._p.append(fld)

out = "/projects/sandbox/DeepFocus/Perioperative_Care_Supplementary_Notes.docx"
doc.save(out)
print("Saved:", out)
print("MCQs generated:", MCQ_N[0])
print("Paragraphs:", len(doc.paragraphs), "Tables:", len(doc.tables))
