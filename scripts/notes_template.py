#!/usr/bin/env python3
"""
Reusable DOCX study-notes template.

Layout spec: LEGAL page (8.5 x 14 in), all margins 1.27 cm, compact spacing,
coloured headings (H1 blue / H2 green / H3 purple), shaded tables with
alternating rows, highlight boxes, MCQ blocks with coloured answers.

Usage:
    from notes_template import *
    h1("Section"); para("text with **bold** and __italic__"); table([...], [...])
    save("/path/out.docx")
"""

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

# ---------------------------------------------------------------- palette
C_H1 = RGBColor(0x1F, 0x3B, 0x73)      # deep blue
C_H2 = RGBColor(0x1B, 0x6B, 0x3A)      # green
C_H3 = RGBColor(0x6A, 0x28, 0x8F)      # purple
C_ANS = RGBColor(0x9C, 0x27, 0x10)     # rust (answers)
C_TXT = RGBColor(0x20, 0x20, 0x20)
C_MUTE = RGBColor(0x55, 0x55, 0x55)

SH_HDR = "1F3B73"      # table header fill
SH_ALT = "EEF3FA"      # alternating row fill
SH_CLIN = "FFF4E5"     # clinical box
SH_MNE = "E8F5E9"      # mnemonic box
SH_HY = "FDECEF"       # high-yield box
SH_GAP = "F3E9FB"      # gap-alert box

BODY = 9.5
TBL = 8.0

USABLE_CM = 19.0       # legal width (21.59) minus 2 x 1.27 cm margins

# ---------------------------------------------------------------- document
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(14)
for _m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(sec, _m, Cm(1.27))

_st = doc.styles["Normal"]
_st.font.name = "Calibri"
_st.font.size = Pt(BODY)
_st.font.color.rgb = C_TXT
_st.paragraph_format.space_before = Pt(0)
_st.paragraph_format.space_after = Pt(1.5)
_st.paragraph_format.line_spacing = 1.0
_st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

MCQ_N = [0]


# ---------------------------------------------------------------- internals
def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _cell_margins(t, top=12, bottom=12, left=60, right=60):
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement("w:" + tag)
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    t._tbl.tblPr.append(mar)


def _borders(t, color="9DB3D0", sz=4):
    b = OxmlElement("w:tblBorders")
    for tag in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + tag)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), color)
        b.append(e)
    t._tbl.tblPr.append(b)


def _rule(p, color, sz="12"):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single")
    bt.set(qn("w:sz"), sz)
    bt.set(qn("w:space"), "2")
    bt.set(qn("w:color"), color)
    pbdr.append(bt)
    pPr.append(pbdr)


# ---------------------------------------------------------------- public API
def rich(p, text, size=BODY, color=None):
    """Inline markup: **bold**, __italic__, ~~bold-italic~~"""
    for t in re.split(r"(\*\*.+?\*\*|__.+?__|~~.+?~~)", text):
        if not t:
            continue
        if t.startswith("**") and t.endswith("**"):
            r = p.add_run(t[2:-2]); r.bold = True
        elif t.startswith("__") and t.endswith("__"):
            r = p.add_run(t[2:-2]); r.italic = True
        elif t.startswith("~~") and t.endswith("~~"):
            r = p.add_run(t[2:-2]); r.bold = True; r.italic = True
        else:
            r = p.add_run(t)
        r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
    return p


def title(main, sub=None, note=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(main)
    r.bold = True; r.font.size = Pt(17); r.font.color.rgb = C_H1
    if sub:
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_after = Pt(1)
        r = q.add_run(sub)
        r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = C_H2
    if note:
        n = doc.add_paragraph()
        n.alignment = WD_ALIGN_PARAGRAPH.CENTER
        n.paragraph_format.space_after = Pt(4)
        r = n.add_run(note)
        r.italic = True; r.font.size = Pt(9); r.font.color.rgb = C_MUTE


def h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    r.font.size = Pt(14.5); r.bold = True; r.font.color.rgb = C_H1
    _rule(p, "1F3B73")
    return p


def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.size = Pt(11.5); r.bold = True; r.font.color.rgb = C_H2
    return p


def h3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.size = Pt(10.5); r.bold = True; r.font.color.rgb = C_H3
    return p


def para(text, size=BODY, before=0, after=1.5, justify=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, text, size=size)
    return p


def bullets(items, size=BODY, indent=0.42, marker="\u25aa"):
    for it in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(indent)
        p.paragraph_format.first_line_indent = Cm(-0.28)
        p.paragraph_format.space_after = Pt(0.8)
        rich(p, marker + "  " + it, size=size)


def numbered(items, size=BODY, indent=0.5):
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(indent)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.space_after = Pt(0.8)
        rich(p, "**%d.**  %s" % (i, it), size=size)


def table(headers, rows, widths=None, size=TBL, header_fill=SH_HDR):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _borders(t)
    _cell_margins(t)
    for i, htxt in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        _shade(c, header_fill)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(htxt)
        r.bold = True; r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            c = cells[ci]
            c.text = ""
            if ri % 2 == 1:
                _shade(c, SH_ALT)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            rich(p, str(val), size=size)
    if widths:
        # narrower tables are fine; overflowing the text block is not
        assert sum(widths) <= USABLE_CM + 0.3, \
            "column widths sum to %.1f cm, usable width is %.1f cm" % (sum(widths), USABLE_CM)
        for ci, w in enumerate(widths):
            for row in t.rows:
                row.cells[ci].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def box(heading, lines, fill=SH_CLIN, border="D9A15B", title_color=None):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    _borders(t, color=border, sz=8)
    _cell_margins(t, top=40, bottom=40, left=100, right=100)
    c = t.rows[0].cells[0]
    c.text = ""
    _shade(c, fill)
    c.width = Cm(USABLE_CM)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(heading)
    r.bold = True; r.font.size = Pt(9.5)
    r.font.color.rgb = title_color or C_H1
    for ln in lines:
        q = c.add_paragraph()
        q.paragraph_format.space_after = Pt(0.5)
        q.paragraph_format.left_indent = Cm(0.15)
        rich(q, ln, size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def clinical(heading, lines):
    return box(heading, lines, fill=SH_CLIN, border="D9A15B")


def mnemonic(heading, lines):
    return box(heading, lines, fill=SH_MNE, border="6DA97A", title_color=C_H2)


def highyield(heading, lines):
    return box(heading, lines, fill=SH_HY, border="D98BA0", title_color=C_ANS)


def alert(heading, lines):
    return box(heading, lines, fill=SH_GAP, border="9B6FC4", title_color=C_H3)


def divider():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    _rule(p, "B9A0D6", sz="6")


def mcq(q, opts, ans, expl):
    """opts: list of option strings; ans: 0-based index of the correct option."""
    MCQ_N[0] += 1
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0.5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("Q%d. " % MCQ_N[0])
    r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = C_H1
    rich(p, q, size=9.5)
    letters = "abcd"
    o = doc.add_paragraph()
    o.paragraph_format.left_indent = Cm(0.45)
    o.paragraph_format.space_after = Pt(0.5)
    o.paragraph_format.keep_with_next = True
    rich(o, "   ".join("(%s) %s" % (letters[i], opts[i]) for i in range(len(opts))), size=9)
    a = doc.add_paragraph()
    a.paragraph_format.left_indent = Cm(0.45)
    a.paragraph_format.space_after = Pt(2)
    ra = a.add_run("Ans: (%s) %s \u2014 " % (letters[ans], opts[ans]))
    ra.bold = True; ra.font.size = Pt(9); ra.font.color.rgb = C_ANS
    rich(a, expl, size=9, color=C_ANS)


def footer(text):
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(text + "  |  page ")
    r.font.size = Pt(7.5); r.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)


def closing(text):
    divider()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    r = p.add_run(text)
    r.italic = True; r.font.size = Pt(9); r.font.color.rgb = C_H1


def save(path):
    doc.save(path)
    words = sum(len(p.text.split()) for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                words += len(c.text.split())
    print("Saved   :", path)
    print("MCQs    :", MCQ_N[0])
    print("Tables  :", len(doc.tables))
    print("Words   : ~%d" % words)
